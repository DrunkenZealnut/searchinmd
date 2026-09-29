"""highlight_pdf_occurrences.py 단위 시험 — python3.13 -m unittest test_highlight_pdf_occurrences

pymupdf(fitz)·openpyxl 이 있는 인터프리터가 필요하다. 개인 자료(data/)는 읽지 않는다 — 모든 PDF 는 fitz 로 만든다.
Vision OCR 통합(OcrToolTests)은 swiftc 가 있을 때만 돈다.
"""
from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path

import fitz

import highlight_pdf_occurrences as HL
import resegment
import semantic_keyword_recount as skr


# ----------------------------------------------------------------------------- 2. 정규화·글자 스트림
class NormTests(unittest.TestCase):
    def test_strip_pattern_equals_resegment(self):
        self.assertEqual(HL.STRIP.pattern, resegment._STRIP.pattern)

    def test_norm_chars_drops_space_markdown_punctuation_and_casefolds(self):
        self.assertEqual(HL.norm_chars("| **안전** 수칙: MSDS(물질) |"), "안전수칙msds물질")

    def test_norm_chars_is_nfc(self):
        decomposed = "한"       # 한 (NFD 자모 3개)
        self.assertEqual(HL.norm_chars(decomposed), "한")

    def test_char_stream_maps_each_normalized_char_to_its_box(self):
        r = lambda x: fitz.Rect(x, 0, x + 10, 10)
        chars = [("안", r(0), 1), (" ", r(10), 1), ("전", r(20), 1), ("|", r(30), 1), ("A", r(40), 2)]
        s = HL.CharStream.from_chars(chars)
        self.assertEqual(s.text, "안전a")
        self.assertEqual([b.x0 for b in s.boxes], [0, 20, 40])
        self.assertEqual(s.line_ids, [1, 1, 2])

    def test_char_stream_casefold_expansion_shares_the_box(self):
        s = HL.CharStream.from_chars([("ß", fitz.Rect(0, 0, 5, 5), 0)])
        self.assertEqual(s.text, "ss")
        self.assertEqual(len(s.boxes), 2)
        self.assertEqual(s.boxes[0], s.boxes[1])

    def test_from_rawdict_reads_a_real_page(self):
        doc = fitz.open()
        page = doc.new_page(width=300, height=200)
        page.insert_text((20, 50), "안전 수칙", fontname="korea", fontsize=14)
        s = HL.CharStream.from_rawdict(page)
        self.assertEqual(s.text, "안전수칙")
        hit = page.search_for("안전")[0]
        self.assertTrue(hit.intersects(s.boxes[0]))
        self.assertTrue(hit.intersects(s.boxes[1]))

    def test_from_ocr_scales_pixels_to_page_points_and_falls_back_to_line_box(self):
        doc = fitz.open()
        page = doc.new_page(width=200, height=100)          # 1 pt = 2 px at 400x200
        ocr = {"width": 400, "height": 200, "lines": [
            {"text": "안전", "conf": 1.0, "box": [20, 40, 60, 60],
             "chars": [{"c": "안", "box": [20, 40, 40, 60]}, {"c": "전", "box": None}]}]}
        s = HL.CharStream.from_ocr(ocr, page)
        self.assertEqual(s.text, "안전")
        self.assertEqual(s.boxes[0], fitz.Rect(10, 20, 20, 30))
        self.assertEqual(s.boxes[1], fitz.Rect(10, 20, 30, 30))       # 글자 box 없음 → 줄 box
        self.assertEqual(s.line_ids, [0, 0])


# ----------------------------------------------------------------------------- 3. PDF 대응
class PdfIndexTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def touch(self, rel: str) -> Path:
        p = self.tmp / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(b"")
        return p

    def test_ncs_index_maps_every_code_in_the_file_name(self):
        a = self.touch("ncs/반도체장비/LM1903060314_18v2, LM1903060315_18v2_반도체 장비 보드 설계.pdf")
        b = self.touch("ncs/반도체재료/LM1903060401_20v2 반도체용 웨이퍼 재료 제조.pdf")
        self.touch("ncs/반도체재료/.DS_Store")
        idx = HL.index_ncs_pdfs(self.tmp / "ncs")
        self.assertEqual(idx, {"LM1903060314": a, "LM1903060315": a, "LM1903060401": b})

    def test_ncs_index_refuses_a_code_in_two_files(self):
        self.touch("ncs/a/LM1903060401_x.pdf")
        self.touch("ncs/b/LM1903060401_y.pdf")
        with self.assertRaises(LookupError):
            HL.index_ncs_pdfs(self.tmp / "ncs")

    def test_textbook_key_strips_timestamp_and_punctuation(self):
        self.assertEqual(HL.textbook_key("20260415_143535_반도체인프라일반_서울시교육청_/20260415_143535_반도체인프라일반_서울시교육청_.md"),
                         "반도체인프라일반서울시교육청")
        self.assertEqual(HL.textbook_key("daegu/20260415_200404_반도체공정기초_렛유인_/20260415_200404_반도체공정기초_렛유인_.md"),
                         "반도체공정기초렛유인")

    def test_textbook_index_keys_by_normalized_name_so_duplicates_stay_apart(self):
        a = self.touch("tb/반도체 포토에칭(에이치앤지).pdf")
        self.touch("tb/반도체산업의유해인자(에피스테메).pdf")
        self.touch("tb/반도체산업의유해인자(에피스테메) 2.pdf")
        idx = HL.index_textbook_pdfs(self.tmp / "tb")
        self.assertEqual(idx["반도체포토에칭에이치앤지"], a)
        self.assertIn("반도체산업의유해인자에피스테메2", idx)

    def test_pdf_for_resolves_both_corpora_and_fails_loudly(self):
        ncs = self.touch("ncs/반도체개발/LM1903060101_23v6_반도체 제품 기획.pdf")
        tb = self.touch("tb/반도체기초기술1(크리아트).pdf")
        ncs_idx, tb_idx = HL.index_ncs_pdfs(self.tmp / "ncs"), HL.index_textbook_pdfs(self.tmp / "tb")
        self.assertEqual(HL.pdf_for("NCS", "반도체개발/LM1903060101_23v6_반도체_제품_기획/x_LM1903060101_y.md", ncs_idx, tb_idx), ncs)
        self.assertEqual(HL.pdf_for("교과서", "20260413_171220_반도체기초기술1_크리아트_/a.md", ncs_idx, tb_idx), tb)
        with self.assertRaises(LookupError):
            HL.pdf_for("NCS", "반도체개발/LM1903060199_없음/a.md", ncs_idx, tb_idx)
        with self.assertRaises(LookupError):
            HL.pdf_for("교과서", "20260413_171220_없는교과서_/a.md", ncs_idx, tb_idx)

    def test_output_path_keeps_the_area_folder_and_adds_the_suffix(self):
        ncs = self.touch("ncs/반도체개발/LM1903060101_23v6_반도체 제품 기획.pdf")
        out = HL.output_path("NCS", ncs, self.tmp / "ncs", self.tmp / "out")
        self.assertEqual(out, self.tmp / "out" / "ncs" / "반도체개발" / "LM1903060101_23v6_반도체 제품 기획_키워드표시.pdf")
        tb = self.touch("tb/반도체기초기술1(크리아트).pdf")
        self.assertEqual(HL.output_path("교과서", tb, self.tmp / "tb", self.tmp / "out"),
                         self.tmp / "out" / "school-text" / "반도체기초기술1(크리아트)_키워드표시.pdf")


# ----------------------------------------------------------------------------- 4. 줄 안 위치 재현
RULES = skr.build_default_rules(list(skr.EXPECTED_KEYWORDS), version=skr.DEFAULT_DICTIONARY)


def _rule(expression: str) -> skr.ExpressionRule:
    return next(r for r in RULES if r.expression == expression)


class LocatorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.compiled = HL.compile_rules(RULES)

    def located(self, line: str, window: str | None = None):
        return [(l.keyword, l.matched_text, l.start, l.end) for l in HL.locate_included(line, window if window is not None else line, self.compiled)]

    def test_held_expression_inside_the_line_is_skipped(self):
        self.assertEqual(self.located("안전성 검토 후 안전 수칙을 지킨다"), [("안전", "안전", 9, 11)])

    def test_exclusion_gives_the_span_to_the_owning_keyword(self):
        out = self.located("물질안전보건자료를 비치한다")
        self.assertEqual([(m, s, e) for _, m, s, e in out], [("물질안전보건자료", 0, 8)])
        self.assertNotEqual(out[0][0], "안전")

    def test_conditional_expression_needs_a_companion_in_the_window(self):
        rule = _rule("장갑")
        self.assertTrue(rule.require_patterns, "장갑 은 조건부 포함 표현이어야 한다 (사전 v2)")
        self.assertEqual(self.located("장갑 두 켤레를 산다"), [])                                   # 동반어(착용·보호·안전…) 없음
        self.assertEqual(self.located("장갑 두 켤레를 산다", "안전을 위해\n장갑 두 켤레를 산다"), [(rule.keyword, "장갑", 0, 2)])

    def test_keyword_order_follows_the_rule_order_and_offsets_are_exact(self):
        out = self.located("안전 교육과 화재 예방")
        self.assertEqual(out, [("안전", "안전", 0, 2), ("화재", "화재", 7, 9)])
        order = list(skr.EXPECTED_KEYWORDS)
        self.assertLess(order.index("안전"), order.index("화재"))

    def test_two_instances_of_one_expression_keep_document_order(self):
        self.assertEqual(self.located("안전 수칙과 안전 장비"), [("안전", "안전", 0, 2), ("안전", "안전", 7, 9)])

    def test_line_offsets_walks_page_blocks_with_the_engine_window(self):
        text = "<!-- page: 1 -->\n안전 수칙\n\n장갑 구매\n<!-- page: 2 -->\n안전을 위한 장갑 구매"
        doc = skr.Document(corpus="NCS", path=Path("x.md"), relative_path="x.md", text=text)
        offsets = HL.line_offsets(doc, self.compiled, wanted_lines={2, 4, 6})
        self.assertEqual({k: [(l.matched_text, l.start) for l in v] for k, v in offsets.items()},
                         {2: [("안전", 0)], 4: [], 6: [("안전", 0), ("장갑", 7)]})

    def test_assign_offsets_uses_engine_when_sequences_agree_else_naive_kth_instance(self):
        rows = [HL.Row("NCS", "x.md", "안전", "안전", "exact", "안전", 2, 1, "안전 수칙과 안전 장비", 1, "미흡·없음", 0),
                HL.Row("NCS", "x.md", "안전", "안전", "exact", "안전", 2, 1, "안전 수칙과 안전 장비", 1, "미흡·없음", 1),
                HL.Row("NCS", "x.md", "안전", "안전", "exact", "안전", 3, 1, "안전성 검사와 안전 수칙", 1, "미흡·없음", 2)]
        offsets = {2: [HL.Located("안전", "안전", "exact", "안전", 0, 2), HL.Located("안전", "안전", "exact", "안전", 8, 10)],
                   3: [HL.Located("안전", "안전", "exact", "안전", 9, 11), HL.Located("안전", "안전", "exact", "안전", 12, 14)]}   # 3행과 어긋남
        counts = HL.assign_offsets(rows, offsets, raw_lines={2: "안전 수칙과 안전 장비", 3: "안전성 검사와 안전 수칙"})
        self.assertEqual([(r.start, r.end, r.locator) for r in rows], [(0, 2, "engine"), (8, 10, "engine"), (0, 2, "naive")])
        self.assertEqual(counts, {"engine": 1, "naive": 1})

    def test_assign_offsets_matches_by_keyword_and_text_when_the_xlsx_is_sorted_differently(self):
        # xlsx 는 같은 줄의 행을 (키워드, 실제 매칭) 정렬로 쓴다; 엔진은 규칙 순서 — 다중집합이 같으면 engine 으로 대응한다
        rows = [HL.Row("NCS", "x.md", "MSDS", "MSDS", "exact", "MSDS", 7, 1, "...", 2, "형식적 언급", 0),
                HL.Row("NCS", "x.md", "위험", "위험", "equivalent", "Hazard", 7, 1, "...", 2, "형식적 언급", 1),
                HL.Row("NCS", "x.md", "위험", "위험", "exact", "위험", 7, 1, "...", 2, "형식적 언급", 2),
                HL.Row("NCS", "x.md", "위험", "위험", "exact", "위험", 7, 1, "...", 2, "형식적 언급", 3)]
        offsets = {7: [HL.Located("위험", "위험", "exact", "위험", 3, 5), HL.Located("위험", "위험", "exact", "위험", 9, 11),
                       HL.Located("위험", "Hazard", "equivalent", "Hazard", 20, 26), HL.Located("MSDS", "MSDS", "exact", "MSDS", 30, 34)]}
        counts = HL.assign_offsets(rows, offsets, raw_lines={7: "raw"})
        self.assertEqual([(r.start, r.end, r.locator) for r in rows], [(30, 34, "engine"), (20, 26, "engine"), (3, 5, "engine"), (9, 11, "engine")])
        self.assertEqual(counts, {"engine": 1, "naive": 0})


# ----------------------------------------------------------------------------- 5. 문맥 창 검색
def _stream(text: str) -> HL.CharStream:
    return HL.CharStream.from_chars([(c, fitz.Rect(i * 10, 0, i * 10 + 10, 10), 0) for i, c in enumerate(text)])


class SearchTests(unittest.TestCase):
    def test_context_window_normalizes_both_sides(self):
        self.assertEqual(HL.context_window("| **안전** 수칙: 준수 |", 4, 6, 3), ("", "안전", "수칙준"))
        self.assertEqual(HL.context_window("가나다 안전 라", 4, 6, 0), ("", "안전", ""))

    def test_find_span_returns_the_core_span_inside_the_window_hit(self):
        self.assertEqual(HL.find_span("안전수칙안전장비", "", "안전", "장비", set()), (4, 6))

    def test_find_span_skips_spans_already_used(self):
        self.assertEqual(HL.find_span("안전수칙안전장비", "", "안전", "", {(0, 2)}), (4, 6))
        self.assertIsNone(HL.find_span("안전수칙안전장비", "", "안전", "", {(0, 2), (4, 6)}))

    def test_resolve_on_page_shrinks_the_window_until_it_hits(self):
        s = _stream("가나다안전라마")
        self.assertEqual(HL.resolve_on_page(s, "xx 안전 yy", 3, 5, set()), (3, 5, 0))          # 앞뒤가 안 맞아 문맥 0자로 맞춘다
        self.assertEqual(HL.resolve_on_page(s, "x가나다 안전 라마", 5, 7, set()), (3, 5, 5))    # 창 3 에서 처음 맞는다: 가나다+라마 = 5자
        self.assertEqual(HL.resolve_on_page(s, "나다 안전 라마", 3, 5, set()), (3, 5, 4))       # 짧은 줄은 창 10 이 그대로 맞는다: 나다+라마 = 4자
        self.assertIsNone(HL.resolve_on_page(s, "xx 없음 yy", 3, 5, set()))

    def test_resolve_row_tries_the_page_then_minus_one_then_plus_one(self):
        streams = {9: _stream("안전"), 10: _stream("없음"), 11: _stream("안전")}
        used = {}
        r = HL.resolve_row(streams.get, 10, "안전", 0, 2, "안전", used)
        self.assertEqual((r.page, r.start, r.end, r.k), (9, 0, 2, 0))
        r2 = HL.resolve_row(streams.get, 10, "안전", 0, 2, "안전", used)
        self.assertEqual(r2.page, 11)                                       # 9쪽 것은 이미 썼다
        r3 = HL.resolve_row(streams.get, 10, "안전", 0, 2, "안전", used)
        self.assertEqual((r3.page, r3.shared), (9, True))                   # 양쪽 다 소진 → 최후 수단: 이웃 쪽에 겹쳐 표시
        self.assertIsNone(HL.resolve_row(streams.get, 5, "안전", 0, 2, "안전", {}))          # 5·4·6쪽 스트림이 없다

    def test_used_spans_are_kept_per_keyword(self):
        streams = {10: _stream("안전보건")}
        used = {}
        first = HL.resolve_row(streams.get, 10, "안전", 0, 2, "안전", used)
        self.assertEqual((first.start, first.shared), (0, False))
        other = HL.resolve_row(streams.get, 10, "안전", 0, 2, "안전보건", used)
        self.assertEqual((other.start, other.shared), (0, False))                                  # 다른 키워드는 겹쳐도 된다
        again = HL.resolve_row(streams.get, 10, "안전", 0, 2, "안전", used)
        self.assertEqual((again.page, again.start, again.shared), (10, 0, True))                  # 소진되면 같은 자리 중복 표시

    def test_resolve_row_shares_on_the_canonical_page_instead_of_moving_to_a_neighbour(self):
        streams = {9: _stream("안전"), 10: _stream("안전"), 11: _stream("안전")}
        used = {}
        first = HL.resolve_row(streams.get, 10, "안전", 0, 2, "안전", used)
        second = HL.resolve_row(streams.get, 10, "안전", 0, 2, "안전", used)
        self.assertEqual((first.page, first.shared), (10, False))
        self.assertEqual((second.page, second.shared), (10, True))                                # 9·11쪽에 자리가 있어도 옮기지 않는다

    def test_far_search_needs_strong_context_and_reports_the_distance(self):
        streams = {9: _stream("없음"), 10: _stream("없음"), 11: _stream("없음"), 13: _stream("가나다라안전마바사아")}
        r = HL.resolve_row(streams.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {})
        self.assertEqual((r.page, r.start, r.end, r.k, r.shared, r.far), (13, 4, 6, 8, False, True))
        self.assertIsNone(HL.resolve_row(streams.get, 10, "xx 안전 yy", 3, 5, "안전", {}))         # 문맥 2자 — 먼 쪽엔 안 간다

    def test_far_search_takes_the_nearest_page_and_stops_at_the_limit(self):
        far = HL.FAR_MAX_PAGES
        streams = {10: _stream("없음"), 10 - 2: _stream("가나다라안전마바사아"), 10 + far: _stream("가나다라안전마바사아"),
                   10 + far + 1: _stream("가나다라안전마바사아")}
        self.assertEqual(HL.resolve_row(streams.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}).page, 8)
        streams.pop(8)
        self.assertEqual(HL.resolve_row(streams.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}).page, 10 + far)
        streams.pop(10 + far)
        self.assertIsNone(HL.resolve_row(streams.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}))

    def test_strong_neighbour_hit_beats_sharing_on_the_canonical_page(self):
        streams = {10: _stream("안전"), 11: _stream("가나다라안전마바사아")}
        used = {(10, "안전"): {(0, 2)}}                                                       # 정본 쪽 자리는 이미 소진
        r = HL.resolve_row(streams.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", used)
        self.assertEqual((r.page, r.k, r.shared, r.far), (11, 8, False, False))            # 문맥 6자+ 가 이웃 쪽에 → 옮긴다
        weak = HL.resolve_row(streams.get, 10, "xx 안전 yy", 3, 5, "안전", used)
        self.assertEqual((weak.page, weak.shared), (10, True))                              # 문맥 근거 없으면 정본 쪽 중복

    def test_ocr_fallback_on_the_canonical_page_outranks_far_only_with_strong_context(self):
        text_layer = {10: _stream("없음"), 13: _stream("가나다라안전마바사아")}
        ocr_strong = {10: _stream("가나다라안전마바사아")}
        r = HL.resolve_row(text_layer.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}, get_ocr_stream=ocr_strong.get)
        self.assertEqual((r.page, r.k, r.ocr, r.far), (10, 8, True, False))              # 정본 쪽 OCR 강한 적중 > 먼 쪽
        ocr_weak = {10: _stream("안전")}
        r = HL.resolve_row(text_layer.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}, get_ocr_stream=ocr_weak.get)
        self.assertEqual((r.page, r.far, r.ocr), (13, True, False))                        # 먼 쪽 강한 적중 > 정본 쪽 OCR 표현만
        text_layer.pop(13)
        r = HL.resolve_row(text_layer.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}, get_ocr_stream=ocr_weak.get)
        self.assertEqual((r.page, r.k, r.ocr), (10, 0, True))                              # 남는 건 OCR 표현만 — 그래도 정본 쪽
        self.assertIsNone(HL.resolve_row(text_layer.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}))   # OCR 제공자 없음

    def test_ocr_on_a_neighbour_page_with_strong_context_comes_before_far_text(self):
        text_layer = {10: _stream("없음"), 11: _stream("없음"), 13: _stream("가나다라안전마바사아")}
        ocr = {11: _stream("가나다라안전마바사아")}                                                # 이웃 쪽 그림 속 글자
        r = HL.resolve_row(text_layer.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}, get_ocr_stream=ocr.get)
        self.assertEqual((r.page, r.k, r.ocr, r.far), (11, 8, True, False))               # ±1쪽 OCR 강한 적중 > 먼 쪽 텍스트
        weak = {11: _stream("안전")}
        r = HL.resolve_row(text_layer.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}, get_ocr_stream=weak.get)
        self.assertEqual((r.page, r.far, r.ocr), (13, True, False))                        # 표현만이면 먼 쪽 텍스트가 먼저
        text_layer.pop(13)
        self.assertIsNone(HL.resolve_row(text_layer.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}, get_ocr_stream=weak.get))
        # 표현만으로는 이웃 쪽 OCR 에 놓지 않는다 — 앞부속 예시 그림의 다른 '위생' 을 잡던 진단 사례

    def test_ocr_off_the_canonical_page_never_places_without_context(self):
        text_layer = {10: _stream("없음"), 11: _stream("안전")}
        elsewhere = HL.CharStream.from_chars([(c, fitz.Rect(500 + i * 10, 0, 510 + i * 10, 10), 0) for i, c in enumerate("안전")])
        used = {(11, "안전"): {(0, 2)}}                                                       # 이웃 쪽 텍스트 자리는 소진
        r = HL.resolve_row(text_layer.get, 10, "xx 안전 yy", 3, 5, "안전", used, get_ocr_stream={11: elsewhere}.get)
        self.assertEqual((r.page, r.shared, r.ocr), (11, True, False))                     # 문맥 없는 OCR 자리보다 텍스트 층 겹쳐 표시
        text_layer[11] = _stream("없음")
        self.assertIsNone(HL.resolve_row(text_layer.get, 10, "xx 안전 yy", 3, 5, "안전", {}, get_ocr_stream={11: elsewhere}.get))
        canonical = HL.resolve_row(text_layer.get, 10, "xx 안전 yy", 3, 5, "안전", {}, get_ocr_stream={10: elsewhere}.get)
        self.assertEqual((canonical.page, canonical.k, canonical.ocr), (10, 0, True))      # 정본 쪽 OCR 만 문맥 없이 받는다

    def test_ocr_hit_on_glyphs_already_highlighted_from_the_text_layer_is_not_a_new_find(self):
        text_layer = {10: _stream("없음"), 11: _stream("가나다라안전마바사아")}
        ocr = {11: _stream("가나다라안전마바사아")}                                                # 같은 글자, 같은 box
        used = {(11, "안전"): {(4, 6)}}                                                       # 텍스트 층 자리는 다른 행이 이미 썼다
        r = HL.resolve_row(text_layer.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", used, get_ocr_stream=ocr.get)
        self.assertEqual((r.page, r.shared, r.ocr), (11, True, False))                     # OCR 이 같은 글자를 다시 읽은 것 → 겹쳐 표시
        self.assertIn((4, 6), used[("ocr", 11, "안전")])                                   # 그 OCR 자리도 쓴 것으로 기록

    def test_far_ocr_needs_strong_context_and_carries_both_flags(self):
        text_layer = {10: _stream("없음"), 13: _stream("없음")}
        ocr = {13: _stream("가나다라안전마바사아")}
        r = HL.resolve_row(text_layer.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}, get_ocr_stream=ocr.get)
        self.assertEqual((r.page, r.k, r.ocr, r.far), (13, 8, True, True))
        self.assertIsNone(HL.resolve_row(text_layer.get, 10, "xx 안전 yy", 3, 5, "안전", {}, get_ocr_stream=ocr.get))   # 문맥 2자 — 먼 쪽 OCR 도 안 간다
        beyond = {10 + HL.FAR_OCR_MAX_PAGES + 1: _stream("가나다라안전마바사아")}
        self.assertIsNone(HL.resolve_row(text_layer.get, 10, "가나다라 안전 마바사아", 5, 7, "안전", {}, get_ocr_stream=beyond.get))

    def test_ocr_hit_is_new_compares_boxes(self):
        text = _stream("가나다라안전마바사아")
        self.assertFalse(HL._ocr_hit_is_new(_stream("가나다라안전마바사아"), (4, 6, 8), text, {(4, 6)}))
        self.assertTrue(HL._ocr_hit_is_new(_stream("가나다라안전마바사아"), (4, 6, 8), text, {(0, 2)}))     # 다른 글자를 썼다
        self.assertTrue(HL._ocr_hit_is_new(_stream("가나다라안전마바사아"), (4, 6, 8), None, {(4, 6)}))
        self.assertTrue(HL._ocr_hit_is_new(_stream("가나다라안전마바사아"), (4, 6, 8), text, None))

    def test_neighbour_sharing_is_the_last_resort(self):
        streams = {10: _stream("없음"), 11: _stream("안전")}
        used = {(11, "안전"): {(0, 2)}}
        r = HL.resolve_row(streams.get, 10, "xx 안전 yy", 3, 5, "안전", used)
        self.assertEqual((r.page, r.shared, r.start), (11, True, 0))
        self.assertIsNone(HL.resolve_row(streams.get, 10, "xx 화재 yy", 3, 5, "화재", used))

    def test_shared_resolution_prefers_the_best_context_match(self):
        streams = {10: _stream("안전수칙안전장비")}
        used = {(10, "안전"): {(0, 2), (4, 6)}}                                                   # 둘 다 이미 씀
        r = HL.resolve_row(streams.get, 10, "xx 안전 장비", 3, 5, "안전", used)
        self.assertEqual((r.start, r.end, r.shared, r.k), (4, 6, True, 2))


# ----------------------------------------------------------------------------- 6. 형광펜 주석
class AnnotTests(unittest.TestCase):
    def test_quads_for_merges_boxes_per_line(self):
        chars = [("안", fitz.Rect(0, 0, 10, 10), 0), ("전", fitz.Rect(10, 0, 20, 10), 0), ("수", fitz.Rect(0, 20, 10, 30), 1), ("칙", fitz.Rect(10, 20, 20, 30), 1)]
        s = HL.CharStream.from_chars(chars)
        quads = HL.quads_for(s, 1, 4)
        self.assertEqual([q.rect for q in quads], [fitz.Rect(10, 0, 20, 10), fitz.Rect(0, 20, 20, 30)])

    def test_add_highlight_sets_grade_color_and_popup(self):
        doc = fitz.open()
        page = doc.new_page(width=300, height=200)
        page.insert_text((20, 50), "안전 수칙", fontname="korea", fontsize=14)
        s = HL.CharStream.from_rawdict(page)
        row = HL.Row("NCS", "x.md", "안전", "안전", "exact", "안전", 4, 10, "안전 수칙", 3, "구체적 대책", 0)
        annot = HL.add_highlight(page, HL.quads_for(s, 0, 2), row, shown_page=9)
        self.assertEqual(annot.type[1], "Highlight")
        self.assertEqual([round(c, 3) for c in annot.colors["stroke"]], [round(c, 3) for c in HL.GRADE_COLORS[3]])
        self.assertEqual(annot.info["title"], "안전 · 등급3 구체적 대책")
        self.assertIn("정본 쪽 10", annot.info["content"])
        self.assertIn("표시 쪽 9", annot.info["content"])
        shared = HL.add_highlight(page, HL.quads_for(s, 0, 2), row, shown_page=10, shared=True)
        self.assertIn("같은 자리 중복", shared.info["content"])
        far = HL.add_highlight(page, HL.quads_for(s, 0, 2), row, shown_page=13, far=True)
        self.assertIn("표시 쪽 13", far.info["content"])
        self.assertIn("3쪽 떨어진", far.info["content"])
        ocr = HL.add_highlight(page, HL.quads_for(s, 0, 2), row, shown_page=10, ocr=True)
        self.assertIn("OCR", ocr.info["content"])
        far_ocr = HL.add_highlight(page, HL.quads_for(s, 0, 2), row, shown_page=13, far=True, ocr=True)
        self.assertIn("3쪽 떨어진", far_ocr.info["content"])
        self.assertIn("OCR", far_ocr.info["content"])
        self.assertTrue(annot.rect.intersects(page.search_for("안전")[0]))

    def test_add_highlight_content_omits_shift_note_when_pages_agree(self):
        doc = fitz.open()
        page = doc.new_page(width=300, height=200)
        page.insert_text((20, 50), "안전 수칙", fontname="korea", fontsize=14)
        s = HL.CharStream.from_rawdict(page)
        row = HL.Row("NCS", "x.md", "안전", "안전", "exact", "안전", 4, 10, "안전 수칙", 1, "미흡·없음", 0)
        annot = HL.add_highlight(page, HL.quads_for(s, 0, 2), row, shown_page=10)
        self.assertNotIn("표시 쪽", annot.info["content"])

    def test_add_legend_puts_a_korean_freetext_on_the_page(self):
        doc = fitz.open()
        page = doc.new_page()
        annot = HL.add_legend(page, "semantic_keyword_recount_20260917.xlsx")
        self.assertEqual(annot.type[1], "FreeText")
        self.assertIn("노랑", annot.info["content"])
        self.assertIn("20260917", annot.info["content"])
        self.assertEqual(len(list(page.annots())), 1)


# ----------------------------------------------------------------------------- 1. Vision OCR 도구
def _text_pdf(path: Path | None = None, scanned: bool = False) -> fitz.Document:
    """1쪽 표지, 2쪽 본문 3줄(안전·화재·감전 아님), 3쪽 부록. scanned=True 면 본문을 이미지로만 넣는다(텍스트 층 없음)."""
    src = fitz.open()
    src.new_page(width=400, height=300).insert_text((50, 60), "표지", fontname="korea", fontsize=20)
    body = src.new_page(width=400, height=300)
    body.insert_text((50, 80), "안전 수칙을 지킨다", fontname="korea", fontsize=18)
    body.insert_text((50, 130), "화재 예방 교육", fontname="korea", fontsize=18)
    src.new_page(width=400, height=300).insert_text((50, 60), "부록", fontname="korea", fontsize=20)
    src.new_page(width=400, height=300)
    src.new_page(width=400, height=300).insert_text((50, 60), "감전 사고를 막기 위한 절연 조치", fontname="korea", fontsize=18)   # 정본은 2쪽이라 하나 실제로는 5쪽
    if not scanned:
        doc = src
    else:
        doc = fitz.open()
        for i in range(src.page_count):
            pix = src[i].get_pixmap(dpi=200)
            page = doc.new_page(width=400, height=300)
            page.insert_image(page.rect, pixmap=pix)
    if path is not None:
        doc.save(str(path))
    return doc


def _figure_page_pdf(path: Path, sentence: str) -> None:
    """_text_pdf 의 3쪽(부록)을 sentence 가 그림으로만 찍힌 쪽으로 바꾼다 — 텍스트 층 없음, 정본은 2쪽이라 한다."""
    src = fitz.open()
    src.new_page(width=400, height=300).insert_text((50, 80), sentence, fontname="korea", fontsize=20)
    pix = src[0].get_pixmap(dpi=200)
    doc = _text_pdf()
    doc.delete_page(2)
    page = doc.new_page(pno=2, width=400, height=300)
    page.insert_image(page.rect, pixmap=pix)
    doc.save(str(path))


HAS_SWIFTC = shutil.which("swiftc") is not None


@unittest.skipUnless(HAS_SWIFTC, "swiftc 가 없다 — Vision OCR 통합은 macOS 에서만")
class OcrToolTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        cls.binary = HL.build_ocr_tool(cls.tmp / "bin")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_build_reuses_the_binary_when_the_source_is_older(self):
        before = self.binary.stat().st_mtime_ns
        self.assertEqual(HL.build_ocr_tool(self.tmp / "bin"), self.binary)
        self.assertEqual(self.binary.stat().st_mtime_ns, before)

    def test_tool_recognizes_korean_with_char_boxes(self):
        doc = _text_pdf()
        png = self.tmp / "p2.png"
        doc[1].get_pixmap(dpi=200).save(str(png))
        (result,) = HL.run_ocr_tool(self.binary, [png])
        self.assertEqual((result["width"], result["height"]), (doc[1].get_pixmap(dpi=200).width, doc[1].get_pixmap(dpi=200).height))
        text = HL.norm_chars("".join(line["text"] for line in result["lines"]))
        self.assertIn("안전", text)
        self.assertIn("화재", text)
        line = next(l for l in result["lines"] if "안전" in l["text"])
        self.assertEqual(len(line["chars"]), len(line["text"]))
        self.assertTrue(all(c["box"] for c in line["chars"] if c["c"].strip()))

    def test_ocr_pages_caches_json_per_page(self):
        doc = _text_pdf()
        cache = self.tmp / "cache"
        pages = HL.ocr_pages(doc, [1], dpi=200, cache_dir=cache, binary=self.binary, stem="시험")
        self.assertEqual(list(pages), [1])
        self.assertTrue((cache / "ocr" / "시험" / "1.json").is_file())
        again = HL.ocr_pages(doc, [1], dpi=200, cache_dir=cache, binary=Path("/nonexistent"), stem="시험")   # 캐시만으로 답한다
        self.assertEqual(again[1]["lines"][0]["text"], pages[1]["lines"][0]["text"])


# ----------------------------------------------------------------------------- 7. 파이프라인
DETAIL_HEADERS = ("파일", "키워드", "등록 표현", "실제 매칭", "계층", "줄", "페이지", "판정 근거", "문맥", "통일 등급", "등급명", "등급사유", "등급 출처")


def _write_xlsx(path: Path, ncs_rows: list[tuple], tb_rows: list[tuple]) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "README"
    for name, rows in (("NCS_매칭상세", ncs_rows), ("교과서_매칭상세", tb_rows)):
        ws = wb.create_sheet(name)
        ws.append(DETAIL_HEADERS)
        for r in rows:
            ws.append(r)
    wb.save(str(path))


def _row(relpath, keyword, matched, line, page, grade, context, label="형식적 언급"):
    return (relpath, keyword, keyword, matched, "exact", line, page, "기존 키워드의 정확 문자열", context, grade, label, "사유", "실제 쪽 판정")


MD_BODY = "<!-- page: 1 -->\n표지\n<!-- page: 2 -->\n안전 수칙을 지킨다\n화재 예방 교육\n감전 주의\n감전 사고를 막기 위한 절연 조치\n<!-- page: 3 -->\n부록\n"


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        for d in ("ncs_pdf/반도체개발", "tb_pdf", "ncs_md", "tb_md", "out"):
            (self.tmp / d).mkdir(parents=True)

    def ncs_fixture(self):
        rel = "반도체개발/LM9999999999_시험_교재/LM9999999999_시험_교재.md"
        md = self.tmp / "ncs_md" / rel
        md.parent.mkdir(parents=True)
        md.write_text(MD_BODY, encoding="utf-8")
        pdf = self.tmp / "ncs_pdf/반도체개발/LM9999999999_시험 교재.pdf"
        _text_pdf(pdf)
        rows = [_row(rel, "안전", "안전", 4, 2, 2, "안전 수칙을 지킨다"),
                _row(rel, "화재", "화재", 5, 2, 2, "화재 예방 교육"),
                _row(rel, "감전", "감전", 6, 2, 2, "감전 주의"),                       # PDF 에는 없는 줄 → 미발견
                _row(rel, "감전", "감전", 7, 2, 2, "감전 사고를 막기 위한 절연 조치")]      # 5쪽에 있음 → 먼 쪽(거리 3)
        return rel, pdf, rows

    def argv(self, xlsx):
        return ["--xlsx", str(xlsx), "--ncs-pdf-root", str(self.tmp / "ncs_pdf"), "--textbook-pdf-root", str(self.tmp / "tb_pdf"),
                "--ncs-md-root", str(self.tmp / "ncs_md"), "--school-md-root", str(self.tmp / "tb_md"), "--out", str(self.tmp / "out"),
                "--skip-corpus-check", "--quiet"]

    def test_load_rows_reads_both_detail_sheets_in_order(self):
        rel, _, rows = self.ncs_fixture()
        xlsx = self.tmp / "t.xlsx"
        _write_xlsx(xlsx, rows, [_row("20260101_000000_시험교과서_출판사_/a.md", "안전", "안전", 2, 1, 3, "안전", "구체적 대책")])
        loaded = HL.load_rows(xlsx)
        self.assertEqual([(r.corpus, r.keyword, r.line, r.page, r.grade, r.order) for r in loaded],
                         [("NCS", "안전", 4, 2, 2, 0), ("NCS", "화재", 5, 2, 2, 1), ("NCS", "감전", 6, 2, 2, 2), ("NCS", "감전", 7, 2, 2, 3), ("교과서", "안전", 2, 1, 3, 4)])
        self.assertEqual(loaded[4].grade_label, "구체적 대책")

    def test_end_to_end_ncs_highlights_resolved_rows_and_logs_the_rest(self):
        import hashlib, json
        rel, pdf, rows = self.ncs_fixture()
        xlsx = self.tmp / "t.xlsx"
        _write_xlsx(xlsx, rows, [])
        before = hashlib.sha256(pdf.read_bytes()).hexdigest()
        self.assertEqual(HL.main(self.argv(xlsx)), 0)
        self.assertEqual(hashlib.sha256(pdf.read_bytes()).hexdigest(), before)                    # 원본은 그대로
        out = self.tmp / "out/ncs/반도체개발/LM9999999999_시험 교재_키워드표시.pdf"
        doc = fitz.open(str(out))
        page1, page2, page5 = doc[0], doc[1], doc[4]
        self.assertEqual([a.type[1] for a in page1.annots()], ["FreeText"])
        far = list(page5.annots())
        self.assertEqual([a.type[1] for a in far], ["Highlight"])
        self.assertIn("3쪽 떨어진", far[0].info["content"])
        highlights = list(page2.annots())
        self.assertEqual([a.type[1] for a in highlights], ["Highlight", "Highlight"])
        self.assertEqual({a.info["title"] for a in highlights}, {"안전 · 등급2 형식적 언급", "화재 · 등급2 형식적 언급"})
        self.assertTrue(all([round(c, 3) for c in a.colors["stroke"]] == [round(c, 3) for c in HL.GRADE_COLORS[2]] for a in highlights))
        self.assertTrue(highlights[0].rect.intersects(page2.search_for("안전")[0]))
        log = json.loads((self.tmp / "out/highlight_log.json").read_text(encoding="utf-8"))
        self.assertEqual(log["totals"], {"books": 1, "rows": 4, "same_page": 2, "shared": 0, "adjacent": 0, "far": 1, "ocr": 0, "ocr_adjacent": 0,
                                         "ocr_far": 0, "unresolved": 1})
        self.assertEqual(log["books"][0]["locator"], {"engine": 4, "naive": 0})
        self.assertEqual(log["xlsx_sha256"], hashlib.sha256(xlsx.read_bytes()).hexdigest())
        csv_text = (self.tmp / "out/unresolved.csv").read_text(encoding="utf-8")
        self.assertEqual(csv_text.count("\n"), 2)                                                # 헤더 + 1행
        self.assertIn("감전", csv_text)
        self.assertNotIn("감전 주의", csv_text)                                                    # 본문 문맥은 싣지 않는다
        self.assertIn("#FFF176", (self.tmp / "out/README.md").read_text(encoding="utf-8"))

    def test_log_records_annotations_the_original_already_had(self):
        import json
        rel, pdf, rows = self.ncs_fixture()
        doc = fitz.open(str(pdf))
        page = doc[1]
        page.add_highlight_annot(page.search_for("예방")[0]).update()                                 # 원본에 이미 있던 형광펜
        doc.save(str(pdf), incremental=True, encryption=fitz.PDF_ENCRYPT_KEEP)
        doc.close()
        xlsx = self.tmp / "t.xlsx"
        _write_xlsx(xlsx, rows, [])
        self.assertEqual(HL.main(self.argv(xlsx)), 0)
        log = json.loads((self.tmp / "out/highlight_log.json").read_text(encoding="utf-8"))
        self.assertEqual(log["books"][0]["preexisting_annots"], {"Highlight": 1})
        out = fitz.open(str(self.tmp / "out/ncs/반도체개발/LM9999999999_시험 교재_키워드표시.pdf"))
        page2 = out[1]
        self.assertEqual(sum(1 for a in page2.annots() if a.type[1] == "Highlight"), 3)             # 원본 1 + 새 2, 원본 주석은 그대로

    def test_main_refuses_when_a_book_has_no_pdf(self):
        rel, pdf, rows = self.ncs_fixture()
        pdf.unlink()
        xlsx = self.tmp / "t.xlsx"
        _write_xlsx(xlsx, rows, [])
        with self.assertRaises(LookupError):
            HL.main(self.argv(xlsx))

    def test_only_filter_writes_a_partial_log(self):
        rel, pdf, rows = self.ncs_fixture()
        xlsx = self.tmp / "t.xlsx"
        _write_xlsx(xlsx, rows, [])
        self.assertEqual(HL.main(self.argv(xlsx) + ["--only", "LM0000000000"]), 1)          # 아무 권도 안 맞음 → 부분 실행, 합계 0
        self.assertFalse((self.tmp / "out/highlight_log.json").exists())
        self.assertTrue((self.tmp / "out/highlight_log.partial.json").exists())

    @unittest.skipUnless(HAS_SWIFTC, "swiftc 가 없다")
    def test_end_to_end_ncs_finds_figure_text_on_a_neighbour_page_by_ocr(self):
        import json
        rel, pdf, rows = self.ncs_fixture()
        sentence = "감전 주의 표지를 부착한 뒤 작업한다"
        (self.tmp / "ncs_md" / rel).write_text(MD_BODY.replace("감전 주의", sentence), encoding="utf-8")
        _figure_page_pdf(pdf, sentence)                                                          # 3쪽 = 그 문장이 그림으로만 있는 쪽
        self.assertEqual(fitz.open(str(pdf))[2].get_text(), "")
        xlsx = self.tmp / "t.xlsx"
        _write_xlsx(xlsx, rows, [])
        self.assertEqual(HL.main(self.argv(xlsx)), 0)
        log = json.loads((self.tmp / "out/highlight_log.json").read_text(encoding="utf-8"))
        self.assertEqual(log["totals"], {"books": 1, "rows": 4, "same_page": 2, "shared": 0, "adjacent": 0, "far": 1, "ocr": 1, "ocr_adjacent": 1,
                                         "ocr_far": 0, "unresolved": 0})
        doc = fitz.open(str(self.tmp / "out/ncs/반도체개발/LM9999999999_시험 교재_키워드표시.pdf"))
        page3 = doc[2]
        (annot,) = list(page3.annots())
        self.assertEqual(annot.info["title"], "감전 · 등급2 형식적 언급")
        self.assertIn("표시 쪽 3", annot.info["content"])
        self.assertIn("OCR", annot.info["content"])
        self.assertTrue(annot.rect.intersects(fitz.Rect(40, 55, 120, 90)))                      # "감전" 이 찍힌 자리 근처
        readme = (self.tmp / "out/README.md").read_text(encoding="utf-8")
        self.assertIn("그림", readme)

    @unittest.skipUnless(HAS_SWIFTC, "swiftc 가 없다")
    def test_end_to_end_textbook_uses_ocr_for_a_scanned_pdf(self):
        rel = "20260101_000000_시험교과서_출판사_/20260101_000000_시험교과서_출판사_.md"
        md = self.tmp / "tb_md" / rel
        md.parent.mkdir(parents=True)
        md.write_text(MD_BODY, encoding="utf-8")
        pdf = self.tmp / "tb_pdf/시험교과서(출판사).pdf"
        _text_pdf(pdf, scanned=True)
        self.assertEqual(fitz.open(str(pdf))[1].get_text(), "")                                   # 텍스트 층 없음
        xlsx = self.tmp / "t.xlsx"
        _write_xlsx(xlsx, [], [_row(rel, "안전", "안전", 4, 2, 3, "안전 수칙을 지킨다", "구체적 대책")])
        self.assertEqual(HL.main(self.argv(xlsx)), 0)
        doc = fitz.open(str(self.tmp / "out/school-text/시험교과서(출판사)_키워드표시.pdf"))
        page = doc[1]                                                                              # Page 를 살려 둬야 주석이 붙어 있다
        highlights = list(page.annots())
        self.assertEqual([a.type[1] for a in highlights], ["Highlight"])
        self.assertEqual([round(c, 3) for c in highlights[0].colors["stroke"]], [round(c, 3) for c in HL.GRADE_COLORS[3]])
        self.assertTrue(highlights[0].rect.intersects(fitz.Rect(50, 60, 90, 85)))                 # "안전" 이 찍힌 자리 근처
        self.assertTrue((self.tmp / "out/.cache/ocr").is_dir())


if __name__ == "__main__":
    unittest.main()
