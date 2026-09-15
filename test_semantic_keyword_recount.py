import tempfile
import unicodedata
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import json
import re
from dataclasses import replace
from unittest import mock

import semantic_keyword_recount as SKR
from semantic_keyword_recount import (
    CandidateDecision,
    DedupRecord,
    Document,
    EXPECTED,
    EXPECTED_KEYWORDS,
    ExpressionRule,
    GradeAssignment,
    InputArtifact,
    KeywordSource,
    aggregate_matches,
    assign_match_grades,
    build_default_rules,
    check_expected,
    check_marker_base,
    dashboard_payload,
    grade_alias_key,
    grade_lookup_key,
    default_candidate_decisions,
    load_documents,
    load_existing_grades,
    public_path,
    read_keyword_workbook,
    run_census,
    run_manifest,
    scan_document,
    select_ncs_documents,
    split_pages,
    summary_metrics,
    summary_payload,
    artifact_manifest,
    write_analysis_pages,
    validate_rules,
    write_report,
    write_dashboard_data,
    write_workbook,
)


class InputReaderTests(unittest.TestCase):
    def test_workbook_reader_keeps_headerless_first_data_row(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "keywords.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "안전"
            ws.append(["number", "영역", "filename", "contents"])
            ws.append([1, "개발", "book-a", "안전 문장"])
            ws.append([2, "제조", "book-b", "안전 문장"])
            no_header = wb.create_sheet("사망")
            no_header.append([3, "장비", "book-c", "사망 문장"])
            no_header.append([4, "재료", "book-d", "사망 문장"])
            wb.save(path)

            sources = {item.keyword: item for item in read_keyword_workbook(path)}

        self.assertEqual(2, len(sources))
        self.assertEqual(2, sources["안전"].search_rows)
        self.assertEqual(2, sources["사망"].search_rows)
        self.assertTrue(sources["안전"].has_header)
        self.assertFalse(sources["사망"].has_header)

    def test_document_paths_are_nfc_and_markers_propagate(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            nfd_name = unicodedata.normalize("NFD", "안전문서.md")
            (root / nfd_name).write_text(
                "머리말\n<!-- page: 3 -->\n첫 문장\n둘째 문장\n"
                "<!-- page: 4 -->\n마지막 문장\n",
                encoding="utf-8",
            )

            docs = load_documents(root, "NCS")

        self.assertEqual(1, len(docs))
        blocks = split_pages(docs[0])
        self.assertEqual("안전문서.md", docs[0].relative_path)
        self.assertEqual([None, 3, 4], [block.page for block in blocks])
        self.assertEqual(("첫 문장", "둘째 문장"), blocks[1].lines)
        self.assertEqual(3, blocks[1].start_line)


class RuleEngineTests(unittest.TestCase):
    def document(self, text):
        return Document("NCS", Path("source.md"), "source.md", text)

    def test_longest_expression_wins_within_one_keyword(self):
        rules = [
            ExpressionRule("보호구", "보호구", "exact", "기존 키워드"),
            ExpressionRule("보호구", "안전 장갑", "specific", "손 보호구"),
            ExpressionRule("보호구", "장갑", "specific", "손 보호구"),
        ]

        records = scan_document(self.document("안전 장갑을 착용한다."), rules)
        included = [record for record in records if record.decision == "included"]

        self.assertEqual(1, len(included))
        self.assertEqual("안전 장갑", included[0].expression)
        self.assertEqual("안전 장갑", included[0].matched_text)

    def test_english_matching_is_case_insensitive(self):
        rules = [ExpressionRule("MSDS", "MSDS", "exact", "기존 키워드")]

        records = scan_document(self.document("msds와 MsDs를 확인한다."), rules)

        self.assertEqual(["msds", "MsDs"], [record.matched_text for record in records])

    def test_explicit_pattern_accepts_spacing_and_punctuation_variant(self):
        rules = [
            ExpressionRule(
                "작업환경",
                "작업 환경",
                "exact",
                "띄어쓰기·기호 변형",
                pattern=r"작업[\s·ㆍ-]*환경",
            )
        ]

        records = scan_document(self.document("작업·환경을 점검한다."), rules)

        self.assertEqual(1, len(records))
        self.assertEqual("작업·환경", records[0].matched_text)

    def test_exact_homonym_is_excluded_without_hiding_valid_occurrence(self):
        rules = [
            ExpressionRule(
                "부상",
                "부상",
                "exact",
                "기존 키워드",
                exclude_patterns=(r"부상(?:하|했|하여|하고|하는|한|할)",),
            )
        ]

        records = scan_document(
            self.document("신산업으로 부상하고 있지만 근로자 부상 사고가 있었다."),
            rules,
        )

        self.assertEqual(2, len(records))
        self.assertEqual(["excluded", "included"], [record.decision for record in records])
        self.assertEqual(["부상", "부상"], [record.matched_text for record in records])

    def test_existing_keyword_cannot_expand_another_keyword(self):
        rules = [
            ExpressionRule("MSDS", "MSDS", "exact", "기존 키워드"),
            ExpressionRule("물질안전보건자료", "물질안전보건자료", "exact", "기존 키워드"),
            ExpressionRule("MSDS", "물질안전보건자료", "equivalent", "금지된 상호 확장"),
        ]

        with self.assertRaisesRegex(ValueError, "기존 키워드 상호 확장"):
            validate_rules(["MSDS", "물질안전보건자료"], rules)


class RegistryAndAggregationTests(unittest.TestCase):
    def test_default_registry_has_an_exact_rule_for_every_keyword(self):
        keywords = ["안전", "사망", "부상"]

        rules = build_default_rules(keywords)

        exact_owners = {rule.keyword for rule in rules if rule.tier == "exact"}
        self.assertEqual(set(keywords), exact_owners)

    def test_candidate_decisions_use_a_closed_status_set(self):
        candidates = [
            CandidateDecision("부상", "상해", "included", "equivalent", "신체 손상"),
            CandidateDecision("추락", "fall", "held", "equivalent", "문맥 혼재"),
            CandidateDecision("중독", "poisoning", "excluded", "equivalent", "공정 용어"),
            CandidateDecision("직업병", "occupational disease", "not-found", "equivalent", "미출현"),
        ]

        self.assertEqual(
            {"included", "held", "excluded", "not-found"},
            {candidate.decision for candidate in candidates},
        )

    def test_full_registry_is_consistent_with_candidate_decisions(self):
        rules = build_default_rules(list(EXPECTED_KEYWORDS))
        candidates = default_candidate_decisions()
        included_candidates = {
            (candidate.keyword, candidate.expression)
            for candidate in candidates
            if candidate.decision == "included"
        }
        added_rules = {
            (rule.keyword, rule.expression) for rule in rules if rule.tier != "exact"
        }

        validate_rules(list(EXPECTED_KEYWORDS), rules)
        self.assertEqual(30, len(EXPECTED_KEYWORDS))
        self.assertEqual(included_candidates, added_rules)
        self.assertGreater(len(added_rules), 20)
        for keyword, expression in added_rules:
            self.assertNotIn(expression.casefold(), {
                other.casefold() for other in EXPECTED_KEYWORDS if other != keyword
            })

    def test_aggregation_reconciles_summary_with_detail_records(self):
        sources = [KeywordSource("안전", 7, True)]
        documents = [
            Document(
                "NCS",
                Path("source.md"),
                "source.md",
                "<!-- page: 2 -->\n안전과 safety를 확인한다.\n",
            )
        ]
        rules = [
            ExpressionRule("안전", "안전", "exact", "기존 키워드"),
            ExpressionRule(
                "안전",
                "safety",
                "equivalent",
                "영문 동등 표현",
                pattern=r"(?<![A-Za-z])safety(?![A-Za-z])",
            ),
        ]

        result = aggregate_matches(sources, documents, rules, [])
        row = next(item for item in result.summary if item.corpus == "NCS")
        included = [record for record in result.matches if record.decision == "included"]

        self.assertEqual(1, row.raw_exact)
        self.assertEqual(0, row.excluded_exact)
        self.assertEqual(1, row.valid_exact)
        self.assertEqual(1, row.equivalent_added)
        self.assertEqual(0, row.specific_added)
        self.assertEqual(2, row.semantic_total)
        self.assertEqual(row.semantic_total, len(included))
        self.assertEqual(1, row.file_count)
        self.assertEqual(1, row.page_count)


class GradeAssignmentTests(unittest.TestCase):
    def test_existing_page_grade_is_inherited_and_new_page_is_regraded(self):
        sources = [KeywordSource("안전", 1, True)]
        documents = [
            Document(
                "NCS",
                Path("source.md"),
                "source.md",
                "머리말 안전\n"
                "<!-- page: 1 -->\n안전 안내\n"
                "<!-- page: 2 -->\n" + "안전 " * 6 + "방지 예방 착용 환기 차단\n",
            )
        ]
        rules = [ExpressionRule("안전", "안전", "exact", "기존 키워드")]
        result = aggregate_matches(sources, documents, rules, [])
        existing = {
            grade_lookup_key("NCS", "source.md", 1): GradeAssignment(
                grade=2,
                label="형식적 언급",
                reason="기존 판정 사유",
                source="existing",
            )
        }

        graded = assign_match_grades(result, existing)
        by_page = {}
        for record in graded.matches:
            by_page.setdefault(record.page, record)

        self.assertEqual(1, by_page[None].grade)
        self.assertEqual("unpaged-context", by_page[None].grade_source)
        self.assertEqual(2, by_page[1].grade)
        self.assertEqual("기존 판정 사유", by_page[1].grade_reason)
        self.assertEqual("existing", by_page[1].grade_source)
        self.assertEqual(3, by_page[2].grade)
        self.assertEqual("new", by_page[2].grade_source)
        summary = next(row for row in graded.summary if row.corpus == "NCS")
        self.assertEqual(1, summary.grade_1)
        self.assertEqual(1, summary.grade_2)
        self.assertEqual(6, summary.grade_3)
        self.assertEqual(0, summary.grade_unpaged)
        self.assertEqual(summary.semantic_total, summary.grade_1 + summary.grade_2 + summary.grade_3 + summary.grade_unpaged)

    def test_unpaged_match_uses_its_occurrence_context(self):
        sources = [KeywordSource("안전", 1, True)]
        documents = [Document("NCS", Path("source.md"), "source.md", "안전\n")]
        result = aggregate_matches(
            sources,
            documents,
            [ExpressionRule("안전", "안전", "exact", "기존 키워드")],
            [],
        )
        graded = assign_match_grades(result, {})
        record = next(r for r in graded.matches if r.decision == "included")
        self.assertEqual((1, "unpaged-context"), (record.grade, record.grade_source))

    def test_paged_record_without_page_text_falls_back_to_grade_one(self):
        result = aggregate_matches(
            [KeywordSource("안전", 1, True)],
            [Document("NCS", Path("a.md"), "반도체개발/LM1903060101_a/a.md", "<!-- page: 1 -->\n안전\n")],
            [ExpressionRule("안전", "안전", "exact", "기존 키워드")],
            [],
        )
        record = next(r for r in result.matches if r.decision == "included")
        orphan = replace(result, matches=(replace(record, page=7),))   # 7쪽 블록은 없다
        graded = assign_match_grades(orphan, {})
        got = graded.matches[0]
        self.assertEqual((1, "unpaged-fallback"), (got.grade, got.grade_source))
        self.assertIn("본문", got.grade_reason)

    def test_legacy_workbooks_are_normalized_to_one_grade_scale(self):
        with tempfile.TemporaryDirectory() as td:
            ncs_path = Path(td) / "ncs.xlsx"
            school_path = Path(td) / "school.xlsx"
            for path, filename, raw_grade in (
                (ncs_path, "LM1903060101_안전", 3),
                (school_path, "20260413_171220_반도체기초기술1_크리아트_.md", 1),
            ):
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.title = "안전"
                worksheet.append(["number", "영역", "filename", "contents", "page", "페이지전체내용", "사고사례여부", "등급", "등급사유"])
                worksheet.append([1, "반도체개발", filename, "안전", 7, "안전 " * 7, "아니오", raw_grade, "원본 사유"])
                workbook.save(path)

            grades = load_existing_grades(ncs_path, school_path)

        ncs = grades[grade_lookup_key("NCS", "반도체개발/LM1903060101_안전.md", 7)]
        school = grades[grade_lookup_key("교과서", "20260413_171220_반도체기초기술1_크리아트_.md", 7)]
        self.assertEqual((3, "구체적 대책", "existing"), (ncs.grade, ncs.label, ncs.source))
        self.assertEqual((2, "형식적 언급", "existing"), (school.grade, school.label, school.source))

    def test_new_page_grading_keeps_duplicate_ncs_codes_separate(self):
        sources = [KeywordSource("안전", 1, True)]
        documents = [
            Document(
                "NCS",
                Path("a.md"),
                "반도체제조/LM1903060205_a.md",
                "<!-- page: 1 -->\n안전\n",
            ),
            Document(
                "NCS",
                Path("b.md"),
                "반도체제조/LM1903060205_b.md",
                "<!-- page: 1 -->\n" + ("안전 " * 6) + ("보호구 착용 " * 5) + "\n",
            ),
        ]
        result = aggregate_matches(
            sources,
            documents,
            [ExpressionRule("안전", "안전", "exact", "기존 키워드")],
            [],
        )

        graded = assign_match_grades(result, {})
        grades_by_path = {
            record.relative_path: record.grade
            for record in graded.matches
            if record.decision == "included"
        }

        self.assertEqual(1, grades_by_path["반도체제조/LM1903060205_a.md"])
        self.assertEqual(3, grades_by_path["반도체제조/LM1903060205_b.md"])

    def test_ambiguous_ncs_code_does_not_inherit_alias_grade(self):
        sources = [KeywordSource("안전", 1, True)]
        documents = [
            Document("NCS", Path("a.md"), "x/LM1903060205_a.md", "<!-- page: 1 -->\n안전\n"),
            Document(
                "NCS",
                Path("b.md"),
                "x/LM1903060205_b.md",
                "<!-- page: 1 -->\n" + ("안전 " * 6) + ("보호구 착용 " * 5) + "\n",
            ),
        ]
        result = aggregate_matches(
            sources,
            documents,
            [ExpressionRule("안전", "안전", "exact", "기존 키워드")],
            [],
        )
        legacy = GradeAssignment(2, "형식적 언급", "기존", "existing")

        graded = assign_match_grades(
            result,
            {grade_alias_key("NCS", "LM1903060205_legacy", 1): legacy},
        )
        by_path = {
            record.relative_path: (record.grade, record.grade_source)
            for record in graded.matches
            if record.decision == "included"
        }

        self.assertEqual((1, "new"), by_path["x/LM1903060205_a.md"])
        self.assertEqual((3, "new"), by_path["x/LM1903060205_b.md"])

    def test_ambiguous_normalized_stem_does_not_inherit_grade(self):
        sources = [KeywordSource("안전", 1, True)]
        documents = [
            Document(
                "NCS",
                Path("space.md"),
                "x/LM1903060205_14v3_MI 장비 운영.md",
                "<!-- page: 1 -->\n안전\n",
            ),
            Document(
                "NCS",
                Path("underscore.md"),
                "x/LM1903060205_14v3_MI_장비_운영.md",
                "<!-- page: 1 -->\n" + ("안전 " * 6) + ("보호구 착용 " * 5) + "\n",
            ),
        ]
        result = aggregate_matches(
            sources,
            documents,
            [ExpressionRule("안전", "안전", "exact", "기존 키워드")],
            [],
        )
        legacy = GradeAssignment(2, "형식적 언급", "기존", "existing")

        graded = assign_match_grades(
            result,
            {grade_lookup_key("NCS", "LM1903060205_14v3_MI 장비 운영", 1): legacy},
        )
        by_path = {
            record.relative_path: (record.grade, record.grade_source)
            for record in graded.matches
            if record.decision == "included"
        }

        self.assertEqual((1, "new"), by_path["x/LM1903060205_14v3_MI 장비 운영.md"])
        self.assertEqual((3, "new"), by_path["x/LM1903060205_14v3_MI_장비_운영.md"])


class OutputTests(unittest.TestCase):
    def sample_result(self):
        sources = [KeywordSource("안전", 3, True)]
        documents = [
            Document(
                "NCS",
                Path("source.md"),
                "source.md",
                "<!-- page: 1 -->\n안전과 safety를 확인한다.\n",
            )
        ]
        rules = [
            ExpressionRule("안전", "안전", "exact", "기존 키워드"),
            ExpressionRule("안전", "safety", "equivalent", "영문 동등 표현"),
        ]
        candidates = [
            CandidateDecision("안전", "safety", "included", "equivalent", "영문 동등 표현"),
            CandidateDecision("안전", "안전성", "held", "equivalent", "문맥 혼재"),
        ]
        return aggregate_matches(sources, documents, rules, candidates)

    def test_workbook_has_required_sheets_and_readable_headers(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "result.xlsx"
            write_workbook(self.sample_result(), path)

            workbook = load_workbook(path, read_only=False, data_only=True)
            try:
                self.assertEqual(
                    [
                        "README",
                        "입력정보",
                        "요약",
                        "포함표현",
                        "보류제외",
                        "NCS_파일별",
                        "교과서_파일별",
                        "NCS_매칭상세",
                        "교과서_매칭상세",
                        "문맥근거",
                    ],
                    workbook.sheetnames,
                )
                self.assertEqual("A2", workbook["요약"].freeze_panes)
                self.assertGreater(workbook["요약"].column_dimensions["A"].width, 5)
                self.assertEqual("말뭉치", workbook["요약"]["A1"].value)
                self.assertIn("등급1 출현", [cell.value for cell in workbook["요약"][1]])
                detail_headers = [cell.value for cell in workbook["NCS_매칭상세"][1]]
                self.assertEqual(
                    ["통일 등급", "등급명", "등급사유", "등급 출처"],
                    detail_headers[-4:],
                )
            finally:
                workbook.close()

    def test_report_and_manifest_are_auditable(self):
        result = self.sample_result()
        manifest = artifact_manifest(result)
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "report.md"
            write_report(result, path)
            report = path.read_text(encoding="utf-8")

        self.assertIn("키워드 독립 원칙", report)
        self.assertIn("키워드별 포함 표현", report)
        self.assertIn("등급1 출현 | 등급2 출현 | 등급3 출현 | 등급 미확정 출현", report)
        self.assertIn("재현성 해시", report)
        self.assertEqual(
            {"source_sha256", "rule_sha256", "detail_sha256", "summary_sha256"},
            set(manifest),
        )
        self.assertTrue(all(len(value) == 64 for value in manifest.values()))

    def test_report_has_ranked_keyword_statistics_per_corpus(self):
        """말뭉치별 키워드 순위 통계 — 출현 내림차순, 비율·누적·등급3 비율·확장분. 교과서에 출현이 없는 키워드도 0 으로 남는다."""
        sources = [KeywordSource("안전", 3, True), KeywordSource("위험", 1, True), KeywordSource("추락", 0, True)]
        documents = [
            Document("NCS", Path("a.md"), "반도체개발/LM1903060101_a/a.md", "<!-- page: 1 -->\n안전 위험 위험\n<!-- page: 2 -->\n안전 safety\n"),
            Document("NCS", Path("b.md"), "반도체개발/LM1903060102_b/b.md", "<!-- page: 1 -->\n위험 위험\n"),
            Document("교과서", Path("s.md"), "s.md", "<!-- page: 1 -->\n위험 위험\n"),
        ]
        rules = [ExpressionRule("안전", "안전", "exact", "기존 키워드"), ExpressionRule("안전", "safety", "equivalent", "영문"),
                 ExpressionRule("위험", "위험", "exact", "기존 키워드"), ExpressionRule("추락", "추락", "exact", "기존 키워드")]
        result = assign_match_grades(aggregate_matches(sources, documents, rules, []), {})
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "report.md"
            write_report(result, path)
            report = path.read_text(encoding="utf-8")
        head, _, _ = report.partition("## 말뭉치별 확장 차이")
        _, _, stats = head.partition("## 키워드 순위 통계")
        self.assertTrue(stats, "키워드 순위 통계 절이 없다")
        ncs, _, school = stats.partition("### 교과서")
        self.assertIn("### NCS", ncs)
        ncs_rows = [l for l in ncs.splitlines() if l.startswith("| ") and not l.startswith("| 순위")]
        self.assertEqual(["위험", "안전", "추락"], [r.split(" | ")[1].strip("`") for r in ncs_rows[:3]])      # 4 > 3 > 0
        self.assertIn("| 1 | `위험` | 4 | 57.1% | 57.1% |", ncs_rows[0])                                  # 4/7 누적 57.1
        self.assertIn("| 2 | `안전` | 3 | 42.9% | 100.0% |", ncs_rows[1])
        self.assertIn("| 3 | `추락` | 0 | 0.0% | 100.0% |", ncs_rows[2])
        self.assertIn("| 합계 | | 7 | 100.0% |", ncs)
        self.assertIn("2/2 (100.0%)", ncs_rows[0])                                                        # 위험: 파일 2/2
        self.assertIn("| 1 | 33.3% |", ncs_rows[1])                                                       # 안전: 확장분 safety 1건 = 1/3
        school_rows = [l for l in school.splitlines() if l.startswith("| ") and not l.startswith("| 순위")]
        self.assertEqual(["위험", "안전", "추락"], [r.split(" | ")[1].strip("`") for r in school_rows[:3]])  # 동률 0 은 원본 키워드 순서
        self.assertIn("| 합계 | | 2 | 100.0% |", school)

    def test_manifest_source_hash_includes_grade_workbook_lineage(self):
        base = self.sample_result()
        first = type(base)(
            base.sources,
            base.documents,
            base.rules,
            base.candidates,
            base.matches,
            base.summary,
            (InputArtifact("NCS 등급 워크북", "/tmp/ncs.xlsx", 1, "a" * 64),),
        )
        second = type(base)(
            base.sources,
            base.documents,
            base.rules,
            base.candidates,
            base.matches,
            base.summary,
            (InputArtifact("NCS 등급 워크북", "/other/ncs.xlsx", 1, "b" * 64),),
        )

        self.assertNotEqual(
            artifact_manifest(first)["source_sha256"],
            artifact_manifest(second)["source_sha256"],
        )

    def test_run_census_applies_legacy_grades_before_writing_outputs(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            ncs_root = root / "ncs"
            school_root = root / "school"
            ncs_root.mkdir()
            school_root.mkdir()
            (ncs_root / "LM1903060101_안전.md").write_text("<!-- page: 1 -->\n안전 안내\n", encoding="utf-8")
            for index in range(1, 86):
                (ncs_root / f"LM19030602{index:02d}_book.md").write_text("", encoding="utf-8")
            for index in range(9):
                (school_root / f"school-{index}.md").write_text("", encoding="utf-8")

            source = root / "source.xlsx"
            workbook = Workbook()
            workbook.remove(workbook.active)
            for keyword in EXPECTED_KEYWORDS:
                worksheet = workbook.create_sheet(keyword)
                worksheet.append(["number", "영역", "filename", "contents", "page", "페이지전체내용", "사고사례여부", "등급", "등급사유"])
                if keyword == "안전":
                    worksheet.append([1, "반도체개발", "LM1903060101_안전", "안전", 1, "안전 안내", "아니오", 2, "기존 등급2"])
            workbook.save(source)

            school_grades = root / "school-grades.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "안전"
            worksheet.append(["number", "영역", "filename", "contents", "page", "페이지전체내용", "사고사례여부", "등급", "등급사유"])
            workbook.save(school_grades)

            result = run_census(
                source,
                ncs_root,
                school_root,
                root / "result.xlsx",
                root / "report.md",
                ncs_grade_workbook=source,
                school_grade_workbook=school_grades,
                dashboard_data_out=root / "dashboard.js",
                expected={"documents": {"NCS": 86, "교과서": 9}},
                git={"commit": "abc1234", "dirty": False}, argv=["x"],
            )

            dashboard_text = (root / "dashboard.js").read_text(encoding="utf-8")

        included = [record for record in result.matches if record.decision == "included"]
        self.assertEqual(1, len(included))
        self.assertEqual((2, "existing"), (included[0].grade, included[0].grade_source))
        self.assertEqual(
            {"키워드 등록 워크북", "NCS 등급 워크북", "교과서 등급 워크북", "NCS Markdown", "교과서 Markdown"},
            {artifact.kind for artifact in result.input_artifacts},
        )
        self.assertIn('"denominator":"occurrences"', dashboard_text)

    def test_dashboard_data_uses_occurrences_as_the_grade_denominator(self):
        sources = [KeywordSource("안전", 1, True)]
        documents = [Document("NCS", Path("source.md"), "반도체개발/source.md", "<!-- page: 1 -->\n안전 안전\n")]
        result = aggregate_matches(
            sources,
            documents,
            [ExpressionRule("안전", "안전", "exact", "기존 키워드")],
            [],
        )
        result = assign_match_grades(
            result,
            {
                grade_lookup_key("NCS", "반도체개발/source.md", 1): GradeAssignment(
                    2, "형식적 언급", "기존", "existing"
                )
            },
        )

        payload = dashboard_payload(result)

        self.assertEqual(2, payload["corpora"]["NCS"]["total"])
        self.assertEqual({"real-page": 0, "existing": 2, "new": 0, "unpaged-context": 0, "unpaged-fallback": 0}, payload["corpora"]["NCS"]["grade_sources"])
        self.assertEqual({"1": 0, "2": 2, "3": 0, "unpaged": 0}, payload["corpora"]["NCS"]["grades"])
        self.assertEqual(2, payload["keywords"][0]["corpora"]["NCS"]["grades"]["2"])
        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "dashboard.js"
            write_dashboard_data(result, output)
            text = output.read_text(encoding="utf-8")
        self.assertTrue(text.startswith("/* Generated"))
        self.assertIn('"denominator":"occurrences"', text.replace(" ", "").replace("\n", ""))


def _doc(corpus, rel, text):
    return Document(corpus, Path(rel), rel, text)


class DictionaryVersionTests(unittest.TestCase):
    """semantic-expression-review §3.1 — 사전 버전(v1/v1fix/v2), 결함 2건, 조건부 규칙, 비정본 버전의 경로 거부."""

    def _scan(self, text, version, keywords=("안전", "PSM", "MSDS")):
        rules = build_default_rules(list(keywords), version=version)
        return scan_document(_doc("NCS", "반도체개발/LM1903060101_a/a.md", text), rules)

    def test_dictionary_versions_and_default(self):
        self.assertEqual(("v1", "v1fix", "v2"), SKR.DICTIONARY_VERSIONS)
        self.assertEqual("v2", SKR.DEFAULT_DICTIONARY)          # 결정 3 (연구책임자 2026-09-14): v2 채택
        self.assertEqual("v2", EXPECTED["dictionary"])
        with self.assertRaises(ValueError):
            build_default_rules(["안전"], version="v9")

    def test_v1fix_english_exact_keywords_need_word_boundary(self):
        text = "<!-- page: 1 -->\nEAPSM 마스크와 Htpsm 공정\nPSM 이행 점검\nMSDS 비치, MSDSX 아님\n"
        v1 = [(r.expression, r.decision, r.line) for r in self._scan(text, "v1") if r.keyword in ("PSM", "MSDS")]
        fix = [(r.expression, r.decision, r.line) for r in self._scan(text, "v1fix") if r.keyword in ("PSM", "MSDS")]
        self.assertEqual(2, sum(1 for e, d, l in v1 if e == "PSM" and d == "included" and l == 2))   # v1: 부분 문자열 2건
        self.assertEqual([], [x for x in fix if x[0] == "PSM" and x[2] == 2])                        # v1fix: 0건
        self.assertEqual(1, sum(1 for e, d, l in fix if e == "PSM" and d == "included" and l == 3))
        self.assertEqual(1, sum(1 for e, d, l in fix if e == "MSDS" and d == "included"))            # MSDSX 는 글자가 이어져 제외

    def test_v1fix_excludes_held_expressions_inside_안전(self):
        text = "<!-- page: 1 -->\n안전성 검토와 안전 마진 확보\n안전 보건 교육\n작업 안전 수칙\n"
        v1 = [r for r in self._scan(text, "v1") if r.keyword == "안전"]
        fix = [r for r in self._scan(text, "v1fix") if r.keyword == "안전"]
        self.assertEqual(3, sum(1 for r in v1 if r.decision == "included"))        # v1: 안전성·안전 마진·작업 안전 (안전 보건은 기존 제외)
        self.assertEqual(1, sum(1 for r in fix if r.decision == "included"))       # v1fix: 작업 안전 만
        held = [r for r in fix if r.decision == "excluded" and r.reason == "보류 표현 내부"]
        self.assertEqual([2, 2], sorted(r.line for r in held))
        self.assertEqual(1, sum(1 for r in fix if r.decision == "excluded" and r.reason != "보류 표현 내부"))   # 안전 보건 은 기존 사유 유지

    def test_v1_rule_content_digest_is_unchanged(self):
        # 2026-09-09 사전의 규칙 내용 지문 — v1 이 은근히 바뀌면 영향표의 기준선이 무너진다
        self.assertEqual(SKR.V1_RULE_CONTENT_SHA256, SKR.rule_content_sha256(build_default_rules(list(EXPECTED_KEYWORDS), version="v1")))

    def test_require_patterns_exclude_without_companion_in_window(self):
        rule = ExpressionRule("보호구", "방진복", "specific", "테스트", require_patterns=(r"착용", r"보호"))
        doc = _doc("NCS", "x/LM1903060101_a/a.md", "<!-- page: 1 -->\n방진복 세탁 주기\n\n방진복 입장\n반드시 착용한다\n<!-- page: 2 -->\n보호 장비\n<!-- page: 3 -->\n방진복 규격\n")
        got = {(r.line, r.decision, r.reason) for r in scan_document(doc, [rule])}
        self.assertIn((2, "excluded", "안전 문맥 동반어 없음"), got)      # 같은 줄·앞뒤 줄에 동반어 없음
        self.assertIn((4, "included", "테스트"), got)                     # 뒷줄 '착용'
        self.assertIn((9, "excluded", "안전 문맥 동반어 없음"), got)      # 앞 블록(7행 '보호')은 페이지 경계 밖
        before = _doc("NCS", "x/LM1903060101_a/a.md", "<!-- page: 1 -->\n반드시 착용한다\n방진복 입장\n")
        self.assertIn((3, "included", "테스트"), {(r.line, r.decision, r.reason) for r in scan_document(before, [rule])})   # 앞줄의 동반어도 창 안

    def test_v2_overrides_apply_held_and_require(self):
        with mock.patch.dict(SKR._V2_OVERRIDES, {("보호구", "방진복"): {"require_patterns": (r"착용",)}, ("인화", "combustible"): {"decision": "held"}}, clear=True):
            v2 = build_default_rules(["보호구", "인화"], version="v2")
            fix = build_default_rules(["보호구", "인화"], version="v1fix")
            v2_cands = default_candidate_decisions(version="v2")
        self.assertEqual((r"착용",), next(r for r in v2 if r.expression == "방진복").require_patterns)
        self.assertEqual((), next(r for r in fix if r.expression == "방진복").require_patterns)
        self.assertFalse(any(r.expression == "combustible" for r in v2))
        self.assertTrue(any(r.expression == "combustible" for r in fix))
        self.assertEqual("held", next(c for c in v2_cands if c.expression == "combustible").decision)

    def test_v2_decision_2_contents(self):
        """결정 2 (연구책임자 2026-09-14, 계층별 처방): 뜻이 다른 표현만 고친다 — 보류 방진화·케미컬, 조건부 방진복·장갑·X선·PSM. 동의어·가연성·combustible 은 그대로."""
        v2 = {(c.keyword, c.expression): c for c in default_candidate_decisions(version="v2")}
        fix = {(c.keyword, c.expression): c for c in default_candidate_decisions(version="v1fix")}
        self.assertEqual("held", v2[("보호구", "방진화")].decision); self.assertEqual("included", fix[("보호구", "방진화")].decision)
        self.assertEqual("held", v2[("화학물질", "케미컬")].decision)
        for key in (("보호구", "방진복"), ("보호구", "장갑"), ("방사선", "X선")):
            self.assertEqual(SKR.SAFETY_COMPANIONS, v2[key].require_patterns, key); self.assertEqual((), fix[key].require_patterns, key)
        for key in (("화학물질", "화학약품"), ("화학물질", "화학 물질"), ("누출", "누설"), ("작업환경", "작업 환경"), ("인화", "가연성"), ("인화", "combustible")):
            self.assertEqual(fix[key], v2[key], key)
        self.assertEqual("held", v2[("방사선", "자외선")].decision)
        rules = {r.expression: r for r in build_default_rules(["PSM", "보호구"], version="v2")}
        self.assertEqual(SKR.SAFETY_COMPANIONS, rules["PSM"].require_patterns)                   # 정확 키워드에도 조건부가 걸린다
        self.assertEqual((), next(r for r in build_default_rules(["PSM"], version="v1fix") if r.expression == "PSM").require_patterns)
        self.assertNotEqual(SKR.rule_content_sha256(build_default_rules(list(SKR.EXPECTED_KEYWORDS), "v1fix")), SKR.rule_content_sha256(build_default_rules(list(SKR.EXPECTED_KEYWORDS), "v2")))

    def test_v2_exact_psm_counts_only_with_companion(self):
        text = "<!-- page: 1 -->\nPSM(phase shift mask) 종류\n감광제 종류\n\nPSM 이행 점검\n위험성 평가 기록\n"
        v2 = [r for r in self._scan(text, "v2", keywords=("PSM",)) if r.keyword == "PSM"]
        self.assertEqual([(2, "excluded", SKR.NO_COMPANION_REASON), (5, "included", "기존 키워드의 정확 문자열")], [(r.line, r.decision, r.reason) for r in v2])
        fix = [r for r in self._scan(text, "v1fix", keywords=("PSM",)) if r.keyword == "PSM"]
        self.assertEqual(2, sum(1 for r in fix if r.decision == "included"))

    def test_non_default_dictionary_refuses_tracked_outputs_and_records_mismatch(self):
        # 거부 경로는 실제 저장소가 아니라 임시 디렉터리 안의 "같은 모양" 경로로 만든다 (HERE 를 임시로 바꿈).
        # 2026-09-14 에 이 테스트가 실제 data/semantic_keyword_recount_20260914.xlsx 를 fixture 로 덮어쓴 적이 있다 —
        # 거부가 실패하면 테스트가 실패해야지, 정본 산출물이 지워지면 안 된다.
        with tempfile.TemporaryDirectory() as td, mock.patch.object(SKR, "HERE", Path(td)):
            kw = census_fixture(Path(td))
            (Path(td) / "docs/03-analysis/data").mkdir(parents=True)
            with self.assertRaises(ValueError) as ctx:
                run_census(**dict(kw, summary_out=Path(td) / "docs/03-analysis/data/semantic_summary.json"),
                           dictionary="v1", expected={"documents": {"NCS": 86, "교과서": 9}}, git={"commit": "x", "dirty": False}, argv=["x"])
            self.assertIn("v1", str(ctx.exception))
            self.assertFalse((Path(td) / "docs/03-analysis/data/semantic_summary.json").exists())
            with self.assertRaises(ValueError):
                run_census(**dict(kw, xlsx_out=Path(td) / "semantic_keyword_recount_20260914.xlsx"), dictionary="v1fix",
                           expected={"documents": {"NCS": 86, "교과서": 9}}, git={"commit": "x", "dirty": False}, argv=["x"])
            self.assertFalse((Path(td) / "semantic_keyword_recount_20260914.xlsx").exists())
            run_census(**kw, dictionary="v1", summary_out=Path(td) / "s.json",
                       expected={"documents": {"NCS": 86, "교과서": 9}, "totals": {"NCS": 999, "교과서": 0}},
                       git={"commit": "x", "dirty": False}, argv=["x"])
            summary = json.loads((Path(td) / "s.json").read_text(encoding="utf-8"))
        self.assertEqual("v1", summary["meta"]["run"]["dictionary"])
        self.assertIsNone(summary["meta"]["run"]["expected"])                    # 비정본 버전: 불일치는 기록만
        self.assertEqual((False, True), (summary["meta"]["run"]["force"], summary["meta"]["run"]["variant"]))   # --force 를 준 적 없다 — 변형 실행이라 기록만 한 것 (적대적 리뷰 F12)
        self.assertTrue(any(m.startswith("totals.NCS") for m in summary["meta"]["run"]["expected_mismatch"]))


def graded_result(text="<!-- page: 1 -->\n안전 안전\n", corpus="NCS", rel="반도체개발/LM1903060101_a/a.md"):
    """등급까지 붙은 최소 결과 — 여러 테스트 클래스가 공유하는 fixture (RemediationTests._graded_result 와 같다)."""
    result = aggregate_matches(
        [KeywordSource("안전", 1, True)],
        [_doc(corpus, rel, text)],
        [ExpressionRule("안전", "안전", "exact", "기존 키워드")],
        [CandidateDecision("안전", "안전성", "held", "equivalent", "문맥 혼재")],
    )
    return assign_match_grades(result, {})


def census_fixture(root, ncs_body="<!-- page: 1 -->\n안전 안내\n"):
    """run_census 용 임시 말뭉치·워크북 kwargs — RemediationTests._census_fixture 의 모듈 수준 이름."""
    return RemediationTests._census_fixture(RemediationTests(), root, ncs_body)


class RemediationTests(unittest.TestCase):
    """2026-09-13 외부감사 시정(semantic-recount-remediation) — 코퍼스 규칙·가드·manifest·산출물."""

    def test_report_path_regex_is_gone(self):
        self.assertFalse(hasattr(SKR, "REPORT_PATH_RE"))

    def test_select_ncs_documents_rejects_file_without_lm_code(self):
        docs = [_doc("NCS", "반도체개발/LM1903060101_a/a.md", "x"), _doc("NCS", "반도체재료/report(1)/r.md", "x")]
        with self.assertRaises(ValueError) as ctx:
            select_ncs_documents(docs)
        self.assertIn("report(1)/r.md", str(ctx.exception))

    def test_select_ncs_documents_keeps_marker_rich_duplicate_and_records_dedup(self):
        sparse = _doc("NCS", "반도체제조/LM1903060205_MI 장비 운영/LM1903060205_MI 장비 운영.md", "본문\n")
        dense = _doc("NCS", "반도체제조/LM1903060205_MI_장비_운영/LM1903060205_MI_장비_운영.md", "<!-- page: 1 -->\n본문\n<!-- page: 2 -->\n")
        other = _doc("NCS", "반도체개발/LM1903060101_a/a.md", "<!-- page: 1 -->\n")
        kept, dedup = select_ncs_documents([sparse, dense, other])
        self.assertEqual([other.relative_path, dense.relative_path], sorted(d.relative_path for d in kept))
        self.assertEqual([DedupRecord("LM1903060205", dense.relative_path, (sparse.relative_path,))], dedup)

    def test_select_ncs_documents_tie_prefers_underscore_path(self):
        a = _doc("NCS", "x/LM1903060205_MI 장비/a.md", "<!-- page: 1 -->\n")
        b = _doc("NCS", "x/LM1903060205_MI_장비/b.md", "<!-- page: 1 -->\n")
        kept, dedup = select_ncs_documents([a, b])
        self.assertEqual([b.relative_path], [d.relative_path for d in kept])
        self.assertEqual((a.relative_path,), dedup[0].dropped)

    def test_check_marker_base_reports_zero_based_files(self):
        zero = _doc("NCS", "LM1903060408_x/x.md", "<!-- page: 0 -->\n본문\n<!-- page: 1 -->\n")
        one = _doc("NCS", "LM1903060101_y/y.md", "<!-- page: 1 -->\n본문\n")
        none = _doc("NCS", "LM1903060102_z/z.md", "본문만\n")
        self.assertEqual(["LM1903060408_x/x.md: 마커 0"], check_marker_base([zero, one, none]))

    def test_expected_pins_unpaged_to_zero_and_documents_to_86_and_9(self):
        self.assertEqual({"NCS": 86, "교과서": 9}, EXPECTED["documents"])
        self.assertEqual(0, EXPECTED["grades"]["NCS"]["unpaged"])
        self.assertEqual(0, EXPECTED["grades"]["교과서"]["unpaged"])

    def test_expected_block_is_fully_pinned(self):
        def leaves(value):
            if isinstance(value, dict):
                for v in value.values():
                    yield from leaves(v)
            else:
                yield value
        self.assertNotIn(None, list(leaves(EXPECTED)))

    def _graded_result(self, text="<!-- page: 1 -->\n안전 안전\n", corpus="NCS", rel="반도체개발/LM1903060101_a/a.md"):
        return graded_result(text, corpus, rel)

    def test_summary_metrics_and_check_expected_list_every_mismatch_and_skip_none(self):
        result = self._graded_result()
        result = replace(result, dedup=(DedupRecord("LM1", "k", ("d",)),))
        metrics = summary_metrics(result, artifact_manifest(result))
        self.assertEqual({"NCS": 1, "교과서": 0}, metrics["documents"])
        self.assertEqual({"NCS": 2, "교과서": 0}, metrics["totals"])
        self.assertEqual({"1": 2, "2": 0, "3": 0, "unpaged": 0}, metrics["grades"]["NCS"])
        self.assertEqual({"real-page": 0, "existing": 0, "new": 2, "unpaged-context": 0, "unpaged-fallback": 0}, metrics["grade_sources"])
        self.assertEqual({"held": 1}, metrics["candidates"])
        self.assertEqual({"LM1": 1}, metrics["dedup"])
        expected = {"documents": {"NCS": 2, "교과서": None}, "totals": {"NCS": 2, "교과서": 0}, "rule_sha256": "nope"}
        bad = check_expected(metrics, expected)
        self.assertEqual(2, len(bad), bad)
        self.assertTrue(any(b.startswith("documents.NCS: 1 != 2") for b in bad), bad)
        self.assertTrue(any(b.startswith("rule_sha256:") for b in bad), bad)
        self.assertEqual([], check_expected(metrics, {"totals": {"NCS": 2}}))

    def test_check_expected_reports_extra_keys_in_strict_groups(self):
        # 적대적 리뷰: 새 중복 코드나 새 판정 상태가 끼어들어도 총계·해시가 그대로면 가드가 못 보던 구멍
        metrics = {"dedup": {"LM1": 1, "LM9": 4}, "candidates": {"held": 1, "weird": 2}, "totals": {"NCS": 2, "교과서": 0}}
        bad = check_expected(metrics, {"dedup": {"LM1": 1}, "candidates": {"held": 1}, "totals": {"NCS": 2}})
        self.assertEqual(2, len(bad), bad)
        self.assertTrue(any(b.startswith("dedup.LM9: 4 != (absent)") for b in bad), bad)
        self.assertTrue(any(b.startswith("candidates.weird: 2 != (absent)") for b in bad), bad)
        self.assertEqual([], check_expected({"grades": {"NCS": {"1": 1}}}, {"grades": {"NCS": {"1": 1}}}))
        self.assertEqual(["grades.NCS.1: None != 1"], check_expected({}, {"grades": {"NCS": {"1": 1}}}))

    def _census_fixture(self, root, ncs_body="<!-- page: 1 -->\n안전 안내\n"):
        ncs_root = root / "ncs"
        school_root = root / "school"
        ncs_root.mkdir()
        school_root.mkdir()
        (ncs_root / "반도체개발").mkdir()
        (ncs_root / "반도체개발" / "LM1903060101_안전.md").write_text(ncs_body, encoding="utf-8")
        for index in range(1, 86):
            (ncs_root / "반도체개발" / f"LM19030602{index:02d}_book.md").write_text("", encoding="utf-8")
        for index in range(9):
            (school_root / f"school-{index}.md").write_text("", encoding="utf-8")
        source = root / "source.xlsx"
        workbook = Workbook()
        workbook.remove(workbook.active)
        for keyword in EXPECTED_KEYWORDS:
            worksheet = workbook.create_sheet(keyword)
            worksheet.append(["number", "영역", "filename", "contents", "page", "페이지전체내용", "사고사례여부", "등급", "등급사유"])
        workbook.save(source)
        school = root / "school-grades.xlsx"
        workbook = Workbook()
        workbook.active.title = "안전"
        workbook.active.append(["number", "영역", "filename", "contents", "page", "페이지전체내용", "사고사례여부", "등급", "등급사유"])
        workbook.save(school)
        return dict(source_workbook=source, ncs_root=ncs_root, school_root=school_root,
                    xlsx_out=root / "out.xlsx", report_out=root / "out.md",
                    ncs_grade_workbook=source, school_grade_workbook=school)

    def test_run_census_checks_document_counts_against_expected(self):
        with tempfile.TemporaryDirectory() as td:
            kw = self._census_fixture(Path(td))
            (kw["ncs_root"] / "반도체개발" / "LM1903060285_book.md").unlink()
            with self.assertRaises(ValueError) as ctx:
                run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}})
            self.assertIn("NCS=85", str(ctx.exception))
            self.assertFalse(kw["xlsx_out"].exists())

    def test_run_census_refuses_zero_based_markers(self):
        with tempfile.TemporaryDirectory() as td:
            kw = self._census_fixture(Path(td), ncs_body="<!-- page: 0 -->\n안전\n")
            with self.assertRaises(ValueError) as ctx:
                run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}})
            self.assertIn("마커 0", str(ctx.exception))

    def test_run_census_refuses_to_write_on_expected_mismatch(self):
        with tempfile.TemporaryDirectory() as td:
            kw = self._census_fixture(Path(td))
            with self.assertRaises(SystemExit):
                run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}, "totals": {"NCS": 999, "교과서": 0}},
                           dashboard_data_out=Path(td) / "d.js", summary_out=Path(td) / "s.json")
            self.assertFalse(kw["xlsx_out"].exists())
            self.assertFalse(kw["report_out"].exists())
            self.assertFalse((Path(td) / "d.js").exists())
            self.assertFalse((Path(td) / "s.json").exists())

    def test_force_writes_and_records_mismatch_in_manifest(self):
        with tempfile.TemporaryDirectory() as td:
            kw = self._census_fixture(Path(td))
            result = run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}, "totals": {"NCS": 999, "교과서": 0}},
                                summary_out=Path(td) / "s.json", force=True, git={"commit": "abc1234", "dirty": False})
            summary = json.loads((Path(td) / "s.json").read_text(encoding="utf-8"))
            self.assertTrue(kw["xlsx_out"].exists())
            self.assertIsNone(summary["meta"]["run"]["expected"])
            self.assertTrue(summary["meta"]["run"]["force"])
            self.assertTrue(any(m.startswith("totals.NCS") for m in summary["meta"]["run"]["expected_mismatch"]))
            self.assertEqual(1, len([r for r in result.matches if r.decision == "included"]))

    def test_run_manifest_has_commit_command_inputs_versions_and_public_paths(self):
        result = self._graded_result()
        result = replace(result, dedup=(DedupRecord("LM1", "k.md", ("d.md",)),))
        run = run_manifest(result, argv=["semantic_keyword_recount.py", "--ncs-root", str(Path.home() / "x" / "ncs"),
                                         f"--school-root={Path.home() / 'y' / 'school'}", "--xlsx-out=/some/other/volume/out.xlsx"],
                           force=False, expected_mismatch=[], git={"commit": "abc1234", "dirty": True},
                           xlsx_out=Path("data/semantic_keyword_recount_20260914.xlsx"))
        self.assertEqual(("abc1234", True), (run["git_commit"], run["git_dirty"]))
        self.assertIn("--ncs-root ~/x/ncs", run["command"])
        self.assertIn("--school-root=~/y/school", run["command"])          # --opt=value 형태도 걷어낸다 (보안 리뷰)
        self.assertIn("--xlsx-out=out.xlsx", run["command"])
        self.assertNotIn(str(Path.home()), run["command"])
        self.assertEqual("semantic_keyword_recount_20260914.xlsx", run["xlsx"])
        self.assertRegex(run["python"], r"^\d+\.\d+")
        self.assertRegex(run["openpyxl"], r"^\d+\.\d+")
        self.assertTrue(run["expected"])
        self.assertEqual((False, False), (run["force"], run["variant"]))
        self.assertEqual([{"code": "LM1", "kept": "k.md", "dropped": ["d.md"]}], run["dedup"])
        variant = run_manifest(result, argv=["x"], force=False, expected_mismatch=["totals.NCS: 1 != 2"], git={"commit": "abc1234", "dirty": True}, variant=True)
        self.assertEqual((None, False, True), (variant["expected"], variant["force"], variant["variant"]))
        self.assertRegex(run["generated_at"], r"^\d{4}-\d{2}-\d{2}T")
        self.assertTrue(all("sha256" in item and "count" in item and "kind" in item for item in run["inputs"]))

    def test_public_path_never_leaks_home_or_repo_root(self):
        self.assertEqual("docs/x.json", public_path(Path(SKR.__file__).parent / "docs" / "x.json"))
        self.assertEqual("~/a/b.md", public_path(Path.home() / "a" / "b.md"))
        self.assertEqual("b.md", public_path("/some/other/volume/b.md"))

    def test_two_runs_on_same_fixture_are_identical(self):
        a, b = self._graded_result(), self._graded_result()
        self.assertEqual(artifact_manifest(a), artifact_manifest(b))
        pa, pb = dashboard_payload(a), dashboard_payload(b)
        pa["meta"].pop("generated"); pb["meta"].pop("generated")
        self.assertEqual(pa, pb)
        with tempfile.TemporaryDirectory() as td:
            sheets = []
            for name, result in (("a", a), ("b", b)):
                out = Path(td) / f"{name}.xlsx"
                write_workbook(result, out, run={"git_commit": "abc1234"})
                book = load_workbook(out, read_only=True)
                sheets.append({ws.title: [list(row) for row in ws.iter_rows(values_only=True)] for ws in book.worksheets})
            sheets[0]["README"] = [r for r in sheets[0]["README"] if r[0] != "생성일"]
            sheets[1]["README"] = [r for r in sheets[1]["README"] if r[0] != "생성일"]
            self.assertEqual(sheets[0], sheets[1])

    def test_summary_json_equals_dashboard_payload_and_copies_previous_basis(self):
        with tempfile.TemporaryDirectory() as td:
            kw = self._census_fixture(Path(td))
            reseg = Path(td) / "reseg_summary.json"
            reseg.write_text(json.dumps({"pages": 2189, "page_g": {"1": 1519, "2": 525, "3": 145}, "books": 86,
                                         "cases_pages": 13, "unresolved": {"pages": 51},
                                         "meta": {"expected": True, "run_at": "2026-09-07T08:48:51+00:00"}}), encoding="utf-8")
            run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}},
                       dashboard_data_out=Path(td) / "d.js", summary_out=Path(td) / "s.json",
                       analysis_dir=Path(td) / "docs",
                       previous_basis=reseg, git={"commit": "abc1234", "dirty": False})
            js = (Path(td) / "d.js").read_text(encoding="utf-8")
            summary = json.loads((Path(td) / "s.json").read_text(encoding="utf-8"))
            self.assertEqual({"NCS_키워드검색결과.html", "keyword-analysis.html", "교과서_키워드검색결과.html"},
                             {p.name for p in (Path(td) / "docs").iterdir()})
            self.assertEqual([], [p.name for p in Path(td).iterdir() if p.name.endswith(".tmp")])   # 원자적 쓰기 잔여물 없음
        payload = json.loads(re.search(r"window\.SEMANTIC_RECOUNT=(.*);\s*$", js, re.S).group(1))
        self.assertEqual(payload, summary)
        self.assertEqual({"source": "reseg_summary.json", "pages": 2189,
                          "page_g": {"1": 1519, "2": 525, "3": 145}, "books": 86, "cases_pages": 13, "unresolved_pages": 51,
                          "source_run_at": "2026-09-07T08:48:51+00:00"},
                         {k: summary["meta"]["previous_basis"][k] for k in ("source", "pages", "page_g", "books", "cases_pages", "unresolved_pages", "source_run_at")})
        self.assertTrue(any(i["kind"] == "이전 기준" and len(i["sha256"]) == 64 for i in summary["meta"]["run"]["inputs"]))
        self.assertIn("run", summary["meta"])
        self.assertIn("manifest", summary["meta"])
        self.assertNotIn("20260909", js.split("\n")[0])
        self.assertNotIn("/Users/", json.dumps(summary))
        self.assertNotIn("relative_path", json.dumps(summary))

    def test_unpaged_context_grade_is_computed_not_defaulted(self):
        # 문맥이 촘촘하면 등급 3 — 등급 1 이 기본값이 아니라 grade_page 의 결과임을 보인다
        text = "안전 안전 안전 안전 안전 안전 보호구 착용 착용 착용 착용 착용 점검 조치\n"
        result = aggregate_matches([KeywordSource("안전", 1, True)], [_doc("NCS", "s.md", text)],
                                   [ExpressionRule("안전", "안전", "exact", "기존 키워드")], [])
        graded = assign_match_grades(result, {})
        record = next(r for r in graded.matches if r.decision == "included")
        self.assertEqual("unpaged-context", record.grade_source)
        self.assertGreaterEqual(record.grade, 2)
        self.assertTrue(record.grade_reason.startswith("페이지 마커 없는 출현의 문맥 기준"))

    def test_unpaged_record_with_empty_context_falls_back_to_grade_one(self):
        result = aggregate_matches([KeywordSource("안전", 1, True)], [_doc("NCS", "s.md", "안전\n")],
                                   [ExpressionRule("안전", "안전", "exact", "기존 키워드")], [])
        record = next(r for r in result.matches if r.decision == "included")
        graded = assign_match_grades(replace(result, matches=(replace(record, context="   "),)), {})
        self.assertEqual((1, "unpaged-fallback"), (graded.matches[0].grade, graded.matches[0].grade_source))
        self.assertIn("문맥", graded.matches[0].grade_reason)

    def test_committed_summary_json_matches_expected(self):
        # CI 가 자기 증명이 되지 않게: 커밋된 정본 요약을 Python EXPECTED 와 직접 대조한다 (적대적 리뷰)
        S = json.loads((Path(SKR.__file__).parent / "docs/03-analysis/data/semantic_summary.json").read_text(encoding="utf-8"))
        C = ("NCS", "교과서")
        run = S["meta"]["run"]
        metrics = {
            "dictionary": run["dictionary"],
            "page_basis": S["meta"]["page_basis"],                                                          # occurrence-real-pages: 정본은 실제 쪽 기준
            "page_maps_sha256": (run.get("page_maps") or {}).get("sha256"),
            "reseg_agreement": run.get("reseg_agreement"),
            "documents": {c: S["corpora"][c]["documents"] for c in C},
            "totals": {c: S["corpora"][c]["total"] for c in C},
            "grades": {c: S["corpora"][c]["grades"] for c in C},
            "grade_sources": {k: sum(S["corpora"][c]["grade_sources"].get(k, 0) for c in C) for k in SKR.GRADE_SOURCES},
            "candidates": S["status"],
            "dedup": {d["code"]: len(d["dropped"]) for d in run["dedup"]},
            **S["meta"]["manifest"],
        }
        self.assertEqual([], check_expected(metrics))
        self.assertTrue(S["meta"]["run"]["expected"])

    def test_force_does_not_bypass_document_count_guard(self):
        with tempfile.TemporaryDirectory() as td:
            kw = self._census_fixture(Path(td))
            (kw["ncs_root"] / "반도체개발" / "LM1903060285_book.md").unlink()
            with self.assertRaises(ValueError):
                run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}}, force=True, git={"commit": "x", "dirty": False}, argv=["x"])
            self.assertFalse(kw["xlsx_out"].exists())

    def test_run_census_rejects_missing_roots_and_workbook_readably(self):
        with tempfile.TemporaryDirectory() as td:
            kw = self._census_fixture(Path(td))
            with self.assertRaises(FileNotFoundError) as ctx:
                run_census(**dict(kw, ncs_root=Path(td) / "nope"), expected={"documents": {"NCS": 86, "교과서": 9}})
            self.assertIn("NCS", str(ctx.exception))
            with self.assertRaises(FileNotFoundError) as ctx:
                run_census(**dict(kw, source_workbook=Path(td) / "missing.xlsx"), expected={"documents": {"NCS": 86, "교과서": 9}})
            self.assertIn("키워드", str(ctx.exception))

    def test_git_info_degrades_when_git_is_unavailable(self):
        from unittest import mock
        with mock.patch.object(SKR.subprocess, "run", side_effect=OSError("no git")):
            self.assertEqual({"commit": "unknown", "dirty": None}, SKR._git_info())

    def test_public_path_exact_base_and_overrides(self):
        self.assertEqual(".", public_path(SKR.HERE))
        self.assertEqual("~", public_path(Path.home()))
        with tempfile.TemporaryDirectory() as td:
            self.assertEqual("x/y.md", public_path(Path(td) / "x" / "y.md", here=Path(td)))
            self.assertEqual("~/z.md", public_path(Path(td) / "z.md", here=Path("/nonexistent-base"), home=Path(td)))

    def test_check_marker_base_reports_any_marker_below_one_and_nonmonotone(self):
        late_zero = _doc("NCS", "LM1903060101_a/a.md", "<!-- page: 1 -->\n본문\n<!-- page: 0 -->\n")
        self.assertEqual(["LM1903060101_a/a.md: 마커 0"], check_marker_base([late_zero]))
        wobbly = _doc("NCS", "LM1903060102_b/b.md", "<!-- page: 3 -->\n가\n<!-- page: 2 -->\n나\n")
        self.assertEqual([], check_marker_base([wobbly]))                       # 비단조는 거부 대상이 아니다 (레거시 목차 마커)
        self.assertEqual(["LM1903060102_b/b.md", "LM1903060101_a/a.md"], SKR.nonmonotone_markers([wobbly, late_zero]))   # 1→0 도 감소

    def test_ncs_code_regex_needs_a_boundary(self):
        self.assertIsNone(SKR._NCS_CODE_RE.search("반도체개발/PLM1903060101_x/x.md"))
        self.assertIsNotNone(SKR._NCS_CODE_RE.search("반도체개발/LM1903060101_x/x.md"))
        kept, dedup = select_ncs_documents([_doc("NCS", "x/lm1903060101_a/a.md", "<!-- page: 1 -->\n"), _doc("NCS", "x/LM1903060101_b/b.md", "")])
        self.assertEqual(1, len(kept)); self.assertEqual("LM1903060101", dedup[0].code)

    def test_summary_payload_has_keyword_groups_and_group_pages(self):
        """hwpx-ncs-section-refresh D1·D2 — 키워드×그룹 총계·등급, 그룹의 교재 실제 쪽수(마커 최대값 합)."""
        docs = [_doc("NCS", "반도체개발/LM1903060101_a/a.md", "<!-- page: 1 -->\n안전 위험\n<!-- page: 3 -->\n안전\n"),
                _doc("NCS", "반도체장비/LM1903060301_b/b.md", "<!-- page: 1 -->\n위험\n<!-- page: 2 -->\n\n"),
                _doc("NCS", "반도체장비/LM1903060302_c/c.md", "안전\n")]          # 마커 없음 → 0쪽, 미확정 출현
        result = assign_match_grades(aggregate_matches([KeywordSource("안전", 1, True), KeywordSource("위험", 1, True)], docs,
                                                       [ExpressionRule("안전", "안전", "exact", "기존"), ExpressionRule("위험", "위험", "exact", "기존")], []), {})
        payload = summary_payload(result)
        groups = {g["name"]: g for g in payload["corpora"]["NCS"]["groups"]}
        self.assertEqual({"반도체개발": 3, "반도체장비": 2}, {n: g["pages"] for n, g in groups.items()})          # max marker 3 / 2 + 0
        kw = {k["name"]: k["corpora"]["NCS"] for k in payload["keywords"]}
        self.assertEqual([("반도체개발", 2), ("반도체장비", 1)], [(g["name"], g["total"]) for g in kw["안전"]["groups"]])
        self.assertEqual([("반도체개발", 1), ("반도체장비", 1)], [(g["name"], g["total"]) for g in kw["위험"]["groups"]])
        for name, c in kw.items():
            self.assertEqual(c["total"], sum(g["total"] for g in c["groups"]), name)
            for grade in ("1", "2", "3", "unpaged"):
                self.assertEqual(c["grades"][grade], sum(g["grades"][grade] for g in c["groups"]), (name, grade))
        self.assertEqual(1, kw["안전"]["groups"][1]["grades"]["unpaged"] + kw["안전"]["groups"][1]["grades"]["1"])   # c.md 의 안전 1건 (unpaged-context 또는 등급1)
        self.assertEqual([g["name"] for g in payload["corpora"]["NCS"]["groups"]], [g["name"] for g in kw["안전"]["groups"]])   # 그룹 순서 동일

    def test_analysis_pages_escape_document_text_and_have_scroll_regions(self):
        result = self._graded_result(text="<!-- page: 1 -->\n안전 <script>&\"x\"\n")
        payload = summary_payload(result)
        with tempfile.TemporaryDirectory() as td:
            write_analysis_pages(result, payload, Path(td))
            ncs = (Path(td) / "NCS_키워드검색결과.html").read_text(encoding="utf-8")
            index = (Path(td) / "keyword-analysis.html").read_text(encoding="utf-8")
        self.assertIn("&lt;script&gt;&amp;&quot;x&quot;", ncs)
        self.assertNotIn("<script>", ncs)
        self.assertNotIn("git", index)                                   # run 없음 → 문서 수만
        self.assertEqual(2, ncs.count('role="region"'))                   # 요약표 1 + 키워드별 상세표 1 (키워드 1개) — 모두 접근 가능한 스크롤 영역
        self.assertIn("@media", ncs)

    def test_analysis_pages_group_detail_rows_per_keyword_in_closed_details(self):
        # 성능 리뷰: 12,536행을 한 표로 그리면 모바일 첫 화면 8.1s — 키워드별 닫힌 <details> 는 레이아웃에서 빠진다 (연구책임자 D1)
        text = "<!-- page: 1 -->\n안전 보건 안전\n"
        result = aggregate_matches([KeywordSource("안전", 1, True), KeywordSource("보건", 1, True)], [_doc("NCS", "반도체개발/LM1903060101_a/a.md", text)],
                                   [ExpressionRule("안전", "안전", "exact", "기존 키워드"), ExpressionRule("보건", "보건", "exact", "기존 키워드")], [])
        result = assign_match_grades(result, {})
        with tempfile.TemporaryDirectory() as td:
            write_analysis_pages(result, summary_payload(result), Path(td))
            ncs = (Path(td) / "NCS_키워드검색결과.html").read_text(encoding="utf-8")
        self.assertEqual(2, ncs.count("<details"))
        self.assertNotIn("<details open", ncs)
        self.assertIn("<summary>안전 (2건)</summary>", ncs)
        self.assertIn("<summary>보건 (1건)</summary>", ncs)
        self.assertIn("검색결과 전체: <strong>3건</strong>", ncs)
        self.assertLess(ncs.index("<summary>안전"), ncs.index("<summary>보건"))    # 출현 많은 키워드가 먼저

    def test_main_passes_new_flags_and_prints_metrics(self):
        from unittest import mock
        seen = {}
        def fake_run_census(*args, **kwargs):
            seen.update(kwargs); return self._graded_result()
        argv = ["semantic_keyword_recount.py", "--source-workbook", "s.xlsx", "--ncs-root", "n", "--school-root", "t",
                "--xlsx-out", "o.xlsx", "--report-out", "r.md", "--summary-out", "s.json", "--analysis-dir", "d",
                "--previous-basis", "p.json", "--force"]
        import io, contextlib
        buf = io.StringIO()
        with mock.patch.object(SKR, "run_census", fake_run_census), mock.patch.object(SKR.sys, "argv", argv), contextlib.redirect_stdout(buf):
            SKR.main()
        self.assertEqual((Path("s.json"), Path("d"), Path("p.json"), True), (seen["summary_out"], seen["analysis_dir"], seen["previous_basis"], seen["force"]))
        self.assertIn("측정값", buf.getvalue())

    def test_analysis_pages_cite_manifest_and_totals(self):
        result = self._graded_result()
        payload = summary_payload(result, run=run_manifest(result, argv=["x"], force=False, expected_mismatch=[],
                                                            git={"commit": "abc1234", "dirty": False},
                                                            xlsx_out=Path("data/semantic_keyword_recount_20260913.xlsx")))
        with tempfile.TemporaryDirectory() as td:
            written = write_analysis_pages(result, payload, Path(td))
            names = sorted(p.name for p in written)
            index = (Path(td) / "keyword-analysis.html").read_text(encoding="utf-8")
            ncs = (Path(td) / "NCS_키워드검색결과.html").read_text(encoding="utf-8")
        self.assertEqual(["NCS_키워드검색결과.html", "keyword-analysis.html", "교과서_키워드검색결과.html"], names)
        self.assertIn("semantic_keyword_recount_20260913.xlsx", index)
        self.assertIn("abc1234", index)
        self.assertIn("NCS 교재 1권", index)
        self.assertIn("검색결과 전체: <strong>2건</strong>", ncs)
        self.assertNotIn("report", index)


class DictionaryVersionAuditTests(unittest.TestCase):
    """ship 커버리지 감사(2026-09-14) — 사전 버전 경로에서 아직 안 짚인 가지: main() 의 --dictionary 전달, 잘못된 버전의 조기 거부,
    변형 거부의 나머지 두 라벨(report_out 기본 이름·analysis_dir 추적 경로), EXPECTED.dictionary 불일치, v1≡v1fix 후보 목록, 빈 문서의 그룹 쪽수."""

    def _main_with(self, argv_tail):
        import io, contextlib
        seen = {}
        def fake_run_census(*args, **kwargs):
            seen.update(kwargs); return graded_result()
        argv = ["semantic_keyword_recount.py", "--source-workbook", "s.xlsx", "--ncs-root", "n", "--school-root", "t",
                "--xlsx-out", "o.xlsx", "--report-out", "r.md"] + argv_tail
        out = io.StringIO()
        with mock.patch.object(SKR, "run_census", fake_run_census), mock.patch.object(SKR.sys, "argv", argv), \
                contextlib.redirect_stdout(out), contextlib.redirect_stderr(io.StringIO()):
            SKR.main()
        self.assertIn(f'"dictionary": "{seen["dictionary"]}"', out.getvalue())      # 측정값 블록도 실제 돌린 사전을 적는다 (ship 리뷰)
        return seen["dictionary"]

    def test_main_passes_dictionary_to_run_census_and_rejects_unknown_choice(self):
        """--dictionary 는 run_census 로 그대로 간다; 생략하면 정본(v2); 목록 밖 값은 argparse 가 막는다."""
        self.assertEqual("v1fix", self._main_with(["--dictionary", "v1fix"]))
        self.assertEqual(SKR.DEFAULT_DICTIONARY, self._main_with([]))
        with self.assertRaises(SystemExit):
            self._main_with(["--dictionary", "v9"])

    def test_run_census_rejects_unknown_version_before_reading_inputs(self):
        """버전 검사가 파일 검사보다 앞이다 — 존재하지 않는 입력으로도 ValueError(버전) 가 먼저 난다."""
        with self.assertRaises(ValueError) as ctx:
            run_census(Path("없는.xlsx"), Path("없는-ncs"), Path("없는-school"), Path("o.xlsx"), Path("r.md"), dictionary="v9")
        self.assertIn("v9", str(ctx.exception))
        with self.assertRaises(ValueError):
            default_candidate_decisions(version="v9")

    def test_variant_refuses_report_default_name_and_analysis_dir_under_docs(self):
        """변형 실행이 거부하는 나머지 두 자리 — report_out 의 정본 기본 이름, analysis_dir 의 docs/ 하위. 거부 뒤에는 아무것도 쓰지 않는다."""
        with tempfile.TemporaryDirectory() as td, mock.patch.object(SKR, "HERE", Path(td)):
            kw = census_fixture(Path(td))
            expected = {"documents": {"NCS": 86, "교과서": 9}}
            with self.assertRaises(ValueError) as ctx:
                run_census(**dict(kw, report_out=Path(td) / "semantic_keyword_recount_20260914_report.md"), dictionary="v1fix",
                           expected=expected, git={"commit": "x", "dirty": False}, argv=["x"])
            self.assertIn("report_out", str(ctx.exception))
            self.assertFalse((Path(td) / "semantic_keyword_recount_20260914_report.md").exists())
            self.assertFalse(kw["xlsx_out"].exists())
            (Path(td) / "docs").mkdir()
            with self.assertRaises(ValueError) as ctx:
                run_census(**kw, dictionary="v1", analysis_dir=Path(td) / "docs" / "pages", expected=expected, git={"commit": "x", "dirty": False}, argv=["x"])
            self.assertIn("analysis_dir", str(ctx.exception))
            self.assertEqual([], list((Path(td) / "docs").iterdir()))
            self.assertFalse(kw["xlsx_out"].exists())

    def test_check_expected_catches_dictionary_mismatch(self):
        """summary_metrics 의 dictionary 는 EXPECTED 와 글자 그대로 비교된다 — v1fix 실행을 v2 정본으로 굳힐 수 없다."""
        result = graded_result()
        metrics = summary_metrics(result, artifact_manifest(result), dictionary="v1fix")
        self.assertEqual("v1fix", metrics["dictionary"])
        self.assertIn("dictionary: v1fix != v2", check_expected(metrics, {"dictionary": "v2"}))
        self.assertEqual([], check_expected(summary_metrics(result, artifact_manifest(result)), {"dictionary": SKR.DEFAULT_DICTIONARY}))

    def test_v1_and_v1fix_share_candidate_decisions_and_v2_differs_only_in_overrides(self):
        """결함 2건(v1fix)은 정확 규칙 쪽이라 후보 목록은 v1 과 같다; v2 는 _V2_OVERRIDES 의 키에서만 다르다."""
        v1, fix, v2 = (default_candidate_decisions(version=v) for v in ("v1", "v1fix", "v2"))
        self.assertEqual(v1, fix)
        self.assertEqual(len(v1), len(v2))
        changed = {(a.keyword, a.expression) for a, b in zip(fix, v2) if a != b}
        self.assertEqual({(k, e) for k, e in SKR._V2_OVERRIDES if k != e}, changed)      # 키워드==표현(PSM) 은 정확 규칙 쪽에 적용된다
        self.assertTrue(all((c.keyword, c.expression) in SKR._V2_OVERRIDES or c.require_patterns == () for c in v2))

    def test_group_pages_zero_for_empty_document_and_max_marker_otherwise(self):
        """그룹 쪽수는 문서별 마커 최댓값의 합 — 본문이 빈 문서는 0 이고 max(()) 로 죽지 않는다; 마커가 역행해도 최댓값을 쓴다."""
        docs = [_doc("NCS", "반도체재료/LM1903060401_a/a.md", ""),
                _doc("NCS", "반도체재료/LM1903060402_b/b.md", "<!-- page: 5 -->\n안전\n<!-- page: 2 -->\n안전\n"),
                _doc("NCS", "반도체제조/LM1903060201_c/c.md", "   \n")]
        result = assign_match_grades(aggregate_matches([KeywordSource("안전", 1, True)], docs, [ExpressionRule("안전", "안전", "exact", "기존")], []), {})
        payload = summary_payload(result)
        groups = {g["name"]: g for g in payload["corpora"]["NCS"]["groups"]}
        self.assertEqual(5, groups["반도체재료"]["pages"])
        self.assertEqual(2, groups["반도체재료"]["documents"])
        self.assertEqual({"documents": 1, "pages": 0, "total": 0}, {k: groups["반도체제조"][k] for k in ("documents", "pages", "total")})   # 출현이 없는 그룹도 문서·쪽수 분모에 남는다 (Codex 적대적 리뷰)
        trailing = _doc("NCS", "반도체개발/LM1903060101_d/d.md", "<!-- page: 1 -->\n안전\n<!-- page: 100 -->\n")
        payload = summary_payload(assign_match_grades(aggregate_matches([KeywordSource("안전", 1, True)], [trailing], [ExpressionRule("안전", "안전", "exact", "기존")], []), {}))
        self.assertEqual(100, payload["corpora"]["NCS"]["groups"][0]["pages"])   # 본문 없는 마지막 마커도 쪽수에 든다
        marker_only = _doc("NCS", "반도체개발/LM1903060102_e/e.md", "<!-- page: 3 -->\n")
        payload = summary_payload(assign_match_grades(aggregate_matches([KeywordSource("안전", 1, True)], [marker_only], [ExpressionRule("안전", "안전", "exact", "기존")], []), {}))
        self.assertEqual(3, payload["corpora"]["NCS"]["groups"][0]["pages"])     # 마커만 있는 문서도 죽지 않는다


if __name__ == "__main__":
    unittest.main()


# ---------------------------------------------------------------- occurrence-real-pages (2026-09-15)
def _real_page_fixture(root: Path):
    """대응 있는 교재 1 + 대응 없는 명시 목록 교재 1 + 대응 없고 목록에도 없는 교재 1 — 설계 §3.1~3.3."""
    ncs = root / "ncs"; (ncs / "반도체개발").mkdir(parents=True); (ncs / "반도체재료").mkdir(parents=True)
    a_text = "<!-- page: 1 -->\n안전 안전 안전\n<!-- page: 2 -->\n안전 위험 보호구 착용 안전 안전 안전 방지 예방 환기 차단 대피 안전모\n안전 위험\n"
    (ncs / "반도체개발" / "LM1903060101_a.md").write_text(a_text, encoding="utf-8")
    b_text = "<!-- page: 1 -->\n안전\n<!-- page: 2 -->\n안전 안전\n"
    (ncs / "반도체재료" / "LM1903060408_b.md").write_text(b_text, encoding="utf-8")            # REAL_PAGE_MARKER_BOOKS 의 하나
    c_text = "<!-- page: 1 -->\n안전\n"
    (ncs / "반도체재료" / "LM1903060499_c.md").write_text(c_text, encoding="utf-8")            # 대응 없음·목록 없음
    maps = root / "maps"; maps.mkdir()
    # a: 6줄(split("\n") 기준, 끝의 빈 문자열 포함) — 블록 2 의 두 줄이 실제 쪽 12·13 으로 갈라진다
    (maps / "LM1903060101.pages.json").write_text(json.dumps({"md": "LM1903060101_a.md", "line_pages": [10, 10, 12, 12, 13, 13]}), encoding="utf-8")
    docs = SKR.load_documents(ncs, "NCS")
    return ncs, maps, docs


class RealPageTests(unittest.TestCase):
    """줄→실제 쪽 대응(occurrence-real-pages): 로딩·검증, 쪽 덧씌우기, 실제 쪽 판정, 이전 기준 결속, 분야 쪽수, manifest."""

    def _rules(self):
        return [ExpressionRule("안전", "안전", "exact", "기존 키워드")], [KeywordSource("안전", 1, True)], [CandidateDecision("안전", "안전성", "held", "equivalent", "문맥 혼재")]

    def test_load_page_maps_validates_and_requires_map_or_listed_book(self):
        with tempfile.TemporaryDirectory() as td:
            ncs, maps, docs = _real_page_fixture(Path(td))
            a, b, c = (next(d for d in docs if code in d.relative_path) for code in ("LM1903060101", "LM1903060408", "LM1903060499"))
            with self.assertRaisesRegex(ValueError, "LM1903060499"):
                SKR.load_page_maps(maps, docs)                                                   # 대응도 목록도 없다
            page_maps, info = SKR.load_page_maps(maps, [a, b])
            self.assertEqual({a.relative_path: (10, 10, 12, 12, 13)}, page_maps)                # b 는 목록 교재 — 표식 그대로; 파일 끝 개행 뒤 빈 원소는 매칭 규약(splitlines)에 없어 빠진다
            self.assertEqual((1, 5), (info.files, len(page_maps[a.relative_path])))
            self.assertRegex(info.sha256, r"^[0-9a-f]{64}$"); self.assertEqual(info.sha256, SKR.load_page_maps(maps, [a, b])[1].sha256)
            bad = Path(td) / "bad"; bad.mkdir()
            for name, payload, msg in (("md", {"md": "other.md", "line_pages": [10, 10, 12, 12, 13, 13]}, "다른 파일"),
                                       ("len", {"md": "LM1903060101_a.md", "line_pages": [10, 10, 12]}, "줄 수"),
                                       ("zero", {"md": "LM1903060101_a.md", "line_pages": [0, 10, 12, 12, 13, 13]}, "1 이상")):
                (bad / "LM1903060101.pages.json").write_text(json.dumps(payload), encoding="utf-8")
                with self.assertRaisesRegex(ValueError, msg, msg=name):
                    SKR.load_page_maps(bad, [a])
            b_gap = _doc("NCS", "반도체재료/LM1903060408_b.md", "<!-- page: 1 -->\n안전\n<!-- page: 3 -->\n안전\n")
            with self.assertRaisesRegex(ValueError, "LM1903060408"):
                SKR.load_page_maps(maps, [b_gap])                                                 # 목록 교재인데 표식이 빠졌다(2 없음)
            no_marker = _doc("NCS", "반도체재료/LM1903060477_empty.md", "")
            self.assertEqual(({}, 0), (SKR.load_page_maps(maps, [no_marker])[0], SKR.load_page_maps(maps, [no_marker])[1].files))   # 표식 없는 문서는 대상 아님
            self.assertEqual(({}, 0), (SKR.load_page_maps(maps, [_doc("교과서", "school.md", "<!-- page: 1 -->\n안전\n")])[0], 0))       # 교과서는 대상 아님

    def test_apply_page_maps_overlays_pages_after_matching(self):
        with tempfile.TemporaryDirectory() as td:
            ncs, maps, docs = _real_page_fixture(Path(td))
            docs = [d for d in docs if "LM1903060499" not in d.relative_path]
            rules, sources, candidates = self._rules()
            page_maps, _ = SKR.load_page_maps(maps, docs)
            plain = aggregate_matches(sources, docs, rules, candidates)
            real = aggregate_matches(sources, docs, rules, candidates, page_maps=page_maps)
            self.assertEqual(len(plain.matches), len(real.matches))                               # 매칭은 표식 블록 위에서 — 총계 불변
            a_pages = sorted({r.page for r in real.matches if "LM1903060101" in r.relative_path})
            self.assertEqual([10, 12, 13], a_pages)                                                # 블록 1→10, 블록 2 의 줄 4→12·줄 5→13
            self.assertEqual([1, 2], sorted({r.page for r in real.matches if "LM1903060408" in r.relative_path}))   # 목록 교재는 표식 그대로
            row = next(r for r in real.summary if r.corpus == "NCS" and r.keyword == "안전")
            self.assertEqual(5, row.page_count)                                                     # 검출 쪽 수는 실제 쪽 기준 (10·12·13 + 1·2)
            self.assertEqual(4, next(r for r in plain.summary if r.corpus == "NCS" and r.keyword == "안전").page_count)

    def test_real_page_grading_uses_page_text_and_ignores_workbook_labels(self):
        with tempfile.TemporaryDirectory() as td:
            ncs, maps, docs = _real_page_fixture(Path(td))
            docs = [d for d in docs if "LM1903060499" not in d.relative_path]
            rules, sources, candidates = self._rules()
            page_maps, _ = SKR.load_page_maps(maps, docs)
            result = aggregate_matches(sources, docs, rules, candidates, page_maps=page_maps)
            a = next(d for d in docs if "LM1903060101" in d.relative_path)
            existing = {SKR.grade_lookup_key("NCS", a.relative_path, 12): GradeAssignment(3, "구체적 대책", "워크북", "existing")}
            graded = assign_match_grades(result, existing, page_maps=page_maps)
            by_page = {}
            for r in graded.matches:
                if "LM1903060101" in r.relative_path: by_page.setdefault(r.page, set()).add((r.grade, r.grade_source))
            self.assertEqual({10: {(1, "real-page")}, 12: {(3, "real-page")}, 13: {(1, "real-page")}}, by_page)   # 쪽 12 는 본문 판정 3, 쪽 13 은 짧아 1 — 워크북 라벨은 무시
            self.assertTrue(all(r.grade_source == "real-page" for r in graded.matches if "LM1903060408" in r.relative_path))   # 목록 교재도 실제 쪽(표식 블록) 판정 — 워크북 라벨 없음
            plain = assign_match_grades(aggregate_matches(sources, docs, rules, candidates), existing)
            self.assertTrue(all(r.grade_source in ("new", "existing") for r in plain.matches))                                       # 대응 없이 부르면 현행 그대로
            self.assertIn("real-page", SKR.GRADE_SOURCES)

    def test_reseg_agreement_counts_shared_pages(self):
        with tempfile.TemporaryDirectory() as td:
            ncs, maps, docs = _real_page_fixture(Path(td))
            docs = [d for d in docs if "LM1903060499" not in d.relative_path]
            rules, sources, candidates = self._rules()
            page_maps, _ = SKR.load_page_maps(maps, docs)
            graded = assign_match_grades(aggregate_matches(sources, docs, rules, candidates, page_maps=page_maps), {}, page_maps=page_maps)
            csv_path = Path(td) / "reseg.csv"
            csv_path.write_text("﻿영역,교재,페이지,등급,출처\n반도체개발,LM1903060101_a,10,1,text\n반도체개발,LM1903060101_a,12,3,text-fallback\n반도체개발,LM1903060101_a,99,2,text\n반도체재료,LM1903060408_b,1,3,label\n", encoding="utf-8")
            self.assertEqual({"pages": 2, "agree": 2, "disagree": []}, SKR.reseg_agreement(graded, csv_path))          # label 출처(라벨 쪽)는 견주지 않는다
            csv_path.write_text("﻿영역,교재,페이지,등급\n반도체개발,LM1903060101_a,10,1\n반도체개발,LM1903060101_a,12,2\n", encoding="utf-8")
            agreement = SKR.reseg_agreement(graded, csv_path)
            self.assertEqual((2, 1), (agreement["pages"], agreement["agree"])); self.assertEqual([{"book": "LM1903060101", "page": 12, "ours": 3, "reseg": 2}], agreement["disagree"])

    def test_group_pages_use_pdf_pages_when_given(self):
        with tempfile.TemporaryDirectory() as td:
            ncs, maps, docs = _real_page_fixture(Path(td))
            docs = [d for d in docs if "LM1903060499" not in d.relative_path]
            rules, sources, candidates = self._rules()
            page_maps, _ = SKR.load_page_maps(maps, docs)
            graded = assign_match_grades(aggregate_matches(sources, docs, rules, candidates, page_maps=page_maps), {}, page_maps=page_maps)
            payload = SKR.dashboard_payload(graded, pdf_pages={"LM1903060101": 40})
            groups = {g["name"]: g["pages"] for g in payload["corpora"]["NCS"]["groups"]}
            self.assertEqual({"반도체개발": 40, "반도체재료": 2}, groups)                          # 대응 교재는 PDF 쪽수, 목록 교재는 표식 최댓값
            self.assertEqual({"반도체개발": 2, "반도체재료": 2}, {g["name"]: g["pages"] for g in SKR.dashboard_payload(graded)["corpora"]["NCS"]["groups"]})
            self.assertEqual({"NCS": "real", "교과서": "marker"}, payload["meta"]["page_basis"])
            self.assertEqual({"NCS": "marker", "교과서": "marker"}, SKR.dashboard_payload(assign_match_grades(aggregate_matches(sources, docs, rules, candidates), {}))["meta"]["page_basis"])
            with self.assertRaisesRegex(ValueError, "LM1903060101"):
                SKR.dashboard_payload(graded, pdf_pages={})                                          # 대응 교재인데 PDF 쪽수를 모른다

    def test_manifest_and_metrics_carry_page_maps_and_agreement(self):
        result = graded_result()
        info = SKR.PageMapsInfo(dir="data/markdown/ncs_paged", files=84, sha256="a" * 64)
        run = SKR.run_manifest(result, argv=["x"], force=False, expected_mismatch=[], git={"commit": "abc1234", "dirty": False},
                               page_maps=info, reseg_agreement={"pages": 3, "agree": 3, "disagree": []})
        self.assertEqual({"dir": "data/markdown/ncs_paged", "files": 84, "sha256": "a" * 64}, run["page_maps"])
        self.assertEqual({"pages": 3, "agree": 3}, run["reseg_agreement"]); self.assertEqual(list(SKR.REAL_PAGE_MARKER_BOOKS), run["real_page_marker_books"])
        metrics = SKR.summary_metrics(result, SKR.artifact_manifest(result), page_maps=info, reseg_agreement=run["reseg_agreement"], page_basis={"NCS": "real", "교과서": "marker"})
        self.assertEqual(("a" * 64, {"pages": 3, "agree": 3}, {"NCS": "real", "교과서": "marker"}), (metrics["page_maps_sha256"], metrics["reseg_agreement"], metrics["page_basis"]))
        self.assertEqual([], SKR.check_expected(metrics, {"page_maps_sha256": "a" * 64, "reseg_agreement": {"pages": 3, "agree": 3}, "page_basis": {"NCS": "real", "교과서": "marker"}}))
        self.assertEqual(["reseg_agreement.agree: 3 != 4"], SKR.check_expected(metrics, {"reseg_agreement": {"pages": 3, "agree": 4}}))
        plain = SKR.summary_metrics(result, SKR.artifact_manifest(result))
        self.assertEqual((None, None, {"NCS": "marker", "교과서": "marker"}), (plain["page_maps_sha256"], plain["reseg_agreement"], plain["page_basis"]))   # 대응 없는 실행은 정본 EXPECTED 와 어긋난다

    def test_run_census_wires_page_maps_agreement_and_pdf_pages(self):
        """run_census(page_maps_dir=, reseg_csv=): 대응 로딩 → 실제 쪽 판정 → 결속 → manifest·metrics·payload; 대응 없는 실행은 정본 EXPECTED(page_basis real)와 어긋난다."""
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            kw = census_fixture(root, ncs_body="<!-- page: 1 -->\n안전 안전 안전\n<!-- page: 2 -->\n안전 위험 보호구 착용 안전 안전 안전 방지 예방 환기 차단 대피 안전모\n안전 위험\n")
            maps = root / "maps"; maps.mkdir()
            (maps / "LM1903060101.pages.json").write_text(json.dumps({"md": "LM1903060101_안전.md", "line_pages": [10, 10, 12, 12, 13, 13]}), encoding="utf-8")
            csv_path = root / "reseg.csv"; csv_path.write_text("﻿영역,교재,페이지,등급\n반도체개발,LM1903060101_안전,10,1\n반도체개발,LM1903060101_안전,12,3\n", encoding="utf-8")
            reseg = root / "reseg_summary.json"
            reseg.write_text(json.dumps({"pages": 2189, "page_g": {"1": 1519, "2": 525, "3": 145}, "books": 86, "cases_pages": 13, "unresolved": {"pages": 51},
                                         "per_book": {"LM1903060101_안전": {"pdf_pages": 40}}, "meta": {"expected": True, "run_at": "2026-09-07T08:48:51+00:00"}}), encoding="utf-8")
            with mock.patch.object(SKR, "REAL_PAGE_MARKER_BOOKS", tuple(f"LM19030602{i:02d}" for i in range(1, 86))):   # 빈 fixture 교재는 표식이 없어 대상 밖이지만, 목록 규칙도 같이 검증
                with self.assertRaisesRegex(ValueError, "previous_basis"):
                    run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}}, page_maps_dir=maps, reseg_csv=csv_path, git={"commit": "x", "dirty": False}, argv=["x"])
                result = run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}}, page_maps_dir=maps, reseg_csv=csv_path, previous_basis=reseg,
                                    summary_out=root / "s.json", git={"commit": "x", "dirty": False}, argv=["x"])
            summary = json.loads((root / "s.json").read_text(encoding="utf-8"))
            run = summary["meta"]["run"]
            self.assertEqual((1, "maps"), (run["page_maps"]["files"], Path(run["page_maps"]["dir"]).name)); self.assertRegex(run["page_maps"]["sha256"], r"^[0-9a-f]{64}$")
            self.assertEqual({"pages": 2, "agree": 2}, run["reseg_agreement"])
            self.assertEqual({"NCS": "real", "교과서": "marker"}, summary["meta"]["page_basis"])
            ncs = summary["corpora"]["NCS"]
            self.assertEqual({"real-page": ncs["total"], "existing": 0, "new": 0, "unpaged-context": 0, "unpaged-fallback": 0}, ncs["grade_sources"])   # 30개 키워드 전부 실제 쪽 판정
            self.assertEqual(40, next(g["pages"] for g in ncs["groups"] if g["name"] == "반도체개발"))
            self.assertEqual(0, ncs["grades"]["unpaged"]); self.assertGreater(ncs["grades"]["3"], 0); self.assertGreater(ncs["grades"]["1"], 0)   # 쪽 12 는 등급3, 쪽 10·13 은 등급1
            metrics = SKR.summary_metrics(result, SKR.artifact_manifest(result))
            self.assertEqual({"NCS": "real", "교과서": "marker"}, metrics["page_basis"])
            # 대응 없이 돌린 실행은 page_basis 가 marker 라 정본 EXPECTED(real) 와 어긋난다 → 쓰지 않는다
            with self.assertRaises(SystemExit):
                run_census(**dict(kw, xlsx_out=root / "o2.xlsx", report_out=root / "o2.md"), expected={"documents": {"NCS": 86, "교과서": 9}, "page_basis": {"NCS": "real", "교과서": "marker"}},
                           git={"commit": "x", "dirty": False}, argv=["x"])
            self.assertFalse((root / "o2.xlsx").exists())

    def test_main_passes_page_maps_and_reseg_csv(self):
        seen = {}
        def fake_run_census(*args, **kwargs):
            seen.update(kwargs); return graded_result()
        argv = ["semantic_keyword_recount.py", "--source-workbook", "s.xlsx", "--ncs-root", "n", "--school-root", "t", "--xlsx-out", "o.xlsx", "--report-out", "r.md",
                "--page-maps", "maps", "--reseg-csv", "r.csv"]
        import io, contextlib
        with mock.patch.object(SKR, "run_census", fake_run_census), mock.patch.object(SKR.sys, "argv", argv), contextlib.redirect_stdout(io.StringIO()):
            SKR.main()
        self.assertEqual((Path("maps"), Path("r.csv")), (seen["page_maps_dir"], seen["reseg_csv"]))

    # ---- 출고 전 커버리지 감사 (2026-09-15): 대응 로딩·판정·결속·분야 쪽수·manifest 의 남은 분기
    def test_document_code_from_directory_or_missing(self):
        """_document_code: 파일명에 없으면 경로에서, 소문자는 대문자로, 글자 뒤에 붙은 PLM… 은 코드가 아니다; 코드 없는 표식 문서는 대응도 목록도 찾을 수 없어 거부."""
        self.assertEqual("LM1903060101", SKR._document_code(_doc("NCS", "반도체개발/LM1903060101_a/a.md", "")))       # 디렉터리에서
        self.assertEqual("LM1903060101", SKR._document_code(_doc("NCS", "반도체개발/lm1903060101_a.md", "")))          # 대문자로 정규화
        self.assertIsNone(SKR._document_code(_doc("NCS", "반도체개발/PLM1903060101_a.md", "")))                      # 글자 뒤에 붙은 PLM… 은 코드가 아니다
        with tempfile.TemporaryDirectory() as td:
            maps = Path(td) / "maps"; maps.mkdir()
            with self.assertRaisesRegex(ValueError, r"\(None\)"):
                SKR.load_page_maps(maps, [_doc("NCS", "반도체개발/nocode/x.md", "<!-- page: 1 -->\n안전\n")])

    def test_load_page_maps_rejects_non_list_bool_and_string_page_values(self):
        with tempfile.TemporaryDirectory() as td:
            ncs, maps, docs = _real_page_fixture(Path(td))
            a = next(d for d in docs if "LM1903060101" in d.relative_path)
            bad = Path(td) / "bad"; bad.mkdir()
            for name, payload, msg in (("notlist", {"md": "LM1903060101_a.md", "line_pages": {"1": 10}}, r"\? ≠ 6"),
                                       ("missing", {"md": "LM1903060101_a.md"}, r"\? ≠ 6"),
                                       ("bool", {"md": "LM1903060101_a.md", "line_pages": [True, 10, 12, 12, 13, 13]}, "1 이상의 정수"),     # JSON true 는 int 의 하위형 — 따로 막는다
                                       ("str", {"md": "LM1903060101_a.md", "line_pages": ["10", 10, 12, 12, 13, 13]}, "1 이상의 정수")):
                (bad / "LM1903060101.pages.json").write_text(json.dumps(payload), encoding="utf-8")
                with self.assertRaisesRegex(ValueError, msg, msg=name):
                    SKR.load_page_maps(bad, [a])

    def test_reseg_agreement_skips_incomplete_rows_and_caps_disagreements(self):
        with tempfile.TemporaryDirectory() as td:
            ncs, maps, docs = _real_page_fixture(Path(td))
            docs = [d for d in docs if "LM1903060499" not in d.relative_path]
            rules, sources, candidates = self._rules()
            page_maps, _ = SKR.load_page_maps(maps, docs)
            graded = assign_match_grades(aggregate_matches(sources, docs, rules, candidates, page_maps=page_maps), {}, page_maps=page_maps)
            csv_path = Path(td) / "reseg.csv"
            csv_path.write_text("영역,교재,페이지,등급,출처\n반도체개발,LM1903060101_a,,1,text\n반도체개발,LM1903060101_a,10,,text\n반도체개발,없는교재,10,1,text\n"
                                "반도체개발,LM1903060101_a,10,3,text\n반도체개발,LM1903060101_a,12,1,text\n반도체재료,LM1903060408_b,1,2,text\n", encoding="utf-8")
            full = SKR.reseg_agreement(graded, csv_path)
            self.assertEqual((3, 0), (full["pages"], full["agree"]))                                     # 쪽·등급이 빈 행과 코드 없는 행은 무시; 공유 세 쪽 모두 불일치
            self.assertEqual([("LM1903060101", 10), ("LM1903060101", 12), ("LM1903060408", 1)], [(d["book"], d["page"]) for d in full["disagree"]])
            capped = SKR.reseg_agreement(graded, csv_path, max_disagree=1)
            self.assertEqual((3, 0, 1), (capped["pages"], capped["agree"], len(capped["disagree"])))     # 예시는 잘라도 집계는 전체

    def test_apply_page_maps_leaves_textbook_and_out_of_range_lines_alone(self):
        rules = [ExpressionRule("안전", "안전", "exact", "기존 키워드")]
        ncs = _doc("NCS", "반도체개발/LM1903060101_a.md", "<!-- page: 1 -->\n안전\n안전\n")
        school = _doc("교과서", "반도체개발/LM1903060101_a.md", "<!-- page: 1 -->\n안전\n안전\n")       # 같은 경로 키라도 교과서는 대응 대상이 아니다
        matches = [r for d in (ncs, school) for r in scan_document(d, rules)]
        out = SKR.apply_page_maps(matches, {"반도체개발/LM1903060101_a.md": (10, 12)})                 # 줄 3 은 대응 범위 밖 (검증을 거치지 않은 짧은 대응)
        self.assertEqual({("NCS", 2): 12, ("NCS", 3): 1, ("교과서", 2): 1, ("교과서", 3): 1}, {(r.corpus, r.line): r.page for r in out})
        self.assertEqual(matches, SKR.apply_page_maps(matches, None)); self.assertIsNot(matches, SKR.apply_page_maps(matches, {}))   # 대응 없음 → 같은 내용의 새 리스트

    def test_real_page_grade_falls_back_to_grade1_when_the_page_has_no_text(self):
        """방어 분기: 대응 교재의 레코드가 본문 없는 쪽에 놓이면(정상 대응에서는 불가능) 등급1·unpaged-fallback 로 기록한다 — 미배정을 두지 않는다."""
        with tempfile.TemporaryDirectory() as td:
            ncs, maps, docs = _real_page_fixture(Path(td))
            docs = [d for d in docs if "LM1903060101" in d.relative_path]
            rules, sources, candidates = self._rules()
            page_maps, _ = SKR.load_page_maps(maps, docs)
            result = aggregate_matches(sources, docs, rules, candidates, page_maps=page_maps)
            stray = replace(result, matches=tuple(replace(r, page=999) if r.line == 2 else r for r in result.matches))
            graded = assign_match_grades(stray, {}, page_maps=page_maps)
            by_line = {r.line: (r.page, r.grade, r.grade_source, r.grade_reason) for r in graded.matches if r.decision == "included"}
            self.assertEqual((999, 1, "unpaged-fallback", "실제 쪽 본문 없음으로 보수적 등급 1 배정"), by_line[2])
            self.assertEqual((12, 3, "real-page"), by_line[4][:3])                                       # 나머지는 실제 쪽 판정 그대로
            self.assertEqual(0, sum(1 for r in graded.matches if r.grade is None))

    def test_pdf_pages_from_previous_basis_filters_unusable_entries(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "reseg_summary.json"
            path.write_text(json.dumps({"per_book": {
                "LM1903060101_a": {"pdf_pages": 40}, "lm1903060102_b": {"pdf_pages": 7},          # 소문자 코드도 대문자 키로
                "LM1903060103_c": {}, "LM1903060104_d": {"pdf_pages": None}, "LM1903060105_e": {"pdf_pages": "12"},
                "LM1903060106_f": {"pdf_pages": 0}, "LM1903060107_g": None, "코드없음": {"pdf_pages": 9}}}), encoding="utf-8")
            self.assertEqual({"LM1903060101": 40, "LM1903060102": 7}, SKR.pdf_pages_from_previous_basis(path))
            path.write_text(json.dumps({"pages": 2189}), encoding="utf-8")
            self.assertEqual({}, SKR.pdf_pages_from_previous_basis(path))                              # per_book 없음

    def test_group_pages_textbook_and_markerless_ncs_ignore_pdf_pages(self):
        rules = [ExpressionRule("안전", "안전", "exact", "기존 키워드")]; sources = [KeywordSource("안전", 1, True)]
        docs = [_doc("NCS", "반도체개발/LM1903060101_a.md", "<!-- page: 1 -->\n안전\n<!-- page: 3 -->\n"),
                _doc("NCS", "반도체개발/LM1903060102_b.md", "안전\n"),                                  # 표식 없음 — 쪽수 0, PDF 쪽수를 묻지 않는다
                _doc("교과서", "school-1.md", "<!-- page: 1 -->\n안전\n<!-- page: 5 -->\n")]
        graded = assign_match_grades(aggregate_matches(sources, docs, rules, []), {})
        payload = SKR.dashboard_payload(graded, pdf_pages={"LM1903060101": 40})
        self.assertEqual({"반도체개발": 40}, {g["name"]: g["pages"] for g in payload["corpora"]["NCS"]["groups"]})
        self.assertEqual([5], [g["pages"] for g in payload["corpora"]["교과서"]["groups"]])           # 교과서는 pdf_pages 가 있어도 표식 최댓값
        self.assertEqual({"NCS": "marker", "교과서": "marker"}, payload["meta"]["page_basis"])           # real-page 판정이 없으면 marker

    def test_run_census_with_maps_but_no_reseg_csv_records_no_agreement(self):
        """reseg_csv 없이 돌린 대응 실행: 결속은 None 으로 기록되고(manifest·metrics·result.run) 정본 EXPECTED 의 reseg_agreement 와 어긋난다; 대응 없는 manifest 는 page_maps None."""
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            kw = census_fixture(root, ncs_body="<!-- page: 1 -->\n안전 안전 안전\n<!-- page: 2 -->\n안전 위험\n")
            maps = root / "maps"; maps.mkdir()
            (maps / "LM1903060101.pages.json").write_text(json.dumps({"md": "LM1903060101_안전.md", "line_pages": [10, 10, 12, 12, 12]}), encoding="utf-8")
            reseg = root / "reseg_summary.json"
            reseg.write_text(json.dumps({"pages": 2189, "page_g": {"1": 1519, "2": 525, "3": 145}, "books": 86, "cases_pages": 13, "unresolved": {"pages": 51},
                                         "per_book": {"LM1903060101_안전": {"pdf_pages": 40}}, "meta": {"expected": True, "run_at": "2026-09-07T08:48:51+00:00"}}), encoding="utf-8")
            with mock.patch.object(SKR, "REAL_PAGE_MARKER_BOOKS", tuple(f"LM19030602{i:02d}" for i in range(1, 86))):
                result = run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}}, page_maps_dir=maps, previous_basis=reseg,
                                    summary_out=root / "s.json", git={"commit": "x", "dirty": False}, argv=["x"])
            run = json.loads((root / "s.json").read_text(encoding="utf-8"))["meta"]["run"]
            self.assertIsNone(run["reseg_agreement"]); self.assertEqual(1, run["page_maps"]["files"])
            self.assertEqual((1, None), (result.run["page_maps"]["files"], result.run["reseg_agreement"]))
            metrics = SKR.summary_metrics(result, SKR.artifact_manifest(result), page_maps=SKR.PageMapsInfo(**result.run["page_maps"]))
            self.assertEqual([f"reseg_agreement.pages: None != {EXPECTED['reseg_agreement']['pages']}", f"reseg_agreement.agree: None != {EXPECTED['reseg_agreement']['agree']}"],
                             check_expected(metrics, {"reseg_agreement": EXPECTED["reseg_agreement"]}))
            legacy = SKR.run_manifest(graded_result(), argv=["x"], force=False, expected_mismatch=[], git={"commit": "x", "dirty": False})
            self.assertEqual((None, None, list(SKR.REAL_PAGE_MARKER_BOOKS)), (legacy["page_maps"], legacy["reseg_agreement"], legacy["real_page_marker_books"]))

    def test_main_prints_disagreements_and_metrics_from_result_run(self):
        run = {"page_maps": {"dir": "data/markdown/ncs_paged", "files": 84, "sha256": "b" * 64},
               "reseg_agreement": {"pages": 3, "agree": 2, "disagree": [{"book": "LM1903060101", "page": 7, "ours": 1, "reseg": 2}]}}
        argv = ["semantic_keyword_recount.py", "--source-workbook", "s.xlsx", "--ncs-root", "n", "--school-root", "t", "--xlsx-out", "o.xlsx", "--report-out", "r.md"]
        import io, contextlib
        out = io.StringIO()
        with mock.patch.object(SKR, "run_census", lambda *a, **k: replace(graded_result(), run=run)), mock.patch.object(SKR.sys, "argv", argv), contextlib.redirect_stdout(out):
            SKR.main()
        text = out.getvalue()
        self.assertIn("이전 기준과 등급이 다른 쪽: [{'book': 'LM1903060101', 'page': 7, 'ours': 1, 'reseg': 2}]", text)
        metrics = json.loads(text.split("측정값 (EXPECTED 고정용):", 1)[1])
        self.assertEqual(("b" * 64, {"pages": 3, "agree": 2}), (metrics["page_maps_sha256"], metrics["reseg_agreement"]))     # result.run 에서 PageMapsInfo 를 되살린다
        quiet = io.StringIO()
        with mock.patch.object(SKR, "run_census", lambda *a, **k: replace(graded_result(), run={"page_maps": None, "reseg_agreement": {"pages": 3, "agree": 3, "disagree": []}})), \
                mock.patch.object(SKR.sys, "argv", argv), contextlib.redirect_stdout(quiet):
            SKR.main()
        self.assertNotIn("이전 기준과 등급이 다른 쪽", quiet.getvalue()); self.assertIn('"page_maps_sha256": null', quiet.getvalue())   # 불일치 없으면 줄 자체가 없다

    def test_workbook_readme_and_run_sheet_state_the_page_basis(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "o.xlsx"
            run = {"git_commit": "abc1234", "page_maps": {"dir": "data/markdown/ncs_paged", "files": 84, "sha256": "c" * 64},
                   "reseg_agreement": {"pages": 3, "agree": 3}, "real_page_marker_books": list(SKR.REAL_PAGE_MARKER_BOOKS)}
            write_workbook(graded_result(), out, run=run)
            book = load_workbook(out, read_only=True)
            readme = {row[0]: row[1] for row in book["README"].iter_rows(values_only=True)}
            self.assertIn("real-page", readme["등급 계보"]); self.assertIn("워크북 라벨(목차 블록) 상속 없음", readme["등급 계보"])
            self.assertIn("page_maps", readme["페이지 기준"]); self.assertIn("per_book.pdf_pages", readme["페이지 기준"])
            info = {row[1]: row[2] for row in book["입력정보"].iter_rows(values_only=True) if row[0] == "실행 정보"}
            self.assertEqual(run["page_maps"], json.loads(info["page_maps"])); self.assertEqual({"pages": 3, "agree": 3}, json.loads(info["reseg_agreement"]))
            self.assertEqual(list(SKR.REAL_PAGE_MARKER_BOOKS), json.loads(info["real_page_marker_books"]))


class ImpactScriptTests(unittest.TestCase):
    """occurrence_real_pages_impact.py — 블록 기준(A) vs 실제 쪽 기준(B): 총계 불변, 이동 행렬, 계보 검증, 산출물에 본문·절대 경로 없음."""

    def _setup(self, td):
        import occurrence_real_pages_impact as IMP
        ncs, maps, docs = _real_page_fixture(Path(td))
        docs = [d for d in docs if "LM1903060499" not in d.relative_path]
        page_maps, _ = SKR.load_page_maps(maps, docs)
        a = next(d for d in docs if "LM1903060101" in d.relative_path)
        existing = {SKR.grade_lookup_key("NCS", a.relative_path, 2): GradeAssignment(3, "구체적 대책", "워크북 라벨(블록)", "existing")}
        return IMP, docs, page_maps, existing

    def test_impact_reports_transition_and_keeps_totals(self):
        with tempfile.TemporaryDirectory() as td:
            IMP, docs, page_maps, existing = self._setup(td)
            out = IMP.compute_impact(docs, ["안전"], existing, page_maps, block_reference=None)
            self.assertEqual(out["totals"]["NCS"], sum(out["grades"]["block"]["NCS"].values())); self.assertEqual(sum(out["grades"]["block"]["NCS"].values()), sum(out["grades"]["real"]["NCS"].values()))
            # LM…0101 블록 2(줄 4 의 5건 + 줄 5 의 1건)는 A 에서 워크북 라벨 3 을 상속, B 에서 쪽 12(5건) → 3, 쪽 13(1건) → 1; 나머지 6건은 양쪽 다 등급1
            self.assertEqual({"1": 6, "2": 0, "3": 6}, out["grades"]["block"]["NCS"]); self.assertEqual({"1": 7, "2": 0, "3": 5}, out["grades"]["real"]["NCS"])
            self.assertEqual((11, 1, 1), (out["transition"]["unchanged"], out["transition"]["moved"], out["transition"]["matrix"]["3->1"]))
            self.assertEqual({"existing": {"occurrences": 6, "unchanged": 5, "unchanged_pct": 83.3}, "new": {"occurrences": 6, "unchanged": 6, "unchanged_pct": 100.0}}, out["by_source"])
            self.assertEqual({"real-marker-nomap", "toc-block"}, set(out["by_book_kind"]))            # LM…0408 은 목록 교재, LM…0101 은 블록 2 → 실제 쪽 3
            self.assertEqual((4, 5), (out["pages"]["block_pages"], out["pages"]["real_pages"]))
            self.assertEqual({"1": 3, "2-3": 6, "4-9": 0, "10+": 0}, out["pages"]["block_width_of_occurrences"])

    def test_impact_refuses_when_block_basis_differs_from_reference(self):
        with tempfile.TemporaryDirectory() as td:
            IMP, docs, page_maps, existing = self._setup(td)
            with self.assertRaisesRegex(ValueError, "계보"):
                IMP.compute_impact(docs, ["안전"], existing, page_maps, block_reference=IMP.BLOCK_BASIS_V2)
            IMP.compute_impact(docs, ["안전"], existing, page_maps, block_reference={"NCS": {"1": 6, "2": 0, "3": 6}, "교과서": {"1": 0, "2": 0, "3": 0}})

    def test_impact_main_writes_json_without_absolute_paths(self):
        import io, contextlib
        import occurrence_real_pages_impact as IMP
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            kw = census_fixture(root, ncs_body="<!-- page: 1 -->\n안전 안전 안전\n<!-- page: 2 -->\n안전 위험 보호구 착용 안전 안전 안전 방지 예방 환기 차단 대피 안전모\n안전 위험\n")
            maps = root / "maps"; maps.mkdir()
            (maps / "LM1903060101.pages.json").write_text(json.dumps({"md": "LM1903060101_안전.md", "line_pages": [10, 10, 12, 12, 13, 13]}), encoding="utf-8")
            out = root / "impact.json"
            argv = ["occurrence_real_pages_impact.py", "--source-workbook", str(kw["source_workbook"]), "--ncs-root", str(kw["ncs_root"]), "--school-root", str(kw["school_root"]),
                    "--school-grade-workbook", str(kw["school_grade_workbook"]), "--page-maps", str(maps), "--out", str(out)]
            buf = io.StringIO()
            with mock.patch.object(IMP.sys, "argv", argv), mock.patch.object(IMP, "BLOCK_BASIS_V2", None), contextlib.redirect_stdout(buf):
                IMP.main()
            text = out.read_text(encoding="utf-8"); data = json.loads(text)
            self.assertNotIn(str(root), text); self.assertNotIn("/Users/", text)
            self.assertEqual(1, data["meta"]["inputs"]["page_maps"]["files"]); self.assertEqual(data["totals"]["NCS"], sum(data["grades"]["real"]["NCS"].values()))
            self.assertIn("이동", buf.getvalue())

    # ---- 출고 전 커버리지 감사 (2026-09-15): book_kind 의 세 값, 짝짓기 거부, 폭 구간 4-9·10+, 교과서 레코드, 정렬 자기 검증 병기
    def test_book_kind_real_marker_and_empty_map(self):
        import occurrence_real_pages_impact as IMP
        dense = _doc("NCS", "반도체개발/LM1903060101_a.md", "<!-- page: 1 -->\n안전\n<!-- page: 2 -->\n안전\n<!-- page: 3 -->\n안전\n<!-- page: 4 -->\n안전\n<!-- page: 5 -->\n안전\n")
        self.assertEqual("real-marker", IMP.book_kind(dense, {dense.relative_path: (1, 1, 2, 2, 3, 3, 4, 4, 5, 6)}))            # 블록 5 / 실제 쪽 6 = 0.83 ≥ 0.8
        self.assertEqual("toc-block", IMP.book_kind(dense, {dense.relative_path: (1, 2, 3, 4, 5, 6, 7, 8, 9, 10)}))            # 5 / 10 = 0.5
        self.assertEqual("toc-block", IMP.book_kind(dense, {dense.relative_path: ()}))                                          # 실제 쪽 0 — 나눗셈 없이 toc-block
        self.assertEqual("real-marker-nomap", IMP.book_kind(dense, {}))
        self.assertEqual(0.8, IMP.REAL_MARKER_RATIO)

    def test_impact_refuses_when_the_map_changes_the_match_set(self):
        with tempfile.TemporaryDirectory() as td:
            IMP, docs, page_maps, existing = self._setup(td)
            original = SKR.apply_page_maps
            with mock.patch.object(SKR, "apply_page_maps", lambda matches, maps: original(matches, maps)[:-1] if maps else original(matches, maps)):     # 실제 쪽 집계만 한 건 잃는다
                with self.assertRaisesRegex(ValueError, "매칭 순서·집합"):
                    IMP.compute_impact(docs, ["안전"], existing, page_maps, block_reference=None)
            with mock.patch.object(SKR, "apply_page_maps", lambda matches, maps: list(reversed(original(matches, maps))) if maps else original(matches, maps)):   # 수는 같고 순서만 다르다
                with self.assertRaisesRegex(ValueError, "매칭 순서·집합"):
                    IMP.compute_impact(docs, ["안전"], existing, page_maps, block_reference=None)

    def test_impact_block_width_buckets_grade3_share_and_textbook_records(self):
        """블록 폭 4-9·10+ 구간, 등급3 비율, 교과서 레코드는 이동 행렬 밖(불변). 블록 2 (안전 ×12) 는 블록 판정 2, 실제 쪽 12개로 흩어지면 각 1."""
        import occurrence_real_pages_impact as IMP
        wide = "<!-- page: 1 -->\n" + "안전\n" * 5 + "<!-- page: 2 -->\n" + "안전\n" * 12
        ncs = _doc("NCS", "반도체개발/LM1903060101_a.md", wide)
        school = _doc("교과서", "school-1.md", "<!-- page: 1 -->\n안전\n")
        line_pages = tuple([10] + [10, 11, 12, 13, 14] + [30] + list(range(30, 42)))                 # 표식 줄은 뒤따르는 쪽; 블록 1 → 실제 쪽 5개(4-9), 블록 2 → 12개(10+)
        out = IMP.compute_impact([ncs, school], ["안전"], {}, {ncs.relative_path: line_pages}, block_reference=None)
        self.assertEqual({"1": 0, "2-3": 0, "4-9": 5, "10+": 12}, out["pages"]["block_width_of_occurrences"])
        self.assertEqual({"NCS": 17, "교과서": 1}, out["totals"])
        self.assertEqual(({"1": 5, "2": 12, "3": 0}, {"1": 17, "2": 0, "3": 0}), (out["grades"]["block"]["NCS"], out["grades"]["real"]["NCS"]))
        self.assertEqual(({"1": 1, "2": 0, "3": 0}, {"1": 1, "2": 0, "3": 0}), (out["grades"]["block"]["교과서"], out["grades"]["real"]["교과서"]))   # 교과서는 대응 밖 — 불변
        self.assertEqual({"NCS": {"block": 0.0, "real": 0.0}, "교과서": {"block": 0.0, "real": 0.0}}, out["grade3_share"])
        self.assertEqual((5, 12, 12), (out["transition"]["unchanged"], out["transition"]["moved"], out["transition"]["matrix"]["2->1"]))   # 교과서 1건은 NCS 행렬에 없다
        self.assertEqual({"new": {"occurrences": 17, "unchanged": 5, "unchanged_pct": 29.4}}, out["by_source"])
        self.assertEqual((2, 17), (out["pages"]["block_pages"], out["pages"]["real_pages"]))
        self.assertEqual({"toc-block": {"books": 1, "occurrences": 17, "block": {"1": 5, "2": 12, "3": 0}, "real": {"1": 17, "2": 0, "3": 0}}}, out["by_book_kind"])
        self.assertEqual({"반도체개발": {"block": {"1": 5, "2": 12, "3": 0}, "real": {"1": 17, "2": 0, "3": 0}}}, out["groups"])

    def test_impact_main_alignment_self_check_follows_previous_basis(self):
        import io, contextlib
        import occurrence_real_pages_impact as IMP
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            kw = census_fixture(root, ncs_body="<!-- page: 1 -->\n안전 안전 안전\n<!-- page: 2 -->\n안전 위험\n")
            maps = root / "maps"; maps.mkdir()
            (maps / "LM1903060101.pages.json").write_text(json.dumps({"md": "LM1903060101_안전.md", "line_pages": [10, 10, 12, 12, 12]}), encoding="utf-8")
            reseg = root / "reseg_summary.json"
            reseg.write_text(json.dumps({"alignment_check": {"books": 23, "overall": {"lines": 100, "exact": 0.9, "near": 0.95, "nogap_exact": 0.92}}, "hybrid_lines": 7}), encoding="utf-8")
            base = ["occurrence_real_pages_impact.py", "--source-workbook", str(kw["source_workbook"]), "--ncs-root", str(kw["ncs_root"]), "--school-root", str(kw["school_root"]),
                    "--school-grade-workbook", str(kw["school_grade_workbook"]), "--page-maps", str(maps)]
            for tag, previous in (("with", reseg), ("without", root / "missing.json")):
                out = root / f"{tag}.json"
                with mock.patch.object(IMP.sys, "argv", base + ["--out", str(out), "--previous-basis", str(previous)]), mock.patch.object(IMP, "BLOCK_BASIS_V2", None), \
                        contextlib.redirect_stdout(io.StringIO()):
                    IMP.main()
                meta = json.loads(out.read_text(encoding="utf-8"))["meta"]
                if tag == "with":
                    check = meta["alignment_self_check"]
                    self.assertEqual((23, 100, 0.9, 0.92, None, 7), (check["books"], check["lines"], check["exact"], check["nogap_exact"], check["all_lines"], check["hybrid_lines"]))   # 없는 키는 None 으로 병기
                    self.assertEqual("reseg_summary.json", Path(check["source"]).name); self.assertNotIn(str(root), json.dumps(meta))
                else:
                    self.assertNotIn("alignment_self_check", meta)                                                                # 이전 기준 파일이 없으면 병기하지 않는다


class LineConventionTests(unittest.TestCase):
    """갭 분석 Act-1 (2026-09-15). G-1: 매칭의 줄 번호는 splitlines() 인데 대응은 split("\\n") 기준 — \\x0c 같은 줄 구분자가 있으면 그 뒤가 한 줄씩 밀린다,
    대응을 매칭 규약으로 옮겨 싣는다. G-4·G-6: CLI 기본값, 제외 레코드·목록 교재·변형 사전에도 대응 적용, 결속 불일치 거부, 대응 적용 결정론, 대응 폴더 부재."""

    def test_page_maps_are_remapped_to_the_matching_line_convention(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); ncs = root / "ncs" / "반도체개발"; ncs.mkdir(parents=True)
            text = "<!-- page: 1 -->\n안전 x\x0cy\n안전\n<!-- page: 2 -->\n안전\n"           # split("\\n") 6줄(끝 빈 문자열 포함) · splitlines 6줄(\\x0c 로 하나 늘고 끝 빈 줄이 빠짐)
            (ncs / "LM1903060101_a.md").write_text(text, encoding="utf-8")
            maps = root / "maps"; maps.mkdir()
            (maps / "LM1903060101.pages.json").write_text(json.dumps({"md": "LM1903060101_a.md", "line_pages": [10, 10, 10, 20, 20, 20]}), encoding="utf-8")
            docs = SKR.load_documents(root / "ncs", "NCS")
            page_maps, _ = SKR.load_page_maps(maps, docs)
            self.assertEqual((10, 10, 10, 10, 20, 20), page_maps[docs[0].relative_path])                  # 매칭 규약(splitlines) 으로 옮겨진 대응
            rules = [ExpressionRule("안전", "안전", "exact", "기존 키워드")]
            result = aggregate_matches([KeywordSource("안전", 1, True)], docs, rules, [], page_maps=page_maps)
            self.assertEqual({2: 10, 4: 10, 6: 20}, {r.line: r.page for r in result.matches if r.decision == "included"})   # 줄 4 의 '안전' 은 쪽 10 (밀리면 20)
            graded = assign_match_grades(result, {}, page_maps=page_maps)
            self.assertTrue(all(r.grade_source == "real-page" for r in graded.matches))
            bad = maps / "bad"; bad.mkdir()
            (bad / "LM1903060101.pages.json").write_text(json.dumps({"md": "LM1903060101_a.md", "line_pages": [10, 10, 10, 20, 20]}), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "줄 수"):
                SKR.load_page_maps(bad, docs)                                                            # 길이 검사는 대응 규약(split("\\n")) 그대로

    def test_excluded_records_and_listed_books_with_maps_and_variant_runs(self):          # Act-1 G-6
        with tempfile.TemporaryDirectory() as td:
            ncs, maps, docs = _real_page_fixture(Path(td))
            docs = [d for d in docs if "LM1903060499" not in d.relative_path]
            b = next(d for d in docs if "LM1903060408" in d.relative_path)
            (maps / "LM1903060408.pages.json").write_text(json.dumps({"md": "LM1903060408_b.md", "line_pages": [30, 30, 31, 31, 31]}), encoding="utf-8")
            page_maps, info = SKR.load_page_maps(maps, docs)
            self.assertEqual((30, 30, 31, 31), page_maps[b.relative_path]); self.assertEqual(2, info.files)              # 목록 교재라도 대응이 있으면 대응을 쓴다
            rules = build_default_rules(["안전"], version="v1fix")                                                         # 변형 사전에도 대응은 그대로 적용
            result = aggregate_matches([KeywordSource("안전", 1, True)], docs, rules, SKR.default_candidate_decisions(version="v1fix"), page_maps=page_maps)
            excluded = [r for r in result.matches if r.decision != "included"]
            self.assertTrue(all(r.page in (10, 12, 13, 30, 31) for r in result.matches if r.corpus == "NCS"), "제외·보류 레코드도 실제 쪽을 받는다")
            self.assertTrue(any(r.page == 31 for r in result.matches if "LM1903060408" in r.relative_path))

    def test_run_census_refuses_agreement_mismatch_and_is_deterministic_with_maps(self):  # Act-1 G-6
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            kw = census_fixture(root, ncs_body="<!-- page: 1 -->\n안전 안전 안전\n<!-- page: 2 -->\n안전 위험 보호구 착용 안전 안전 안전 방지 예방 환기 차단 대피 안전모\n안전 위험\n")
            maps = root / "maps"; maps.mkdir()
            (maps / "LM1903060101.pages.json").write_text(json.dumps({"md": "LM1903060101_안전.md", "line_pages": [10, 10, 12, 12, 13, 13]}), encoding="utf-8")
            reseg = root / "reseg_summary.json"
            reseg.write_text(json.dumps({"pages": 2189, "page_g": {"1": 1519, "2": 525, "3": 145}, "books": 86, "cases_pages": 13, "unresolved": {"pages": 51},
                                         "per_book": {"LM1903060101_안전": {"pdf_pages": 40}}, "meta": {"expected": True, "run_at": "2026-09-07T08:48:51+00:00"}}), encoding="utf-8")
            csv_path = root / "reseg.csv"; csv_path.write_text("﻿영역,교재,페이지,등급,출처\n반도체개발,LM1903060101_안전,12,1,text\n", encoding="utf-8")   # 쪽 12 는 본문 판정 3 — 이전 기준이 1 이면 결속 실패
            with mock.patch.object(SKR, "REAL_PAGE_MARKER_BOOKS", tuple(f"LM19030602{i:02d}" for i in range(1, 86))):
                with self.assertRaises(SystemExit):
                    run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}, "reseg_agreement": {"pages": 1, "agree": 1}}, page_maps_dir=maps, reseg_csv=csv_path, previous_basis=reseg,
                               git={"commit": "x", "dirty": False}, argv=["x"])
                self.assertFalse(kw["xlsx_out"].exists())
                csv_path.write_text("﻿영역,교재,페이지,등급,출처\n반도체개발,LM1903060101_안전,12,3,text\n", encoding="utf-8")
                hashes = []
                for n in (1, 2):
                    run_census(**dict(kw, xlsx_out=root / f"o{n}.xlsx", report_out=root / f"o{n}.md"), expected={"documents": {"NCS": 86, "교과서": 9}, "reseg_agreement": {"pages": 1, "agree": 1}},
                               page_maps_dir=maps, reseg_csv=csv_path, previous_basis=reseg, summary_out=root / f"s{n}.json", git={"commit": "x", "dirty": False}, argv=["x"])
                    hashes.append(json.loads((root / f"s{n}.json").read_text(encoding="utf-8"))["meta"]["manifest"])
                self.assertEqual(hashes[0], hashes[1])                                                                     # 대응을 적용해도 같은 fixture → 같은 해시
                with self.assertRaisesRegex(FileNotFoundError, "대응 폴더"):
                    run_census(**dict(kw, xlsx_out=root / "o3.xlsx", report_out=root / "o3.md"), expected={"documents": {"NCS": 86, "교과서": 9}}, page_maps_dir=root / "none", previous_basis=reseg,
                               git={"commit": "x", "dirty": False}, argv=["x"])

    def test_main_defaults_to_the_paged_dir_and_reseg_csv(self):                                # Act-1 G-4
        seen = {}
        def fake_run_census(*args, **kwargs):
            seen.update(kwargs); return graded_result()
        argv = ["semantic_keyword_recount.py", "--source-workbook", "s.xlsx", "--ncs-root", "n", "--school-root", "t", "--xlsx-out", "o.xlsx", "--report-out", "r.md"]
        import io, contextlib
        with mock.patch.object(SKR, "run_census", fake_run_census), mock.patch.object(SKR.sys, "argv", argv), contextlib.redirect_stdout(io.StringIO()):
            SKR.main()
        self.assertEqual((SKR.DEFAULT_PAGE_MAPS_DIR, SKR.DEFAULT_RESEG_CSV), (seen["page_maps_dir"], seen["reseg_csv"]))

    # ---- 출고 전 커버리지 감사 (2026-09-15): 줄 규약의 남은 변형, 결속 불일치 때의 stderr 힌트
    def test_to_matching_lines_crlf_no_trailing_newline_and_unknown_separator(self):
        """\\r\\n 은 split("\\n") 조각 끝의 \\r 이 splitlines 에서 사라질 뿐 줄 수가 같고, 끝 개행이 없으면 빈 원소를 빼지 않는다; 빈 본문(splitlines 0줄 vs split 1줄)은 거부."""
        self.assertEqual((10, 20, 20), SKR._to_matching_lines("a\r\nb\r\nc", [10, 20, 20]))          # 끝 개행 없음: 3 원소 그대로
        self.assertEqual((10, 20), SKR._to_matching_lines("a\r\nb\r\n", [10, 20, 30]))              # 끝 개행: 마지막 빈 원소(30) 는 splitlines 에 없다
        self.assertEqual((10, 10, 20), SKR._to_matching_lines("a\rb\nc", [10, 20]))                  # 홀로 선 \r 도 splitlines 는 줄로 가른다
        with self.assertRaisesRegex(ValueError, "줄 규약"):
            SKR._to_matching_lines("", [1])

    def test_run_census_mismatch_prints_agreement_hint_with_examples(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            kw = census_fixture(root, ncs_body="<!-- page: 1 -->\n안전 안전 안전\n<!-- page: 2 -->\n안전 위험\n")
            maps = root / "maps"; maps.mkdir()
            (maps / "LM1903060101.pages.json").write_text(json.dumps({"md": "LM1903060101_안전.md", "line_pages": [10, 10, 12, 12, 12]}), encoding="utf-8")
            reseg = root / "reseg_summary.json"
            reseg.write_text(json.dumps({"pages": 2189, "page_g": {"1": 1519, "2": 525, "3": 145}, "books": 86, "cases_pages": 13, "unresolved": {"pages": 51},
                                         "per_book": {"LM1903060101_안전": {"pdf_pages": 40}}, "meta": {"expected": True, "run_at": "2026-09-07T08:48:51+00:00"}}), encoding="utf-8")
            csv_path = root / "reseg.csv"; csv_path.write_text("영역,교재,페이지,등급\n반도체개발,LM1903060101_안전,10,3\n", encoding="utf-8")   # 쪽 10 은 본문 판정 1
            import io, contextlib
            listed = tuple(f"LM19030602{i:02d}" for i in range(1, 86))
            err = io.StringIO()
            with mock.patch.object(SKR, "REAL_PAGE_MARKER_BOOKS", listed), contextlib.redirect_stderr(err), self.assertRaises(SystemExit):
                run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}, "reseg_agreement": {"pages": 1, "agree": 1}}, page_maps_dir=maps, reseg_csv=csv_path,
                           previous_basis=reseg, git={"commit": "x", "dirty": False}, argv=["x"])
            self.assertIn("reseg_agreement.agree: 0 != 1", err.getvalue())
            self.assertIn("줄→쪽 대응(page_maps)이 바뀌었거나 regrade 기준선이 바뀐 것", err.getvalue())
            self.assertIn("'book': 'LM1903060101', 'page': 10, 'ours': 1, 'reseg': 3", err.getvalue())                          # 어느 쪽인지 바로 보이게 예시를 붙인다
            self.assertFalse(kw["xlsx_out"].exists())
            err2 = io.StringIO()                                                                                                 # 결속은 맞고 다른 항목이 어긋나면 힌트가 없다
            with mock.patch.object(SKR, "REAL_PAGE_MARKER_BOOKS", listed), contextlib.redirect_stderr(err2), self.assertRaises(SystemExit):
                run_census(**kw, expected={"documents": {"NCS": 86, "교과서": 9}, "reseg_agreement": {"pages": 1, "agree": 0}, "totals": {"NCS": 999999}}, page_maps_dir=maps,
                           reseg_csv=csv_path, previous_basis=reseg, git={"commit": "x", "dirty": False}, argv=["x"])
            self.assertIn("totals.NCS", err2.getvalue()); self.assertNotIn("regrade 기준선", err2.getvalue())
