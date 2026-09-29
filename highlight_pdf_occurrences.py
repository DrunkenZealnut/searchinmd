#!/usr/bin/env python3.13
"""정본 재검산 xlsx 의 출현 키워드를 원본 PDF 사본에 등급별 형광펜 주석으로 표시한다.

    python3.13 highlight_pdf_occurrences.py                       # 기본 경로 (data/…), 산출물 data/highlighted/
    python3.13 highlight_pdf_occurrences.py --only LM1903060329   # 한 권만 (부분 실행 — 로그는 .partial)

두 경로: NCS PDF 는 텍스트 층이 있어 PyMuPDF rawdict 글자 좌표를, 교과서 PDF 는 스캔본이라 Apple Vision OCR
(outputs/vision_ocr_chars.swift, 실행 시 swiftc 로 빌드) 글자 좌표를 쓴다. 쪽 안 위치는 마크다운 줄의 앞뒤 문맥 창으로
찾고, 줄 안 매칭 오프셋은 정본 규칙 엔진(semantic_keyword_recount)의 포함 판정을 재현해 얻는다.
설계: docs/02-design/features/pdf-keyword-highlight.design.md
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import fitz

import semantic_keyword_recount as skr

HERE = Path(__file__).resolve().parent

STRIP = re.compile(r'[\s#*_>|\[\]()!`\-–·•○●◆◇■□▶▷※,.:;]+')      # resegment._STRIP 과 동일 (시험이 고정)
OUT_SUFFIX = "_키워드표시"
GRADE_HEX = {1: "#FFF176", 2: "#FFB74D", 3: "#81C784"}              # 연구 책임자 D2 (2026-09-26): 노랑·주황·초록
GRADE_COLORS = {g: tuple(int(h[i:i + 2], 16) / 255 for i in (1, 3, 5)) for g, h in GRADE_HEX.items()}

_TEXTBOOK_KEY_RE = re.compile(r"[^0-9A-Za-z가-힣]+")
_TIMESTAMP_RE = re.compile(r"^\d{8}_\d{6}_")


def nfc(s: str | None) -> str:
    return unicodedata.normalize("NFC", s or "")


def norm_chars(s: str | None) -> str:
    """resegment.norm_text + casefold — 비교 전용 (NFC, 공백·마크다운 기호·문장부호 제거)."""
    return STRIP.sub("", nfc(s)).casefold()


# ----------------------------------------------------------------------------- 글자 스트림
@dataclass
class CharStream:
    """쪽의 글자를 추출 순서대로 정규화한 문자열과, 정규화 글자마다의 PDF 좌표·줄 id."""

    text: str = ""
    boxes: list[fitz.Rect] = field(default_factory=list)
    line_ids: list[int] = field(default_factory=list)

    @classmethod
    def from_chars(cls, chars) -> "CharStream":
        text, boxes, line_ids = [], [], []
        for ch, rect, line_id in chars:
            for out in norm_chars(ch):                       # casefold 는 늘어날 수 있다(ß → ss) — 같은 box 를 나눠 갖는다
                text.append(out)
                boxes.append(fitz.Rect(rect))
                line_ids.append(line_id)
        return cls("".join(text), boxes, line_ids)

    @classmethod
    def from_rawdict(cls, page: fitz.Page) -> "CharStream":
        chars = []
        line_id = 0
        for block in page.get_text("rawdict")["blocks"]:
            if block.get("type") != 0:
                continue
            for line in block["lines"]:
                for span in line["spans"]:
                    for ch in span["chars"]:
                        chars.append((ch["c"], fitz.Rect(ch["bbox"]), line_id))
                line_id += 1
        return cls.from_chars(chars)

    @classmethod
    def from_ocr(cls, ocr_page: dict, page: fitz.Page) -> "CharStream":
        """Vision OCR JSON(픽셀, 좌상단 원점) → 쪽 좌표. 글자 box 가 없으면 줄 box 를 쓴다. 회전 쪽은 되돌린다."""
        sx = page.rect.width / ocr_page["width"]
        sy = page.rect.height / ocr_page["height"]
        derot = page.derotation_matrix

        def to_rect(box) -> fitz.Rect:
            r = fitz.Rect(box[0] * sx + page.rect.x0, box[1] * sy + page.rect.y0, box[2] * sx + page.rect.x0, box[3] * sy + page.rect.y0)
            return r * derot

        chars = []
        for line_id, line in enumerate(ocr_page.get("lines", [])):
            line_rect = to_rect(line["box"])
            for ch in line["chars"]:
                chars.append((ch["c"], to_rect(ch["box"]) if ch.get("box") else line_rect, line_id))
        return cls.from_chars(chars)


# ----------------------------------------------------------------------------- PDF 대응
def index_ncs_pdfs(root: Path) -> dict[str, Path]:
    """파일명의 LM 코드 전부 → 경로. 한 파일이 두 코드를 지니기도 한다(LM…314, LM…315 합본); 한 코드가 두 파일이면 거부."""
    index: dict[str, Path] = {}
    for f in sorted(Path(root).rglob("*.pdf")):
        for code in skr._NCS_CODE_RE.findall(nfc(f.name)):
            code = code.upper()
            if code in index and index[code] != f:
                raise LookupError(f"LM 코드 {code} 가 두 PDF 에 있습니다: {index[code]} / {f}")
            index[code] = f
    return index


def textbook_key(relpath: str) -> str:
    """xlsx 파일 경로의 교과서 폴더명 → 이름 키 (타임스탬프·기호 제거)."""
    folder = Path(nfc(relpath)).parent.name
    return _TEXTBOOK_KEY_RE.sub("", _TIMESTAMP_RE.sub("", folder))


def index_textbook_pdfs(root: Path) -> dict[str, Path]:
    return {_TEXTBOOK_KEY_RE.sub("", nfc(f.stem)): f for f in sorted(Path(root).glob("*.pdf"))}


def pdf_for(corpus: str, relpath: str, ncs_index: dict[str, Path], textbook_index: dict[str, Path]) -> Path:
    if corpus == "NCS":
        m = skr._NCS_CODE_RE.search(nfc(relpath))
        if not m or m.group(0).upper() not in ncs_index:
            raise LookupError(f"NCS PDF 를 찾을 수 없습니다: {relpath}")
        return ncs_index[m.group(0).upper()]
    key = textbook_key(relpath)
    if key not in textbook_index:
        raise LookupError(f"교과서 PDF 를 찾을 수 없습니다: {relpath} (키 {key})")
    return textbook_index[key]


def output_path(corpus: str, pdf: Path, pdf_root: Path, out_root: Path) -> Path:
    name = pdf.stem + OUT_SUFFIX + ".pdf"
    if corpus == "NCS":
        return Path(out_root) / "ncs" / pdf.relative_to(pdf_root).parent / name
    return Path(out_root) / "school-text" / name


# ----------------------------------------------------------------------------- 줄 안 위치 재현 (정본 엔진의 포함 판정)
_IMAGE_ONLY_RE = re.compile(r"\s*!\[[^]]*]\([^)]*\)\s*")            # scan_document 가 건너뛰는 줄


@dataclass(frozen=True)
class Located:
    keyword: str
    expression: str
    tier: str
    matched_text: str
    start: int
    end: int


@dataclass
class Row:
    """xlsx 매칭상세 한 행. text/start/end/locator 는 assign_offsets 가 채운다."""

    corpus: str
    relpath: str
    keyword: str
    expression: str
    tier: str
    matched: str
    line: int
    page: int
    context: str
    grade: int
    grade_label: str
    order: int
    text: str | None = None
    start: int | None = None
    end: int | None = None
    locator: str | None = None


def compile_rules(rules: list[skr.ExpressionRule]) -> dict[str, list]:
    by_keyword: dict[str, list] = {}
    for index, rule in enumerate(rules):
        by_keyword.setdefault(rule.keyword, []).append((index, rule, skr._compile_rule(rule), skr._compile_companions(rule)))
    return by_keyword


def locate_included(line: str, window: str, compiled: dict[str, list]) -> list[Located]:
    """scan_document 의 포함 판정을 그대로 재현해 줄 안 오프셋을 돌려준다 — 키워드는 규칙 순서, 안에서는 시작 위치 순.

    제외·보류·동반어 규칙과 겹침 커서(같은 키워드 안에서 앞 매칭과 겹치면 버림)가 정본과 같아야 xlsx 행 순서열과 1:1 이 된다.
    """
    out: list[Located] = []
    for keyword, compiled_rules in compiled.items():
        included = []
        for rule_index, rule, pattern, companions in compiled_rules:
            matches_on_line = list(pattern.finditer(line))
            if not matches_on_line:
                continue
            exclusion_spans = [(m.start(), m.end()) for ex in rule.exclude_patterns for m in re.finditer(ex, line, re.IGNORECASE)]
            held_spans = [(m.start(), m.end()) for held in rule.held_patterns for m in re.finditer(held, line, re.IGNORECASE)]
            companion = companions is None or companions.search(window) is not None
            for match in matches_on_line:
                span = (match.start(), match.end())
                if any(skr._overlaps(span, held) for held in held_spans):
                    continue
                if any(skr._overlaps(span, ex) for ex in exclusion_spans):
                    continue
                if not companion:
                    continue
                included.append((match.start(), match.end(), rule_index, rule, match.group(0)))
        included.sort(key=lambda item: (item[0], -(item[1] - item[0]), item[2]))
        cursor = -1
        for start, end, _, rule, matched_text in included:
            if start < cursor:
                continue
            out.append(Located(keyword, rule.expression, rule.tier, matched_text, start, end))
            cursor = end
    return out


def line_offsets(document: skr.Document, compiled: dict[str, list], wanted_lines: set[int] | None = None) -> dict[int, list[Located]]:
    """문서의 (원하는) 줄마다 포함 매칭 오프셋 — 정본과 같은 페이지 블록·동반어 창(같은 줄 ±1, 블록 안)."""
    out: dict[int, list[Located]] = {}
    for block in skr.split_pages(document):
        for line_offset, line in enumerate(block.lines):
            line_number = block.start_line + line_offset
            if wanted_lines is not None and line_number not in wanted_lines:
                continue
            if not line or _IMAGE_ONLY_RE.fullmatch(line):
                out[line_number] = []
                continue
            window = "\n".join(block.lines[max(0, line_offset - 1):line_offset + 2])
            out[line_number] = locate_included(line, window, compiled)
    return out


def assign_offsets(rows: list[Row], offsets: dict[int, list[Located]], raw_lines: dict[int, str]) -> dict[str, int]:
    """행마다 text/start/end 를 채운다. 줄의 (키워드, 실제 매칭) 순서열이 엔진과 같으면 engine, 아니면 naive(k번째 인스턴스)."""
    counts = {"engine": 0, "naive": 0}
    by_line: dict[int, list[Row]] = {}
    for row in rows:
        by_line.setdefault(row.line, []).append(row)
    for line_number, group in by_line.items():
        group.sort(key=lambda r: r.order)
        located = offsets.get(line_number, [])
        raw = raw_lines.get(line_number)
        # xlsx 는 같은 줄의 행을 (키워드, 실제 매칭) 정렬로 쓴다 — 다중집합이 같으면 같은 (키워드, 실제 매칭)끼리 문서 순서로 대응
        if raw is not None and sorted((r.keyword, r.matched) for r in group) == sorted((l.keyword, l.matched_text) for l in located):
            queues: dict[tuple[str, str], list[Located]] = {}
            for loc in located:
                queues.setdefault((loc.keyword, loc.matched_text), []).append(loc)
            for row in group:
                loc = queues[(row.keyword, row.matched)].pop(0)
                row.text, row.start, row.end, row.locator = raw, loc.start, loc.end, "engine"
            counts["engine"] += 1
            continue
        counts["naive"] += 1
        seen: dict[str, int] = {}
        for row in group:
            k = seen.get(row.matched.casefold(), 0)
            seen[row.matched.casefold()] = k + 1
            hits = list(re.finditer(re.escape(row.matched), row.context, re.IGNORECASE))
            row.text, row.locator = row.context, "naive"
            if k < len(hits):
                row.start, row.end = hits[k].start(), hits[k].end()
            else:
                row.start = row.end = None
    return counts


# ----------------------------------------------------------------------------- 문맥 창 검색
CONTEXT_WINDOWS = (10, 6, 3, 0)
NEIGHBOR_MIN_CONTEXT = 6    # 정본 쪽 자리가 소진됐을 때 이웃 쪽(±1)이 중복 표시를 이기려면 맞아야 하는 문맥 글자 수
FAR_MAX_PAGES = 10          # 정본 쪽·±1쪽에 없을 때 ±2 ~ 이 거리까지 재탐색 (2026-09-26 연구 책임자 "재탐색"; 진단상 그 너머 적중 0)
FAR_MIN_CONTEXT = 8         # 먼 쪽은 앞뒤 문맥이 합쳐 이만큼은 맞아야 받아들인다 — 표현만으로는 절대 옮기지 않는다
FAR_OCR_MAX_PAGES = FAR_MAX_PAGES   # OCR 도 같은 거리까지 (2026-09-26 연구 책임자 "그림 속 글자 포함"; 진단: 개발 분야 33권의 앞부속 그림 쪽이 정본 쪽에서 3쪽)


@dataclass(frozen=True)
class Resolution:
    page: int          # 표시한 쪽 (1-based)
    start: int         # 스트림 안 정규화 글자 범위
    end: int
    k: int             # 실제로 함께 맞춘 앞뒤 문맥 글자 수 (0 = 표현만으로 맞춤 — 같은 표현이 여럿이면 순서로 배정된 것)
    shared: bool = False   # 정본 쪽의 인스턴스가 소진돼 이미 표시한 자리에 겹쳐 표시 (정본 출현 수 > PDF 본문 표현 수)
    far: bool = False      # 정본 쪽·±1쪽에 없어 ±2 ~ FAR_MAX_PAGES 쪽에서 강한 문맥(FAR_MIN_CONTEXT 자 이상)으로 찾음
    ocr: bool = False      # 텍스트 층에 없어 정본 쪽을 OCR 로 다시 읽어 찾음 (좌표는 OCR 스트림 기준) — NCS 대체 경로


def context_window(text: str, start: int, end: int, k: int) -> tuple[str, str, str]:
    pre, post = norm_chars(text[:start]), norm_chars(text[end:])
    return (pre[-k:] if k else "", norm_chars(text[start:end]), post[:k] if k else "")


def find_span(stream_text: str, pre: str, core: str, post: str, used: set[tuple[int, int]]) -> tuple[int, int] | None:
    """창(pre+core+post)의 첫 적중 중 아직 안 쓴 것의 core 범위."""
    if not core:
        return None
    needle = pre + core + post
    i = stream_text.find(needle)
    while i >= 0:
        span = (i + len(pre), i + len(pre) + len(core))
        if span not in used:
            return span
        i = stream_text.find(needle, i + 1)
    return None


def _window_ladder(text: str, start: int, end: int):
    """양쪽 문맥(10·6·3) → 뒤쪽만 → 앞쪽만 (10·6·3) → 표현만. 표 셀 경계처럼 한쪽 문맥만 이어지는 경우를 잡는다."""
    for k in CONTEXT_WINDOWS[:-1]:
        yield context_window(text, start, end, k)
    for k in CONTEXT_WINDOWS[:-1]:
        pre, core, post = context_window(text, start, end, k)
        yield ("", core, post)
        yield (pre, core, "")
    yield context_window(text, start, end, 0)


def resolve_on_page(stream: CharStream, text: str, start: int, end: int, used: set[tuple[int, int]]) -> tuple[int, int, int] | None:
    tried: set[tuple[str, str, str]] = set()
    for pre, core, post in _window_ladder(text, start, end):
        if (pre, core, post) in tried:
            continue
        tried.add((pre, core, post))
        span = find_span(stream.text, pre, core, post, used)
        if span:
            return (span[0], span[1], len(pre) + len(post))
    return None


def _try_page(stream: CharStream | None, spans: set, text: str, start: int, end: int, min_ctx: int):
    if stream is None:
        return None
    hit = resolve_on_page(stream, text, start, end, spans)
    return hit if hit and hit[2] >= min_ctx else None


def _ocr_hit_is_new(ocr_stream: CharStream, hit, text_stream: CharStream | None, text_spans: set | None) -> bool:
    """OCR 적중이 텍스트 층에서 이미 표시한 글자를 다시 읽은 것이면 새 발견이 아니다 — box 가 겹치면 같은 글자다.

    OCR 스트림은 텍스트 층의 글자도 함께 읽으므로, 텍스트 층 자리가 소진된 행이 OCR 경로에서 같은 글자를 "안 쓴 자리" 로 다시 찾을 수 있다.
    그런 적중은 겹쳐 표시(shared) 로 남겨야 집계가 뜻을 잃지 않는다."""
    if text_stream is None or not text_spans:
        return True
    boxes = ocr_stream.boxes[hit[0]:hit[1]]
    for i, j in text_spans:
        for text_box in text_stream.boxes[i:j]:
            if any(box.intersects(text_box) for box in boxes):
                return False
    return True


def resolve_row(get_stream, page: int, text: str, start: int, end: int, keyword: str, used: dict, get_ocr_stream=None) -> Resolution | None:
    """행 하나를 쪽 위에 놓는다 — 근거가 강한 순서로 (2026-09-26 진단·연구 책임자 "재탐색"·"그림 속 글자 포함" 반영):

    1. 정본 쪽, 안 쓴 자리 (문맥 불문)
    2. ±1쪽, 안 쓴 자리, 문맥 NEIGHBOR_MIN_CONTEXT 자 이상 — 줄→쪽 대응이 한 쪽 어긋난 행
    3. 정본 쪽, 표현은 있으나 자리가 소진 → 겹쳐 표시(shared) — 같은 문장의 사본이 한 쪽에 몰린 행
    4. ±1쪽, 안 쓴 자리 (문맥 불문) — 정본 쪽에 표현이 아예 없을 때
    5. 정본 쪽 OCR(텍스트 층에 없는 글자 — 표 셀·그림·스크린샷), 안 쓴 자리, 문맥 6자 이상 (get_ocr_stream 이 있을 때)
    6. ±1쪽 OCR, 안 쓴 자리, 문맥 6자 이상 — 이웃 쪽 그림 속 글자
    7. ±2 ~ FAR_MAX_PAGES 쪽 텍스트, 안 쓴 자리, 문맥 FAR_MIN_CONTEXT 자 이상 (far)
    8. ±2 ~ FAR_OCR_MAX_PAGES 쪽 OCR, 안 쓴 자리, 문맥 FAR_MIN_CONTEXT 자 이상 (far + ocr)
    9. 정본 쪽 OCR, 안 쓴 자리 (문맥 불문) — 정본 쪽 라벨 자체가 근거
    10. ±1쪽, 표현은 있으나 자리가 소진 → 겹쳐 표시(shared, 이웃 쪽)
    정본 쪽이 아닌 쪽에 문맥 없이 OCR 로 놓는 단계는 두지 않는다 — 앞부속의 예시 그림(네일숍 '위생 서비스', 정본 쪽에서 4~5쪽)이 있는 책은
    이웃 쪽 다른 그림의 '위생가운' 을 잡게 되기 때문(2026-09-26 진단). 그런 행은 미발견으로 남긴다.
    used 는 (쪽, 키워드) → 이미 표시한 범위; OCR 스트림은 ("ocr", 쪽, 키워드) 로 따로 센다 (키워드가 다르면 겹쳐도 된다).
    OCR 적중이 텍스트 층에서 이미 표시한 글자와 겹치면(_ocr_hit_is_new) 새 발견으로 치지 않고 다음 자리를 본다.
    """
    canonical = get_stream(page) if page >= 1 else None
    neighbours = [c for c in (page - 1, page + 1) if c >= 1]
    core = norm_chars(text[start:end])

    def take(candidate: int, hit, **flags) -> Resolution:
        used.setdefault((candidate, keyword), set()).add((hit[0], hit[1]))
        return Resolution(candidate, hit[0], hit[1], hit[2], **flags)

    def try_ocr(candidate: int, min_ctx: int, **flags) -> Resolution | None:
        if get_ocr_stream is None or candidate < 1:
            return None
        stream = get_ocr_stream(candidate)
        if stream is None:
            return None
        spans = used.setdefault(("ocr", candidate, keyword), set())
        while True:
            hit = _try_page(stream, spans, text, start, end, min_ctx)
            if hit is None:
                return None
            spans.add((hit[0], hit[1]))                                    # 새 발견이든 텍스트 층과 겹친 글자든, 이 OCR 자리는 쓴 것
            if _ocr_hit_is_new(stream, hit, get_stream(candidate), used.get((candidate, keyword))):
                return Resolution(candidate, hit[0], hit[1], hit[2], ocr=True, **flags)

    hit = _try_page(canonical, used.setdefault((page, keyword), set()), text, start, end, 0)                        # 1
    if hit:
        return take(page, hit)
    for candidate in neighbours:                                                                                       # 2
        hit = _try_page(get_stream(candidate), used.setdefault((candidate, keyword), set()), text, start, end, NEIGHBOR_MIN_CONTEXT)
        if hit:
            return take(candidate, hit)
    if canonical is not None and core and core in canonical.text:                                                      # 3
        hit = resolve_on_page(canonical, text, start, end, set())
        return Resolution(page, hit[0], hit[1], hit[2], shared=True)
    for candidate in neighbours:                                                                                       # 4
        hit = _try_page(get_stream(candidate), used.setdefault((candidate, keyword), set()), text, start, end, 0)
        if hit:
            return take(candidate, hit)
    resolution = try_ocr(page, NEIGHBOR_MIN_CONTEXT)                                                                  # 5
    if resolution:
        return resolution
    for candidate in neighbours:                                                                                       # 6
        resolution = try_ocr(candidate, NEIGHBOR_MIN_CONTEXT)
        if resolution:
            return resolution
    for distance in range(2, FAR_MAX_PAGES + 1):                                                                       # 7
        for candidate in (page - distance, page + distance):
            if candidate < 1:
                continue
            hit = _try_page(get_stream(candidate), used.setdefault((candidate, keyword), set()), text, start, end, FAR_MIN_CONTEXT)
            if hit:
                return take(candidate, hit, far=True)
    for distance in range(2, FAR_OCR_MAX_PAGES + 1):                                                                   # 8
        for candidate in (page - distance, page + distance):
            resolution = try_ocr(candidate, FAR_MIN_CONTEXT, far=True)
            if resolution:
                return resolution
    resolution = try_ocr(page, 0)                                                                                      # 9
    if resolution:
        return resolution
    for candidate in neighbours:                                                                                       # 10
        stream = get_stream(candidate)
        if stream is not None and core and core in stream.text:
            hit = resolve_on_page(stream, text, start, end, set())
            return Resolution(candidate, hit[0], hit[1], hit[2], shared=True)
    return None


# ----------------------------------------------------------------------------- 형광펜 주석
LEGEND_TEXT = ("형광펜 범례 — 노랑: 등급1 미흡·없음 · 주황: 등급2 형식적 언급 · 초록: 등급3 구체적 대책\n"
               "정본 {source} 의 출현 키워드를 그 쪽의 등급 색으로 표시. 형광펜 팝업에 키워드·등급·정본 쪽.")


def quads_for(stream: CharStream, i: int, j: int) -> list[fitz.Quad]:
    """스트림 [i, j) 글자의 box 를 줄별로 합친 quad 목록 (줄이 바뀌면 새 quad)."""
    quads: list[fitz.Quad] = []
    current_line, current_rect = None, None
    for idx in range(i, j):
        line_id, box = stream.line_ids[idx], stream.boxes[idx]
        if line_id != current_line:
            if current_rect is not None:
                quads.append(current_rect.quad)
            current_line, current_rect = line_id, fitz.Rect(box)
        else:
            current_rect |= box
    if current_rect is not None:
        quads.append(current_rect.quad)
    return quads


def add_highlight(page: fitz.Page, quads: list[fitz.Quad], row: Row, shown_page: int, shared: bool = False, far: bool = False,
                  ocr: bool = False) -> fitz.Annot:
    annot = page.add_highlight_annot(quads)
    annot.set_colors(stroke=GRADE_COLORS[row.grade])
    content = f"등록 표현: {row.expression} / 실제 매칭: {row.matched} / 계층: {row.tier}\n정본 쪽 {row.page}, 마크다운 {row.line}줄"
    if far:
        content += (f"\n표시 쪽 {shown_page} — 정본 쪽 {row.page}에서 {abs(shown_page - row.page)}쪽 떨어진 곳에서 앞뒤 문맥으로 찾음"
                    " (등급은 정본 쪽의 것)")
    elif shown_page != row.page:
        content += f"\n표시 쪽 {shown_page} — 정본 쪽에서 못 찾아 이웃 쪽에 표시 (등급은 정본 쪽의 것)"
    if shared:
        content += "\n같은 자리 중복 표시 — 정본 출현 수가 이 쪽 PDF 본문의 표현 수보다 많음 (표·병합 셀 중복 가능)"
    if ocr:
        content += "\nPDF 텍스트 층에 없어 이 쪽을 OCR 로 다시 읽어 찾음 (좌표는 OCR 기준)"
    annot.set_info(title=f"{row.keyword} · 등급{row.grade} {row.grade_label}", content=content)
    annot.update()
    return annot


def add_legend(page: fitz.Page, source_label: str) -> fitz.Annot:
    r = page.rect
    rect = fitz.Rect(r.x0 + 8, r.y0 + 8, min(r.x1 - 8, r.x0 + 440), r.y0 + 44)
    annot = page.add_freetext_annot(rect, LEGEND_TEXT.format(source=source_label), fontsize=8, fontname="korea",
                                    text_color=(0, 0, 0), fill_color=(1, 0.98, 0.85))          # border_color 는 rich_text 전용 (PyMuPDF 1.27)
    annot.set_info(title="범례")
    annot.update()
    return annot


# ----------------------------------------------------------------------------- Vision OCR 도구 (교과서 스캔본)
OCR_SWIFT_SOURCE = HERE / "outputs" / "vision_ocr_chars.swift"
OCR_BATCH = 16


class OcrError(RuntimeError):
    """Vision OCR 도구의 빌드·실행 실패. 실행 중에 나면 그 권만 건너뛰고 나머지는 계속한다(설계 §5)."""


def build_ocr_tool(cache_dir: Path, source: Path = OCR_SWIFT_SOURCE) -> Path:
    """swiftc 로 빌드해 cache_dir 에 둔다. 소스보다 새 바이너리가 있으면 그대로 쓴다."""
    cache_dir = Path(cache_dir)
    cache_dir.mkdir(parents=True, exist_ok=True)
    binary = cache_dir / "vision_ocr_chars"
    if binary.is_file() and binary.stat().st_mtime_ns >= source.stat().st_mtime_ns:
        return binary
    swiftc = shutil.which("swiftc")
    if not swiftc:
        raise OcrError("swiftc 가 없습니다 — 교과서 OCR 은 macOS Swift 툴체인이 필요합니다 (xcode-select --install)")
    tmp = binary.with_name(binary.name + ".tmp")
    proc = subprocess.run([swiftc, "-O", "-o", str(tmp), str(source)], capture_output=True, text=True)
    if proc.returncode != 0:
        raise OcrError(f"OCR 도구 빌드 실패: {proc.stderr.strip()[:500]}")
    os.replace(tmp, binary)
    return binary


def run_ocr_tool(binary: Path, images: list[Path]) -> list[dict]:
    proc = subprocess.run([str(binary)] + [str(p) for p in images], capture_output=True)
    if proc.returncode != 0:
        raise OcrError(f"OCR 도구 실행 실패 ({proc.returncode}): {proc.stderr.decode('utf-8', 'replace')[:500]}")
    try:
        results = json.loads(proc.stdout)
    except ValueError as exc:
        raise OcrError(f"OCR 도구 출력이 JSON 이 아닙니다: {exc}") from exc
    if not isinstance(results, list):
        raise OcrError(f"OCR 도구 출력이 목록이 아닙니다: {type(results).__name__}")
    return results


def ocr_pages(doc: fitz.Document, indices: list[int], dpi: int, cache_dir: Path, binary: Path, stem: str) -> dict[int, dict]:
    """쪽 인덱스(0-based) → OCR JSON. cache_dir/ocr/<stem>/<idx>.json 에 없는(또는 dpi 가 다른) 쪽만 렌더해 도구를 부른다."""
    folder = Path(cache_dir) / "ocr" / stem
    folder.mkdir(parents=True, exist_ok=True)
    out: dict[int, dict] = {}
    missing: list[int] = []
    for idx in indices:
        f = folder / f"{idx}.json"
        if f.is_file():
            cached = json.loads(f.read_text(encoding="utf-8"))
            if cached.get("dpi") == dpi:
                out[idx] = cached
                continue
        missing.append(idx)
    if not missing:
        return out
    with tempfile.TemporaryDirectory(prefix="ocr_") as td:
        for batch_start in range(0, len(missing), OCR_BATCH):
            batch = missing[batch_start:batch_start + OCR_BATCH]
            pngs = []
            for idx in batch:
                png = Path(td) / f"{idx}.png"
                doc[idx].get_pixmap(dpi=dpi).save(str(png))
                pngs.append(png)
            results = run_ocr_tool(binary, pngs)
            if len(results) != len(batch):
                raise OcrError(f"OCR 결과 수가 다릅니다: {stem} {len(results)} != {len(batch)}")
            for idx, result in zip(batch, results):
                if "error" in result:
                    raise OcrError(f"OCR 실패: {stem} {idx + 1}쪽: {result['error']}")
                result["dpi"] = dpi
                (folder / f"{idx}.json").write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
                out[idx] = result
            for png in pngs:
                png.unlink(missing_ok=True)
    return out


# ----------------------------------------------------------------------------- 파이프라인
DETAIL_SHEETS = (("NCS", "NCS_매칭상세"), ("교과서", "교과서_매칭상세"))
DETAIL_COLUMNS = ("파일", "키워드", "등록 표현", "실제 매칭", "계층", "줄", "페이지", "문맥", "통일 등급", "등급명")
DEFAULT_XLSX = HERE / "data" / "semantic_keyword_recount_20260917.xlsx"
DEFAULT_SUMMARY = HERE / "docs" / "03-analysis" / "data" / "semantic_summary.json"


def load_rows(xlsx: Path) -> list[Row]:
    import openpyxl

    workbook = openpyxl.load_workbook(str(xlsx), read_only=True)
    rows: list[Row] = []
    for corpus, sheet in DETAIL_SHEETS:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"{xlsx.name} 에 {sheet} 시트가 없습니다")
        it = workbook[sheet].iter_rows(values_only=True)
        header = next(it, ())
        col = {h: i for i, h in enumerate(header)}
        missing = [h for h in DETAIL_COLUMNS if h not in col]
        if missing:
            raise ValueError(f"{sheet}: 열이 없습니다 {missing}")
        for values in it:
            if values[col["파일"]] is None:
                continue
            rows.append(Row(corpus, nfc(str(values[col["파일"]])), str(values[col["키워드"]]), str(values[col["등록 표현"]]),
                            str(values[col["계층"]]), str(values[col["실제 매칭"]]), int(values[col["줄"]]), int(values[col["페이지"]]),
                            nfc(str(values[col["문맥"]] or "")), int(values[col["통일 등급"]]), str(values[col["등급명"]]), len(rows)))
    workbook.close()
    return rows


def check_corpus(summary_path: Path, xlsx: Path, ncs_docs: list, school_docs: list, dictionary: str) -> dict:
    """정본 manifest(semantic_summary.json meta.run)와 입력이 같은지 — 마크다운 두 세트의 sha256, 사전 버전, xlsx 이름."""
    run = json.loads(Path(summary_path).read_text(encoding="utf-8"))["meta"]["run"]
    want = {item["kind"]: item["sha256"] for item in run.get("inputs", [])}
    got = {"NCS Markdown": skr._document_set_sha256(ncs_docs), "교과서 Markdown": skr._document_set_sha256(school_docs)}
    bad = [kind for kind, sha in got.items() if want.get(kind) != sha]
    if bad:
        raise ValueError("마크다운 코퍼스가 정본 실행과 다릅니다(줄 안 위치 재현이 어긋날 수 있음): " + ", ".join(bad) + " — --skip-corpus-check 로 무시")
    if run.get("dictionary") != dictionary:
        raise ValueError(f"사전 버전이 정본({run.get('dictionary')})과 다릅니다: {dictionary}")
    m = re.search(r"--xlsx-out[= ]([^ ]+)", run.get("command", ""))
    if m and os.path.basename(m.group(1)) != xlsx.name:
        raise ValueError(f"정본 실행이 쓴 xlsx 는 {os.path.basename(m.group(1))} 인데 {xlsx.name} 을 받았습니다 — --skip-corpus-check 로 무시")
    return run


@dataclass
class BookStats:
    corpus: str
    relpath: str
    pdf: str
    out: str
    rows: int = 0
    same_page: int = 0
    shared: int = 0
    adjacent: int = 0
    far: int = 0
    ocr: int = 0               # OCR 로 찾은 행 전부 (정본 쪽 + 아래 둘)
    ocr_adjacent: int = 0      # 그중 ±1쪽 OCR
    ocr_far: int = 0           # 그중 ±2~FAR_OCR_MAX_PAGES 쪽 OCR (문맥 FAR_MIN_CONTEXT 자 이상)
    unresolved: int = 0
    locator: dict = field(default_factory=lambda: {"engine": 0, "naive": 0})
    context: dict = field(default_factory=lambda: {"0": 0, "1-3": 0, "4-9": 0, "10+": 0})
    pdf_pages: int = 0
    ocr_pages: int = 0
    preexisting_annots: dict = field(default_factory=dict)     # 원본 PDF 에 이미 있던 주석 (종류별 수) — 그대로 둔다
    seconds: float = 0.0
    unresolved_rows: list = field(default_factory=list)

    def as_json(self) -> dict:
        d = {k: v for k, v in self.__dict__.items() if k != "unresolved_rows"}
        return d


@dataclass
class FailedBook:
    """OCR 실패로 건너뛴 권. 그 출현은 미발견 목록에 사유와 함께 남긴다."""
    corpus: str
    relpath: str
    pdf: str
    rows: list
    error: str

    def as_json(self) -> dict:
        return {"corpus": self.corpus, "relpath": self.relpath, "pdf": self.pdf, "rows": len(self.rows), "error": self.error}


FAILED_REASON = "OCR 실패로 권 건너뜀"


def _context_bucket(k: int) -> str:
    return "0" if k == 0 else "1-3" if k <= 3 else "4-9" if k <= 9 else "10+"


def process_book(corpus: str, relpath: str, rows: list[Row], document: skr.Document, compiled: dict, pdf: Path, out: Path,
                 dpi: int, cache_dir: Path, binary: Path | None, source_label: str) -> BookStats:
    """한 권: 줄 안 오프셋 → 사본 복사 → 쪽 스트림(rawdict / OCR) → 행마다 해석·형광펜 → 범례 → 증분 저장."""
    t0 = time.time()
    stats = BookStats(corpus, relpath, str(pdf), str(out), rows=len(rows))
    raw_lines = dict(enumerate(document.text.splitlines(), start=1))
    offsets = line_offsets(document, compiled, wanted_lines={r.line for r in rows})
    stats.locator = assign_offsets(rows, offsets, raw_lines)

    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(pdf, out)
    doc = fitz.open(str(out))
    try:
        stats.pdf_pages = doc.page_count
        for index in range(doc.page_count):
            for annot in doc[index].annots():
                stats.preexisting_annots[annot.type[1]] = stats.preexisting_annots.get(annot.type[1], 0) + 1
        streams: dict[int, CharStream] = {}
        ocr_streams: dict[int, CharStream] = {}
        stem = pdf.stem

        def get_stream(page_no: int) -> CharStream | None:
            idx = page_no - 1
            if idx < 0 or idx >= doc.page_count:
                return None
            if page_no not in streams:
                page = doc[idx]
                if corpus == "교과서":
                    result = ocr_pages(doc, [idx], dpi, cache_dir, binary, stem)[idx]
                    stats.ocr_pages += 1
                    streams[page_no] = CharStream.from_ocr(result, page)
                else:
                    streams[page_no] = CharStream.from_rawdict(page)
            return streams[page_no]

        def get_ocr_stream(page_no: int) -> CharStream | None:          # NCS 대체 경로: 텍스트 층에 없는 글자를 OCR 로 (binary 가 있을 때만)
            idx = page_no - 1
            if binary is None or idx < 0 or idx >= doc.page_count:
                return None
            if page_no not in ocr_streams:
                result = ocr_pages(doc, [idx], dpi, cache_dir, binary, stem)[idx]
                stats.ocr_pages += 1
                ocr_streams[page_no] = CharStream.from_ocr(result, doc[idx])
            return ocr_streams[page_no]

        if corpus == "교과서":
            canonical = sorted({r.page - 1 for r in rows if 0 <= r.page - 1 < doc.page_count})
            ocr_pages(doc, canonical, dpi, cache_dir, binary, stem)          # 정본 쪽은 한꺼번에 (이웃 쪽은 필요할 때만)

        used: dict = {}
        for row in sorted(rows, key=lambda r: r.order):
            resolution = None
            if row.start is not None and row.text is not None:
                resolution = resolve_row(get_stream, row.page, row.text, row.start, row.end, row.keyword, used,
                                         get_ocr_stream=get_ocr_stream if corpus == "NCS" else None)
            if resolution is None:
                stats.unresolved += 1
                stats.unresolved_rows.append(row)
                continue
            source = ocr_streams if resolution.ocr else streams
            add_highlight(doc[resolution.page - 1], quads_for(source[resolution.page], resolution.start, resolution.end), row, resolution.page,
                          shared=resolution.shared, far=resolution.far, ocr=resolution.ocr)
            if resolution.shared:
                stats.shared += 1
            elif resolution.ocr:
                stats.ocr += 1
                if resolution.far:
                    stats.ocr_far += 1
                elif resolution.page != row.page:
                    stats.ocr_adjacent += 1
            elif resolution.far:
                stats.far += 1
            elif resolution.page == row.page:
                stats.same_page += 1
            else:
                stats.adjacent += 1
            stats.context[_context_bucket(resolution.k)] += 1

        add_legend(doc[0], source_label)
        try:
            doc.save(str(out), incremental=True, encryption=fitz.PDF_ENCRYPT_KEEP)
        except Exception:                                                     # 증분 저장이 안 되는 파일(손상된 xref 등) — 전체 저장으로
            tmp = out.with_name(out.name + ".tmp")
            doc.save(str(tmp), garbage=0)
            doc.close()
            os.replace(tmp, out)
        else:
            doc.close()
    except BaseException:
        if not doc.is_closed:
            doc.close()
        out.unlink(missing_ok=True)                                      # 형광펜이 덜 들어간 사본은 남기지 않는다
        raise
    stats.seconds = round(time.time() - t0, 1)
    return stats


def _unresolved_reason(row: Row, pdf_pages: int) -> str:
    if row.start is None:
        return "줄 안 위치 없음"
    if row.page > pdf_pages:
        return "PDF 쪽수 초과"
    return "정본 쪽·이웃 쪽에서 미발견"


def write_outputs(out_root: Path, books: list[BookStats], totals: dict, meta: dict, partial: bool,
                  failed: list[FailedBook] | tuple = ()) -> None:
    out_root.mkdir(parents=True, exist_ok=True)
    log = {"generated_at": meta["generated_at"], "xlsx": meta["xlsx"], "xlsx_sha256": meta["xlsx_sha256"], "dpi": meta["dpi"],
           "colors": GRADE_HEX, "canonical_run": meta.get("canonical_run"), "partial": partial, "totals": totals,
           "books": [b.as_json() for b in books], "failed_books": [f.as_json() for f in failed]}
    name = "highlight_log.partial.json" if partial else "highlight_log.json"
    (out_root / name).write_text(json.dumps(log, ensure_ascii=False, indent=1), encoding="utf-8")
    with (out_root / ("unresolved.partial.csv" if partial else "unresolved.csv")).open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(("말뭉치", "파일", "키워드", "등록 표현", "실제 매칭", "줄", "정본 쪽", "등급", "위치 방식", "사유"))
        for b in books:
            for row in b.unresolved_rows:
                writer.writerow((row.corpus, row.relpath, row.keyword, row.expression, row.matched, row.line, row.page, row.grade,
                                 row.locator or "", _unresolved_reason(row, b.pdf_pages)))
        for f in failed:
            for row in f.rows:
                writer.writerow((row.corpus, row.relpath, row.keyword, row.expression, row.matched, row.line, row.page, row.grade,
                                 row.locator or "", FAILED_REASON))
    if not partial:
        (out_root / "README.md").write_text(readme_text(totals, meta, books, failed), encoding="utf-8")


TOTAL_KEYS = ("books", "rows", "same_page", "shared", "adjacent", "far", "ocr", "ocr_adjacent", "ocr_far", "unresolved")


def book_totals(books: list[BookStats]) -> dict:
    return {"books": len(books), **{k: sum(getattr(b, k) for b in books) for k in TOTAL_KEYS[1:]}}


def readme_text(totals: dict, meta: dict, books: list[BookStats], failed: list[FailedBook] | tuple = ()) -> str:
    by_corpus: dict[str, dict] = {}
    for b in books:
        c = by_corpus.setdefault(b.corpus, {k: 0 for k in TOTAL_KEYS})
        c["books"] += 1
        for k in TOTAL_KEYS[1:]:
            c[k] += getattr(b, k)
    lines = [
        "# 출현 키워드 형광펜 PDF",
        "",
        f"- 생성 {meta['generated_at']} · 정본 `{meta['xlsx']}` (sha256 `{meta['xlsx_sha256'][:16]}…`) · 도구 `highlight_pdf_occurrences.py`",
        "- 원본 PDF 는 건드리지 않았다. 사본(`…_키워드표시.pdf`)에 형광펜·범례 **주석**만 얹었다 — 뷰어에서 주석을 숨기면 원본 그대로다.",
        "- 등급은 쪽 속성이라 한 쪽의 형광펜은 모두 같은 색이다. 형광펜 팝업(제목)에 `키워드 · 등급N 등급명`, 내용에 등록 표현·실제 매칭·정본 쪽.",
        "",
        "| 등급 | 뜻 | 형광펜 |",
        "|:---:|---|---|",
        f"| 1 | 미흡·없음 | 노랑 `{GRADE_HEX[1]}` |",
        f"| 2 | 형식적 언급 | 주황 `{GRADE_HEX[2]}` |",
        f"| 3 | 구체적 대책 | 초록 `{GRADE_HEX[3]}` |",
        "",
        f"| 말뭉치 | 권 | 출현 | 정본 쪽 자기 자리 | 중복 자리 | 이웃 쪽(±1) | 먼 쪽(±2~{FAR_MAX_PAGES}) | OCR 대체 | 미발견 |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for corpus, c in by_corpus.items():
        lines.append(f"| {corpus} | {c['books']} | {c['rows']:,} | {c['same_page']:,} | {c['shared']:,} | {c['adjacent']:,} | {c['far']:,} | {c['ocr']:,} | {c['unresolved']:,} |")
    lines += [
        f"| 합계 | {totals['books']} | {totals['rows']:,} | {totals['same_page']:,} | {totals['shared']:,} | {totals['adjacent']:,} | {totals['far']:,} | {totals['ocr']:,} | {totals['unresolved']:,} |",
        "",
    ]
    if failed:
        lines += [f"**OCR 실패로 만들지 못한 권 {len(failed)}권** — 출현 {sum(len(f.rows) for f in failed):,}건은 위 표에 없고 `unresolved.csv` 에 "
                  f"사유 `{FAILED_REASON}` 로, 오류는 `highlight_log.json` 의 `failed_books` 에 있다. 이 실행은 완결본이 아니다.", ""]
        lines += [f"- {f.corpus} `{Path(f.pdf).name}` (출현 {len(f.rows):,}건)" for f in failed]
        lines.append("")
    lines += [
        "- NCS: PDF 텍스트 층(PyMuPDF rawdict)에서 마크다운 줄의 앞뒤 문맥 창으로 위치를 찾았다.",
        f"- 교과서: 스캔본이라 Apple Vision OCR({meta['dpi']} dpi, `outputs/vision_ocr_chars.swift`)로 글자 좌표를 얻었다. 마크다운 변환기(surya)와 다른 OCR 이라 일부는 못 찾는다.",
        "- 정본 쪽 중복 자리: 정본 출현 수가 그 쪽 PDF 본문의 표현 수보다 많아(표·병합 셀 중복 등) 이미 표시한 자리에 겹쳐 표시한 출현 — 팝업에 적혀 있다.",
        "- 이웃 쪽에 표시된 출현: 정본 쪽에 표현이 아예 없고(줄→쪽 대응이 ±1쪽 어긋남) 이웃 쪽에 있는 경우 — 팝업에 `표시 쪽 M` 을 적었고 색은 정본 쪽 등급이다.",
        f"- 먼 쪽에 표시된 출현: 정본 쪽·±1쪽에 없어 ±2~{FAR_MAX_PAGES}쪽을 앞뒤 문맥 {FAR_MIN_CONTEXT}자 이상이 맞을 때만 받아들여 찾은 경우 — 팝업에 거리를 적었다.",
        f"- OCR 대체(NCS): PDF 텍스트 층에 없는 글자(표 셀·그림·스크린샷 속 글자)를 Vision OCR 로 다시 읽어 찾은 경우 — 정본 쪽 {totals['ocr'] - totals['ocr_adjacent'] - totals['ocr_far']:,} · "
        f"±1쪽 {totals['ocr_adjacent']:,} · 먼 쪽(±2~{FAR_OCR_MAX_PAGES}, 문맥 {FAR_MIN_CONTEXT}자 이상) {totals['ocr_far']:,}. 좌표는 OCR 기준, 팝업에 적었다. "
        "그림·스크린샷 속 글자도 출현으로 센다(연구 책임자, 2026-09-26) — 변환기가 그 글자를 본문처럼 넣었고 정본은 그것을 셌기 때문이다.",
        "- 중복 자리 가운데 이웃 쪽에 겹친 것: 정본 쪽에 표현이 없고 이웃 쪽의 표현 자리가 모두 소진된 경우(최후 수단) — 팝업에 두 사유가 함께 적혀 있다.",
        "- 미발견 출현은 `unresolved.csv` (본문 문맥 없음). 권별 수치는 `highlight_log.json` — 원본에 이미 있던 주석은 그대로 두었고 `preexisting_annots` 에 종류별 수를 적었다.",
        "",
        "재생성: `python3.13 highlight_pdf_occurrences.py` (기본 경로) — OCR 결과는 `.cache/ocr/` 에 남아 재실행이 빠르다.",
        "",
    ]
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="정본 xlsx 의 출현 키워드를 원본 PDF 사본에 등급별 형광펜으로 표시")
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--ncs-pdf-root", type=Path, default=HERE / "data" / "markdown" / "ncs" / "pdf")
    ap.add_argument("--textbook-pdf-root", type=Path, default=HERE / "data" / "markdown" / "school-text" / "전공교과서(pdf)")
    ap.add_argument("--ncs-md-root", type=Path, default=HERE / "data_source" / "markdown" / "ncs")
    ap.add_argument("--school-md-root", type=Path, default=HERE / "data_source" / "markdown" / "school-text")
    ap.add_argument("--out", type=Path, default=HERE / "data" / "highlighted")
    ap.add_argument("--summary", type=Path, default=DEFAULT_SUMMARY, help="정본 manifest — 마크다운 sha·사전·xlsx 이름을 대조")
    ap.add_argument("--skip-corpus-check", action="store_true")
    ap.add_argument("--dictionary", default=skr.DEFAULT_DICTIONARY)
    ap.add_argument("--only", action="append", default=[], help="파일 경로에 이 문자열이 든 권만 (부분 실행 — 로그는 .partial)")
    ap.add_argument("--dpi", type=int, default=200)
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args(argv)

    def say(msg: str) -> None:
        if not args.quiet:
            print(msg, flush=True)

    rows = load_rows(args.xlsx)
    ncs_docs, _dedup = skr.select_ncs_documents(skr.load_documents(args.ncs_md_root, "NCS"))
    school_docs = skr.load_documents(args.school_md_root, "교과서")
    canonical_run = None
    if not args.skip_corpus_check:
        run = check_corpus(args.summary, args.xlsx, ncs_docs, school_docs, args.dictionary)
        canonical_run = {k: run.get(k) for k in ("generated_at", "dictionary", "git_commit")}
    documents = {(d.corpus, d.relative_path): d for d in ncs_docs + school_docs}
    compiled = compile_rules(skr.build_default_rules(list(skr.EXPECTED_KEYWORDS), version=args.dictionary))
    ncs_index, textbook_index = index_ncs_pdfs(args.ncs_pdf_root), index_textbook_pdfs(args.textbook_pdf_root)

    grouped: dict[tuple[str, str], list[Row]] = {}
    for row in rows:
        grouped.setdefault((row.corpus, row.relpath), []).append(row)
    partial = bool(args.only)
    plan = []
    for (corpus, relpath), book_rows in grouped.items():
        if args.only and not any(token in relpath for token in args.only):
            continue
        document = documents.get((corpus, relpath))
        if document is None:
            raise LookupError(f"마크다운이 없습니다: {corpus} {relpath}")
        pdf = pdf_for(corpus, relpath, ncs_index, textbook_index)
        root = args.ncs_pdf_root if corpus == "NCS" else args.textbook_pdf_root
        plan.append((corpus, relpath, book_rows, document, pdf, output_path(corpus, pdf, root, args.out)))

    cache_dir = args.out / ".cache"
    try:
        binary = build_ocr_tool(cache_dir / "bin")                    # 교과서 필수, NCS 는 대체 경로(텍스트 층에 없는 글자)
    except RuntimeError as exc:
        if any(p[0] == "교과서" for p in plan):
            raise
        binary = None
        say(f"OCR 도구 없음 — NCS 의 OCR 대체 경로는 건너뜁니다: {exc}")
    source_label = args.xlsx.name
    books: list[BookStats] = []
    failed: list[FailedBook] = []
    for n, (corpus, relpath, book_rows, document, pdf, out) in enumerate(plan, start=1):
        try:
            stats = process_book(corpus, relpath, book_rows, document, compiled, pdf, out, args.dpi, cache_dir, binary, source_label)
        except OcrError as exc:                                          # 설계 §5: Vision 실패는 그 권만 건너뛰고 나머지는 계속
            failed.append(FailedBook(corpus, relpath, str(pdf), book_rows, str(exc)[:500]))
            say(f"[{n}/{len(plan)}] {corpus} {pdf.name}: OCR 실패로 건너뜀 — {exc}")
            continue
        books.append(stats)
        say(f"[{n}/{len(plan)}] {corpus} {pdf.name}: 출현 {stats.rows} → 그 쪽 {stats.same_page} (+중복 {stats.shared}), ±1쪽 {stats.adjacent}, 먼 쪽 {stats.far}, OCR {stats.ocr}, 미발견 {stats.unresolved}"
            f" (engine {stats.locator['engine']}줄 / naive {stats.locator['naive']}줄{', OCR ' + str(stats.ocr_pages) + '쪽' if stats.ocr_pages else ''}, {stats.seconds}s)")

    totals = book_totals(books)
    meta = {"generated_at": dt.datetime.now().astimezone().isoformat(timespec="seconds"), "xlsx": skr.public_path(args.xlsx),
            "xlsx_sha256": hashlib.sha256(args.xlsx.read_bytes()).hexdigest(), "dpi": args.dpi, "canonical_run": canonical_run}
    write_outputs(args.out, books, totals, meta, partial, failed)
    complete = not partial and not failed and totals["rows"] == len(rows) and totals["books"] == len(grouped)
    say(f"합계: {totals['books']}권 {totals['rows']:,}건 — 그 쪽 {totals['same_page']:,} (+중복 {totals['shared']:,}), ±1쪽 {totals['adjacent']:,}, 먼 쪽 {totals['far']:,}, OCR {totals['ocr']:,}, 미발견 {totals['unresolved']:,}"
        + (f" · OCR 실패로 건너뛴 권 {len(failed)}권(출현 {sum(len(f.rows) for f in failed):,}건)" if failed else "")
        + ("" if complete else " (완결본 아님 — 정본 합계와 대조하지 않음)"))
    return 0 if complete else 1


if __name__ == "__main__":
    sys.exit(main())
